from io import BytesIO

from docx import Document
from openpyxl import Workbook
from pydantic import BaseModel

from utils.sop_uplift.analysis_engine import generate_rule_based_suggestions, quality_gate_suggestions
from utils.sop_uplift.case_index import build_case_local_index, search_case_local_index
from utils.sop_uplift.case_chat import capture_context_from_message
from utils.sop_uplift.corpus_map import build_case_corpus_map
from utils.sop_uplift.extraction_router import plan_extraction
from utils.sop_uplift.extractors import extract_sop_structure
from utils.sop_uplift.markdown_ingestion import convert_bytes_to_markdown, looks_corrupt_markdown
from utils.sop_uplift.preview_renderer import build_preview_model
from utils.sop_uplift.retrieval import retrieve_context
from utils.sop_uplift.schema_validation import validate_payload
from utils.sop_uplift.tagging import suggest_document_tag
from utils.sop_uplift.task_store import InMemoryTaskStore


class DemoSchema(BaseModel):
    name: str


def test_planned_sop_uplift_modules_expose_deterministic_v1_behaviors():
    conversion = convert_bytes_to_markdown(b"# SOP\n\nOwner reviews access.", "access_sop.md")
    tag = suggest_document_tag({"filename": "access_sop.md", "bucket": "sops"}, conversion.markdown)
    extraction_plan = plan_extraction(tag["confirmed_tag"])
    sop = extract_sop_structure(conversion.markdown, anchors=[{"anchor_id": "a1", "text": "Owner reviews access."}])
    corpus = build_case_corpus_map(
        {"process_name": "Access reviews"},
        document_tags=[tag],
        controls=[],
        risks=[],
        requirements=[],
        evidence_items=[],
        issues_findings=[],
        diagram_references=[],
        case_chat_context=[],
    )
    retrieved = retrieve_context("owner access", [{"chunk_id": "c1", "content": "Owner reviews access quarterly.", "anchor_ids": ["a1"]}])
    suggestions = generate_rule_based_suggestions([{"anchor_id": "a1", "text": "Operations reviews access exceptions before user access is approved."}])
    preview = build_preview_model([{"document_id": "d1", "markdown": conversion.markdown}], [{"anchor_id": "a1", "text": "Owner reviews access."}], suggestions)
    captured = capture_context_from_message("Operations Risk reviews exceptions weekly.")
    validated = validate_payload({"name": "ok"}, DemoSchema)
    tasks = InMemoryTaskStore()
    task = tasks.create("case-1", "analysis")

    assert conversion.status == "converted"
    assert tag["confirmed_tag"] == "sop"
    assert extraction_plan["extractors"] == ["sop_structure", "procedure_steps"]
    assert sop["sections"][0]["anchor_id"] == "a1"
    assert corpus["primary_process"] == "Access reviews"
    assert retrieved[0]["chunk_id"] == "c1"
    assert suggestions[0]["anchor_id"] == "a1"
    assert preview["highlights"][0]["anchor_id"] == "a1"
    assert captured["type"] == "frequency"
    assert validated.name == "ok"
    assert tasks.get(task["task_id"])["status"] == "running"


def test_document_tagging_uses_strong_filename_signals_before_bucket_fallback():
    tag = suggest_document_tag(
        {"file_id": "rcm-1", "filename": "WM_Risk_Controls_Matrix.xlsx", "bucket": "sops"},
        "Control ID | Risk ID | Frequency",
    )

    assert tag["confirmed_tag"] == "risk_control_matrix"
    assert tag["confidence"] == "medium"


def test_rule_based_suggestions_are_selective_and_use_supporting_context():
    suggestions = generate_rule_based_suggestions(
        [
            {
                "anchor_id": "sop-1",
                "file_id": "sop-file",
                "text": "Operations reviews high risk client onboarding exceptions before account opening.",
                "section_path": ["Risk assessment"],
            },
            {
                "anchor_id": "rcm-1",
                "file_id": "rcm-file",
                "text": "Risk ID | Control ID | Frequency | Evidence | Owner",
                "section_path": ["Control matrix"],
            },
        ],
        context_chunks=[
            {
                "chunk_id": "ctx-1",
                "document_id": "doc-rcm",
                "content": "Control C-12 requires Compliance approval evidence for high risk onboarding exceptions.",
                "anchor_ids": ["rcm-anchor-1"],
            }
        ],
    )

    assert len(suggestions) == 1
    assert suggestions[0]["anchor_id"] == "sop-1"
    assert suggestions[0]["title"] != "Make SOP step testable"
    assert suggestions[0]["source_references"][0]["anchor_id"] == "sop-1"
    assert suggestions[0]["source_references"][1]["anchor_id"] == "rcm-anchor-1"
    assert "Compliance approval evidence" in suggestions[0]["suggested_text"]
    assert "Source text:" not in suggestions[0]["suggested_text"]


def test_rule_based_suggestions_use_supporting_raci_for_owner_language():
    suggestions = generate_rule_based_suggestions(
        [
            {
                "anchor_id": "sop-1",
                "file_id": "sop-file",
                "text": "Review source of funds and source of wealth before onboarding high risk clients.",
                "section_path": ["KYC assessment"],
            }
        ],
        context_chunks=[
            {
                "chunk_id": "ctx-1",
                "document_id": "doc-raci",
                "content": "RACI matrix: Responsible - Relationship Manager; Accountable - Branch Operations; Consulted - Compliance.",
                "anchor_ids": ["raci-anchor-1"],
            }
        ],
    )

    assert suggestions
    assert suggestions[0]["type"] == "ownership_gap"
    assert "Relationship Manager shall" in suggestions[0]["suggested_text"]
    assert suggestions[0]["source_references"][1]["anchor_id"] == "raci-anchor-1"


def test_rule_based_suggestions_use_supporting_frequency_in_insert_ready_language():
    suggestions = generate_rule_based_suggestions(
        [
            {
                "anchor_id": "sop-1",
                "file_id": "sop-file",
                "text": "Owner reviews access control evidence and risk exceptions.",
                "section_path": ["Access review"],
            }
        ],
        context_chunks=[
            {
                "chunk_id": "case-context-1",
                "document_id": "case-chat",
                "content": "Case context frequency: weekly.",
                "anchor_ids": ["case-context-1"],
            }
        ],
    )

    assert suggestions
    assert suggestions[0]["type"] == "frequency_gap"
    assert "weekly" in suggestions[0]["suggested_text"]
    assert "before the process advances" not in suggestions[0]["suggested_text"]


def test_rule_based_suggestions_do_not_embed_raw_markdown_tables_in_language():
    suggestions = generate_rule_based_suggestions(
        [
            {
                "anchor_id": "sop-1",
                "file_id": "sop-file",
                "text": "Review source of funds and source of wealth before onboarding clients.",
                "section_path": ["KYC assessment"],
            }
        ],
        context_chunks=[
            {
                "chunk_id": "ctx-1",
                "document_id": "doc-risk-register",
                "content": "| Wealth Client Onboarding Risk Register | | | |\n| --- | --- | --- | --- |\n| Risk ID | Risk event | Existing controls | Evidence |\n| R-1 | Incomplete KYC | NAAF completeness review | Signed NAAF and approval workflow record |",
                "anchor_ids": ["risk-register-anchor-1"],
            }
        ],
    )

    assert suggestions
    suggested_text = suggestions[0]["suggested_text"]
    assert "supporting record:" not in suggested_text
    assert "| --- |" not in suggested_text
    assert "Wealth Client Onboarding Risk Register | |" not in suggested_text
    assert not suggested_text.lower().startswith("revise the step")


def test_quality_gate_removes_placeholder_and_raw_table_suggestions():
    suggestions = quality_gate_suggestions(
        [
            {
                "suggestion_id": "bad-table",
                "type": "evidence_gap",
                "status": "open",
                "created_from": "analysis",
                "anchor_id": "a1",
                "summary": "Name evidence.",
                "suggested_text": "This should align to the supporting record: | Risk | Control | Evidence | --- | --- | --- |",
            },
            {
                "suggestion_id": "bad-placeholder",
                "type": "evidence_gap",
                "status": "open",
                "created_from": "analysis",
                "anchor_id": "a2",
                "summary": "Name evidence.",
                "suggested_text": "The performer shall retain the relevant approval, report, ticket, or system record.",
            },
            {
                "suggestion_id": "good",
                "type": "evidence_gap",
                "status": "open",
                "created_from": "analysis",
                "anchor_id": "a3",
                "summary": "Retain the named approval.",
                "suggested_text": "The Relationship Manager shall retain the signed NAAF approval workflow record in the onboarding case file.",
            },
        ]
    )

    assert [item["suggestion_id"] for item in suggestions] == ["good"]


def test_quality_gate_requires_insert_ready_sop_language():
    suggestions = quality_gate_suggestions(
        [
            {
                "suggestion_id": "advice",
                "type": "evidence_gap",
                "status": "open",
                "created_from": "llm",
                "anchor_id": "a1",
                "summary": "Add evidence language.",
                "suggested_text": "Clarify this item so it is easier to test.",
            },
            {
                "suggestion_id": "sop-language",
                "type": "evidence_gap",
                "status": "open",
                "created_from": "llm",
                "anchor_id": "a2",
                "summary": "Add retained evidence language.",
                "suggested_text": "The Relationship Manager shall retain the signed NAAF approval workflow record in the onboarding case file.",
                "style_match_notes": "Uses the SOP's shall-based procedural tone.",
                "source_references": [{"anchor_id": "rcm-1"}],
            },
        ]
    )

    assert [item["suggestion_id"] for item in suggestions] == ["sop-language"]


def test_quality_gate_drops_non_sop_anchor_and_missing_source_references():
    suggestions = quality_gate_suggestions(
        [
            {
                "suggestion_id": "rcm-anchor",
                "type": "evidence_gap",
                "status": "open",
                "created_from": "llm",
                "anchor_id": "rcm-1",
                "summary": "Add retained evidence language.",
                "suggested_text": "The Relationship Manager shall retain the signed NAAF approval workflow record in the onboarding case file.",
                "source_references": [{"anchor_id": "rcm-1"}],
            },
            {
                "suggestion_id": "missing-source",
                "type": "evidence_gap",
                "status": "open",
                "created_from": "llm",
                "anchor_id": "sop-1",
                "summary": "Add retained evidence language.",
                "suggested_text": "The Relationship Manager shall retain the signed NAAF approval workflow record in the onboarding case file.",
                "source_references": [],
            },
            {
                "suggestion_id": "good",
                "type": "evidence_gap",
                "status": "open",
                "created_from": "llm",
                "anchor_id": "sop-1",
                "summary": "Add retained evidence language.",
                "suggested_text": "The Relationship Manager shall retain the signed NAAF approval workflow record in the onboarding case file.",
                "source_references": [{"anchor_id": "rcm-2"}],
            },
        ],
        eligible_anchor_ids={"sop-1"},
        concrete_terms=["signed NAAF approval workflow record"],
        anchor_text_by_id={"sop-1": "Relationship Manager reviews source of wealth before onboarding."},
    )

    assert [item["suggestion_id"] for item in suggestions] == ["good"]


def test_quality_gate_drops_advice_language_and_ungrounded_text():
    suggestions = quality_gate_suggestions(
        [
            {
                "suggestion_id": "advice-start",
                "type": "evidence_gap",
                "status": "open",
                "created_from": "llm",
                "anchor_id": "sop-1",
                "summary": "Add retained evidence language.",
                "suggested_text": "Add the retained evidence artifact to this SOP step.",
                "source_references": [{"anchor_id": "rcm-1"}],
            },
            {
                "suggestion_id": "ungrounded",
                "type": "frequency_gap",
                "status": "open",
                "created_from": "llm",
                "anchor_id": "sop-1",
                "summary": "Add timing language.",
                "suggested_text": "Compliance shall complete this procedure before the process advances.",
                "source_references": [{"anchor_id": "rcm-1"}],
            },
            {
                "suggestion_id": "good",
                "type": "frequency_gap",
                "status": "open",
                "created_from": "llm",
                "anchor_id": "sop-1",
                "summary": "Add timing language.",
                "suggested_text": "Branch Operations shall validate the NAAF before account opening.",
                "source_references": [{"anchor_id": "rcm-1"}],
            },
        ],
        eligible_anchor_ids={"sop-1"},
        concrete_terms=["Branch Operations", "NAAF", "before account opening"],
        anchor_text_by_id={"sop-1": "Validate the onboarding file."},
    )

    assert [item["suggestion_id"] for item in suggestions] == ["good"]


def test_rule_based_suggestions_include_supporting_document_improvement_areas():
    suggestions = generate_rule_based_suggestions(
        [
            {
                "anchor_id": "rcm-1",
                "file_id": "rcm-file",
                "text": "Risk ID R1 | Control ID C1 | Control description: Review high risk onboarding exceptions",
                "section_path": ["Controls"],
            }
        ],
        include_supporting_documents=True,
    )

    assert suggestions
    assert suggestions[0]["anchor_id"] == "rcm-1"
    assert suggestions[0]["type"] in {"ownership_gap", "frequency_gap", "evidence_gap", "mapping_gap", "weak_control_description"}
    assert "Source text:" not in suggestions[0]["suggested_text"]


def test_document_tagging_covers_primary_document_types():
    examples = [
        ({"filename": "risk_register.xlsx", "bucket": "sops"}, "risk_register"),
        ({"filename": "control_inventory.xlsx", "bucket": "supporting_material"}, "control_inventory"),
        ({"filename": "test_evidence.zip", "bucket": "supporting_material"}, "evidence"),
        ({"filename": "process_swimlane.drawio", "bucket": "supporting_material"}, "process_diagram"),
        ({"filename": "audit_findings.pdf", "bucket": "supporting_material"}, "audit_report"),
        ({"filename": "policy_manual.docx", "bucket": "procedures"}, "policy"),
        ({"filename": "operating_sop.docx", "bucket": "sops"}, "sop"),
    ]

    for metadata, expected in examples:
        assert suggest_document_tag(metadata)["confirmed_tag"] == expected


def test_markdown_conversion_preserves_docx_paragraphs_and_tables():
    document = Document()
    document.add_heading("Client onboarding", level=1)
    document.add_paragraph("The operations team validates identity documents.")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Control ID"
    table.rows[0].cells[1].text = "Frequency"
    table.rows[1].cells[0].text = "C-1"
    table.rows[1].cells[1].text = "Monthly"
    payload = BytesIO()
    document.save(payload)

    conversion = convert_bytes_to_markdown(payload.getvalue(), "client_onboarding.docx")

    assert conversion.converter == "python-docx"
    assert "# Client onboarding" in conversion.markdown
    assert "| Control ID | Frequency |" in conversion.markdown
    assert "| C-1 | Monthly |" in conversion.markdown


def test_markdown_conversion_preserves_xlsx_rows_as_markdown_tables():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "RCM"
    sheet.append(["Risk ID", "Control ID", "Frequency"])
    sheet.append(["R-1", "C-1", "Quarterly"])
    payload = BytesIO()
    workbook.save(payload)

    conversion = convert_bytes_to_markdown(payload.getvalue(), "risk_controls_matrix.xlsx")

    assert conversion.converter == "openpyxl"
    assert "## RCM" in conversion.markdown
    assert "| Risk ID | Control ID | Frequency |" in conversion.markdown
    assert "| R-1 | C-1 | Quarterly |" in conversion.markdown


def test_corrupt_office_binary_text_is_detected():
    assert looks_corrupt_markdown("PK\x03\x04 random bytes word/styles.xml ��" * 3)


def test_case_local_index_searches_only_uploaded_case_material():
    case = {
        "case_id": "case-1",
        "process_name": "Vendor onboarding",
        "uploaded_files": [
            {"file_id": "file-1", "filename": "vendor_sop.md", "bucket": "sops"},
            {"file_id": "file-2", "filename": "evidence_pack.pdf", "bucket": "evidence"},
        ],
        "document_tags": [
            {"file_id": "file-1", "confirmed_tag": "sop", "confidence": "high"},
            {"file_id": "file-2", "confirmed_tag": "evidence", "confidence": "medium"},
        ],
        "chunks": [
            {
                "chunk_id": "chunk-1",
                "document_id": "doc_file-1",
                "content": "Business Owner submits vendor onboarding request and validates completeness.",
                "anchor_ids": ["anchor-1"],
                "section_paths": [["Initial onboarding"]],
            },
            {
                "chunk_id": "chunk-2",
                "document_id": "doc_file-2",
                "content": "Evidence pack contains due diligence review records and approval logs.",
                "anchor_ids": ["anchor-2"],
                "section_paths": [["Evidence"]],
            },
        ],
        "case_chat": [
            {"message_id": "msg-1", "role": "user", "content": "Operations Risk approves high-risk exceptions weekly."}
        ],
    }

    index = build_case_local_index(case)
    results = search_case_local_index("weekly exception approval", index, limit=3)

    assert index["case_id"] == "case-1"
    assert index["source_scope"] == "case_uploads_only"
    assert index["stats"]["chunks"] == 2
    assert index["stats"]["chat_messages"] == 1
    assert results[0]["source_type"] == "chat_message"
    assert results[0]["source_id"] == "msg-1"
    assert all(result["case_id"] == "case-1" for result in results)
