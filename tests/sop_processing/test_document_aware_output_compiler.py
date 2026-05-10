from __future__ import annotations

import re
import zipfile
from io import BytesIO
from typing import Any
from xml.etree import ElementTree as ET

import pytest
from docx import Document

from utils.services.schemas import LLMResult
from utils.sop_processing.case_store import DocumentUpliftCaseStore


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W_NS}


def _qn(tag: str) -> str:
    return f"{{{W_NS}}}{tag}"


def _visible_text(element: ET.Element) -> str:
    parts: list[str] = []
    for node in element.iter():
        if node.tag in {_qn("t"), _qn("delText")}:
            parts.append(node.text or "")
        elif node.tag == _qn("tab"):
            parts.append("\t")
        elif node.tag == _qn("br"):
            parts.append("\n")
    return "".join(parts)


def _insertion_text(element: ET.Element) -> str:
    parts: list[str] = []
    for insertion in element.findall(".//w:ins", NS):
        parts.append(
            "".join(node.text or "" for node in insertion.findall(".//w:t", NS)).strip()
        )
    return " ".join(part for part in parts if part)


def _docx_package_xml(docx_bytes: bytes) -> tuple[ET.Element, str]:
    with zipfile.ZipFile(BytesIO(docx_bytes)) as package:
        document_root = ET.fromstring(package.read("word/document.xml"))
        comments_xml = package.read("word/comments.xml").decode("utf-8")
    return document_root, comments_xml


def _body_paragraph_insertions(root: ET.Element) -> list[str]:
    parent_by_child = {child: parent for parent in root.iter() for child in parent}
    insertions: list[str] = []
    for paragraph in root.findall(".//w:p", NS):
        current = paragraph
        is_table_paragraph = False
        while current in parent_by_child:
            current = parent_by_child[current]
            if current.tag == _qn("tbl"):
                is_table_paragraph = True
                break
        if is_table_paragraph:
            continue
        text = _insertion_text(paragraph)
        if text:
            insertions.append(re.sub(r"\s+", " ", text).strip())
    return insertions


def _role_table_docx_bytes() -> bytes:
    document = Document()
    document.add_heading("2. Roles and Responsibilities", level=1)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Role"
    table.cell(0, 1).text = "Responsibility"
    rows = [
        (
            "Business Process Operations Owner",
            "Coordinates process governance and periodic procedure refresh.",
        ),
        (
            "Technology Operations Team",
            "Maintains the workflow platform and supporting operational records.",
        ),
        (
            "Independent Review Function",
            "Tests process design and operating effectiveness.",
        ),
    ]
    for role, responsibility in rows:
        cells = table.add_row().cells
        cells[0].text = role
        cells[1].text = responsibility
    document.add_heading("3. Operating Procedure", level=1)
    document.add_paragraph(
        "The Business Process Operations Owner launches the quarterly workflow review."
    )
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _sectioned_docx_bytes() -> bytes:
    document = Document()
    document.add_heading("5. Monitoring and Review", level=1)
    document.add_paragraph("The operations lead reviews exceptions monthly.")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _bullet_responsibility_docx_bytes() -> bytes:
    document = Document()
    document.add_heading("2. Roles and Responsibilities", level=1)
    document.add_paragraph(
        "Business Process Operations Owner: Coordinates process governance and periodic procedure refresh.",
        style="List Bullet",
    )
    document.add_paragraph(
        "Technology Operations Team: Maintains the workflow platform and supporting operational records.",
        style="List Bullet",
    )
    document.add_paragraph(
        "Independent Review Function: Tests process design and operating effectiveness.",
        style="List Bullet",
    )
    document.add_heading("3. Operating Procedure", level=1)
    document.add_paragraph(
        "The Business Process Operations Owner launches the quarterly workflow review."
    )
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _reference_catalog_docx_bytes() -> bytes:
    document = Document()
    document.add_heading("5. Operating Procedure", level=1)
    document.add_paragraph(
        "The Operations Team reviews access requests before account activation."
    )
    document.add_heading("9. Reference Documents", level=1)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Reference"
    table.cell(0, 1).text = "Description"
    rows = [
        ("POL-OPS-001", "Operations policy"),
        ("PROC-SCREEN-001", "Approved screening platform procedure"),
        ("STD-LOG-001", "Logging and evidence standard"),
    ]
    for reference, description in rows:
        cells = table.add_row().cells
        cells[0].text = reference
        cells[1].text = description
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _memory_store(case_id: str = "case-1") -> DocumentUpliftCaseStore:
    store = DocumentUpliftCaseStore.__new__(DocumentUpliftCaseStore)
    store._use_memory_fallback()
    created = store.create_case({"title": "Document-aware output", "process_name": "Generic"})
    original_id = created["case_id"]
    store._memory[case_id] = store._memory.pop(original_id)
    store._memory[case_id]["case_id"] = case_id
    return store


def _fake_llm_unavailable(prompt: str, schema_name: str, **_kwargs: Any) -> LLMResult:
    if schema_name == "sop_section_rewrite":
        return LLMResult(status="failed", error="offline")
    if schema_name == "swimlane_extraction":
        return LLMResult(status="success", output={"title": "Invalid", "lanes": [], "steps": []})
    raise AssertionError(f"Unexpected schema_name {schema_name}")


def _generate_docx(
    monkeypatch: pytest.MonkeyPatch,
    original_sop_bytes: bytes,
    suggestions: list[dict[str, Any]],
) -> bytes:
    from utils.sop_processing import output_generator

    store = _memory_store()
    monkeypatch.setattr(output_generator, "get_store", lambda: store)
    monkeypatch.setattr(output_generator, "call_llm", _fake_llm_unavailable)

    result = output_generator.generate_outputs(
        case_id="case-1",
        original_sop_bytes=original_sop_bytes,
        suggestions=suggestions,
        process_steps=[],
        corpus_map={"general_relationships": [{"relationship_type": "support_to_procedure"}]},
        style_profile={"body_font": "Arial"},
        pipeline_id="case-1:stage2",
        stage1_cost=None,
    )

    assert result.status == "success"
    output_bytes = store.get_output_content("case-1", result.docx_file_id or "")
    assert output_bytes
    return output_bytes


def test_document_aware_output_updates_responsibility_cell_not_role_cell(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    docx_bytes = _generate_docx(
        monkeypatch,
        _role_table_docx_bytes(),
        [
            {
                "suggestion_id": "bundle-1",
                "suggestion_type": "mapping_gap",
                "severity": "high",
                "title": "Add service access review coverage",
                "detail": "The supporting evidence requires a documented access review activity.",
                "proposed_text": "Coordinate the document updates.",
                "review_status": "accepted",
                "source_references": [
                    {
                        "document_id": "supporting-record",
                        "filename": "Access_Review_Evidence.xlsx",
                        "sheet_name": "Reviews",
                        "row_index": 4,
                    }
                ],
                "edit_targets": [
                    {
                        "target_id": "bundle-1:procedure",
                        "target_type": "procedure_step",
                        "title": "Procedure update",
                        "target_text": (
                            "The Business Process Operations Owner launches the quarterly workflow review."
                        ),
                        "proposed_text": (
                            "The Technology Operations Team reviews service access exceptions "
                            "before workflow closure. Retained evidence includes review log, "
                            "exception notes, and closure approval."
                        ),
                    },
                    {
                        "target_id": "bundle-1:responsibility",
                        "target_type": "role_responsibility",
                        "title": "Responsibility update",
                        "target_text": (
                            "| Role | Responsibility |\n"
                            "| Business Process Operations Owner | Coordinates process governance and periodic procedure refresh. |\n"
                            "| Technology Operations Team | Maintains the workflow platform and supporting operational records. |\n"
                            "| Independent Review Function | Tests process design and operating effectiveness. |"
                        ),
                        "proposed_text": (
                            "The Technology Operations Team is responsible for ensuring that "
                            "service access exceptions are reviewed before workflow closure. "
                            "Retained evidence includes review log, exception notes, and closure approval."
                        ),
                    },
                ],
            }
        ],
    )

    root, _comments_xml = _docx_package_xml(docx_bytes)
    role_table = root.findall(".//w:tbl", NS)[0]
    rows = role_table.findall("./w:tr", NS)
    business_role_cell = rows[1].findall("./w:tc", NS)[0]
    technology_role_cell = rows[2].findall("./w:tc", NS)[0]
    technology_responsibility_cell = rows[2].findall("./w:tc", NS)[1]

    assert _insertion_text(business_role_cell) == ""
    assert _insertion_text(technology_role_cell) == ""

    responsibility_insert = _insertion_text(technology_responsibility_cell)
    assert "service access" in responsibility_insert
    assert responsibility_insert.startswith("Reviews service access exceptions")
    assert len(responsibility_insert) <= 220
    assert "Retained evidence includes" not in responsibility_insert
    assert "responsible for ensuring that" not in responsibility_insert

    body_insertions = _body_paragraph_insertions(root)
    assert any("Retained evidence includes" in text for text in body_insertions)


def test_document_aware_output_reuses_existing_section_and_splits_heading_from_body(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    docx_bytes = _generate_docx(
        monkeypatch,
        _sectioned_docx_bytes(),
        [
            {
                "suggestion_id": "section-1",
                "suggestion_type": "process_improvement",
                "severity": "medium",
                "title": "Clarify monitoring evidence",
                "detail": "The supporting evidence requires the review date and reviewer to be retained.",
                "proposed_text": (
                    "5. Monitoring and Review The operations lead records the review date, "
                    "reviewer, exception count, and remediation status after each monthly review."
                ),
                "review_status": "accepted",
                "source_references": [
                    {"document_id": "monitoring-log", "filename": "Monitoring_Log.csv"}
                ],
            }
        ],
    )

    root, _comments_xml = _docx_package_xml(docx_bytes)
    insertions = _body_paragraph_insertions(root)

    assert any("records the review date" in text for text in insertions)
    assert all(not text.startswith("5. Monitoring and Review") for text in insertions)
    assert all("Review The operations lead" not in text for text in insertions)


def test_document_aware_procedure_addition_does_not_anchor_in_reference_catalog(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    docx_bytes = _generate_docx(
        monkeypatch,
        _reference_catalog_docx_bytes(),
        [
            {
                "suggestion_id": "reference-anchor-1",
                "suggestion_type": "process_improvement",
                "severity": "medium",
                "title": "Specify approved screening platform",
                "detail": (
                    "The evidence requires the process to identify the approved "
                    "screening platform and match disposition."
                ),
                "proposed_text": (
                    "The Operations Team performs screening using the approved "
                    "platform (PROC-SCREEN-001). Potential matches are placed on "
                    "hold until a documented final disposition is recorded."
                ),
                "review_status": "accepted",
                "source_references": [
                    {"document_id": "evidence", "filename": "Screening_Evidence.docx"}
                ],
            }
        ],
    )

    root, _comments_xml = _docx_package_xml(docx_bytes)
    tables = root.findall(".//w:tbl", NS)
    assert tables
    reference_table_insertions = _insertion_text(tables[0])
    body_insertions = _body_paragraph_insertions(root)

    assert "PROC-SCREEN-001" not in reference_table_insertions
    assert any("PROC-SCREEN-001" in text for text in body_insertions)


def test_document_aware_output_strips_numbered_heading_with_colon() -> None:
    from utils.sop_processing import output_generator

    cleaned = output_generator._clean_rewritten_text(
        "5.2 Access Review Process: Technology Operations must verify access "
        "exceptions before workflow closure."
    )

    assert cleaned == (
        "Technology Operations must verify access exceptions before workflow closure."
    )


def test_document_aware_comments_use_specific_change_titles_not_generic_target_labels(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    docx_bytes = _generate_docx(
        monkeypatch,
        _role_table_docx_bytes(),
        [
            {
                "suggestion_id": "comment-1",
                "suggestion_type": "mapping_gap",
                "severity": "high",
                "title": "Add service access review coverage",
                "detail": "The supporting evidence requires a documented access review activity.",
                "proposed_text": "Coordinate the document updates.",
                "review_status": "accepted",
                "source_references": [
                    {
                        "document_id": "supporting-record",
                        "filename": "Access_Review_Evidence.xlsx",
                        "sheet_name": "Reviews",
                        "row_index": 4,
                    }
                ],
                "edit_targets": [
                    {
                        "target_id": "comment-1:procedure",
                        "target_type": "procedure_step",
                        "title": "Procedure update",
                        "target_text": (
                            "The Business Process Operations Owner launches the quarterly workflow review."
                        ),
                        "proposed_text": (
                            "The Technology Operations Team reviews service access exceptions "
                            "before workflow closure."
                        ),
                    },
                    {
                        "target_id": "comment-1:responsibility",
                        "target_type": "role_responsibility",
                        "title": "Responsibility update",
                        "target_text": (
                            "| Role | Responsibility |\n"
                            "| Technology Operations Team | Maintains the workflow platform and supporting operational records. |"
                        ),
                        "proposed_text": (
                            "Owns service access exception review before workflow closure."
                        ),
                    },
                ],
            }
        ],
    )

    _root, comments_xml = _docx_package_xml(docx_bytes)

    assert "Add service access review coverage" in comments_xml
    assert "Access_Review_Evidence.xlsx - Reviews row 4" in comments_xml
    assert "Procedure update. Why this was suggested" not in comments_xml
    assert "Responsibility update. Why this was suggested" not in comments_xml
    assert "anchor" not in comments_xml


def test_document_aware_output_updates_bullet_responsibility_area(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    docx_bytes = _generate_docx(
        monkeypatch,
        _bullet_responsibility_docx_bytes(),
        [
            {
                "suggestion_id": "bullet-1",
                "suggestion_type": "mapping_gap",
                "severity": "medium",
                "title": "Add workflow exception ownership",
                "detail": "The supporting evidence requires exception review ownership.",
                "proposed_text": "Coordinate the document updates.",
                "review_status": "accepted",
                "source_references": [
                    {"document_id": "supporting-record", "filename": "Workflow_Review.csv"}
                ],
                "edit_targets": [
                    {
                        "target_id": "bullet-1:responsibility",
                        "target_type": "role_responsibility",
                        "title": "Responsibility update",
                        "target_text": (
                            "Business Process Operations Owner: Coordinates process governance and periodic procedure refresh.\n"
                            "Technology Operations Team: Maintains the workflow platform and supporting operational records.\n"
                            "Independent Review Function: Tests process design and operating effectiveness."
                        ),
                        "proposed_text": (
                            "The Technology Operations Team is responsible for ensuring that "
                            "workflow exceptions are reviewed before closure. Retained evidence "
                            "includes exception review log."
                        ),
                    }
                ],
            }
        ],
    )

    root, _comments_xml = _docx_package_xml(docx_bytes)
    paragraphs = root.findall(".//w:p", NS)
    paragraph_texts = [re.sub(r"\s+", " ", _visible_text(p)).strip() for p in paragraphs]
    technology_index = next(
        index
        for index, text in enumerate(paragraph_texts)
        if text.startswith("Technology Operations Team:")
    )
    procedure_index = next(
        index
        for index, text in enumerate(paragraph_texts)
        if text.startswith("The Business Process Operations Owner launches")
    )
    target_paragraph = paragraphs[technology_index]
    insertion = _insertion_text(target_paragraph)

    assert insertion
    assert insertion.startswith("Reviews workflow exceptions")
    assert "workflow exceptions" in insertion
    assert len(insertion) <= 220
    assert "Retained evidence includes" not in insertion
    assert technology_index < procedure_index


def test_document_aware_responsibility_text_converts_passive_records_to_active_voice() -> None:
    from utils.sop_processing import output_generator

    concise = output_generator._concise_responsibility_text(
        "The Technology Operations Team is responsible for ensuring that all "
        "system access records are stored in the approved repository within "
        "5 business days of creation. Records are retained for seven years. "
        "Retained evidence includes repository confirmation."
    )

    assert concise == (
        "Maintains system access records in the approved repository within "
        "5 business days of creation."
    )
    assert "Owns all" not in concise
    assert "are stored" not in concise


def test_document_aware_responsibility_text_truncates_at_clause_boundary() -> None:
    from utils.sop_processing import output_generator

    concise = output_generator._concise_responsibility_text(
        "The Operations Review Team is responsible for ensuring that all service "
        "accounts are re-screened against access lists, exception logs, privileged "
        "access reports, and application inventories on a monthly basis and upon "
        "trigger events such as incidents, urgent advisories, emergency access, "
        "vendor changes, platform changes, or control failures. Retained evidence "
        "includes review log."
    )

    assert len(concise) <= 220
    assert not concise.endswith(" public.")
    assert concise.endswith(".")
