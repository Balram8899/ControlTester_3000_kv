from __future__ import annotations

import zipfile
import warnings
from io import BytesIO
from typing import Any
from xml.etree import ElementTree as ET

import pytest
from docx import Document
from docx.shared import Inches
from pyparsing import PyparsingDeprecationWarning

warnings.filterwarnings("ignore", category=PyparsingDeprecationWarning)

from utils.services.schemas import LLMResult
from utils.sop_processing.case_store import DocumentUpliftCaseStore


def _docx_bytes() -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    document.add_paragraph("KPMG")
    document.add_paragraph("SOP-WM-KYC-ONBOARD-003 | v3.1 | 01 Feb 2025")
    document.add_heading("3. CLIENT ONBOARDING", level=1)
    document.add_heading("3.1 Know Your Client (KYC) and Account Opening", level=2)
    document.add_paragraph(
        "Verify proof of address not older than 6 months before account opening."
    )
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _process_docx_bytes() -> bytes:
    document = Document()
    document.add_heading("4. Procedure", level=1)
    document.add_paragraph(
        "The Relationship Manager collects and verifies client identity documentation "
        "using approved onboarding methods."
    )
    document.add_paragraph(
        "The Compliance Officer screens the client against sanctions and watchlist sources "
        "before account opening."
    )
    document.add_paragraph("The Service Desk logs incident tickets before triage.")
    document.add_paragraph("The Incident Manager reviews major incidents daily.")
    document.add_paragraph("The Problem Manager tracks recurring incident themes.")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _docx_with_short_header_and_matching_body() -> bytes:
    document = Document()
    document.add_heading("Document history", level=1)
    document.add_paragraph("Date")
    document.add_heading("KYC refresh", level=1)
    document.add_paragraph("The Investment Advisor reviews KYC information annually.")
    document.add_paragraph("Author")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _docx_with_responsibility_and_procedure_sections() -> bytes:
    document = Document()
    document.add_heading("2. Roles and Responsibilities", level=1)
    document.add_paragraph(
        "Technology Operations monitors incidents and escalates critical events."
    )
    document.add_heading("3. Procedure", level=1)
    document.add_paragraph("The Service Desk logs incident tickets before triage.")
    document.add_paragraph("The Incident Manager reviews major incidents daily.")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _memory_store(case_id: str = "case-1") -> DocumentUpliftCaseStore:
    store = DocumentUpliftCaseStore.__new__(DocumentUpliftCaseStore)
    store._use_memory_fallback()
    created = store.create_case({"title": "KYC uplift", "process_name": "KYC"})
    original_id = created["case_id"]
    store._memory[case_id] = store._memory.pop(original_id)
    store._memory[case_id]["case_id"] = case_id
    return store


def _suggestion(status: str = "accepted") -> dict[str, Any]:
    return {
        "suggestion_id": "sug-1",
        "suggestion_type": "evidence_requirement",
        "severity": "high",
        "title": "Clarify recency requirement for address proof",
        "detail": "The RCM requires current and verifiable address evidence.",
        "original_text": "not older than 6 months",
        "proposed_text": "not older than 3 months",
        "review_status": status,
        "source_references": [
            {"document_id": "SOP.docx", "anchor_id": "anc-1"},
            {"document_id": "RCM.xlsx", "sheet_name": "Controls", "row_index": 14},
        ],
    }


def _fake_llm(prompt: str, schema_name: str, **_kwargs: Any) -> LLMResult:
    if schema_name == "sop_section_rewrite":
        assert "not older than 6 months" in prompt
        assert "not older than 3 months" in prompt
        return LLMResult(
            status="success",
            output={"rewritten_text": "not older than 3 months"},
        )
    if schema_name == "swimlane_extraction":
        assert "not older than 3 months" in prompt
        return LLMResult(
            status="success",
            output={
                "title": "Access review process swimlane",
                "lanes": [
                    {
                        "lane_id": "l1",
                        "label": "Business owner",
                        "colour_hex": "#1E49E2",
                    }
                ],
                "steps": [
                    {
                        "step_id": "s1",
                        "lane_id": "l1",
                        "label": "Start",
                        "step_type": "start",
                        "next_steps": [],
                        "branch_labels": {},
                    }
                ],
            },
        )
    raise AssertionError(f"Unexpected schema_name {schema_name}")


def _fake_llm_without_swimlane(prompt: str, schema_name: str, **_kwargs: Any) -> LLMResult:
    if schema_name == "sop_section_rewrite":
        return LLMResult(
            status="success",
            output={"rewritten_text": "not older than 3 months"},
        )
    if schema_name == "swimlane_extraction":
        return LLMResult(status="success", output={"title": "Invalid", "lanes": [], "steps": []})
    raise AssertionError(f"Unexpected schema_name {schema_name}")


def _fake_llm_swimlane_from_as_is_sop(prompt: str, schema_name: str, **_kwargs: Any) -> LLMResult:
    if schema_name == "sop_section_rewrite":
        raise AssertionError("Rejected suggestions should not be rewritten")
    if schema_name == "swimlane_extraction":
        assert "not older than 6 months" in prompt
        assert "not older than 3 months" not in prompt
        return LLMResult(
            status="success",
            output={
                "title": "As-is SOP swimlane",
                "lanes": [
                    {
                        "lane_id": "owner",
                        "label": "Process owner",
                        "colour_hex": "#1E49E2",
                    }
                ],
                "steps": [
                    {
                        "step_id": "start",
                        "lane_id": "owner",
                        "label": "Start",
                        "step_type": "start",
                        "next_steps": [],
                        "branch_labels": {},
                    }
                ],
            },
        )
    raise AssertionError(f"Unexpected schema_name {schema_name}")


def _fake_llm_rewrite_unavailable(prompt: str, schema_name: str, **_kwargs: Any) -> LLMResult:
    if schema_name == "sop_section_rewrite":
        return LLMResult(status="failed", error="offline")
    if schema_name == "swimlane_extraction":
        return LLMResult(status="success", output={"title": "Invalid", "lanes": [], "steps": []})
    raise AssertionError(f"Unexpected schema_name {schema_name}")


def _fake_llm_markdown_rewrite(prompt: str, schema_name: str, **_kwargs: Any) -> LLMResult:
    if schema_name == "sop_section_rewrite":
        return LLMResult(
            status="success",
            output={
                "rewritten_text": (
                    "### Procedure Step: Evidence review "
                    "Owner: Operations "
                    "Procedure: Review exception tickets before closure. "
                    "Retained Evidence: * Exception ticket * Review log"
                )
            },
        )
    if schema_name == "swimlane_extraction":
        return LLMResult(status="success", output={"title": "Invalid", "lanes": [], "steps": []})
    raise AssertionError(f"Unexpected schema_name {schema_name}")


def test_generate_outputs_retains_docx_shell_and_writes_track_changes_with_rationale(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from utils.sop_processing import output_generator

    store = _memory_store()
    monkeypatch.setattr(output_generator, "get_store", lambda: store)
    monkeypatch.setattr(output_generator, "call_llm", _fake_llm)

    result = output_generator.generate_outputs(
        case_id="case-1",
        original_sop_bytes=_docx_bytes(),
        suggestions=[_suggestion()],
        process_steps=[],
        corpus_map={"risk_to_control_map": [{"risk_id": "R-1"}]},
        style_profile={"primary_colour_hex": "#00338D", "body_font": "Arial"},
        pipeline_id="case-1:stage2",
        stage1_cost=None,
    )

    assert result.status == "success"
    assert result.docx_file_id
    assert result.diagram_png_file_id
    assert result.diagram_pdf_file_id
    assert result.sections_rewritten == 1
    assert result.swimlane_spec is not None
    output_bytes = store.get_output_content("case-1", result.docx_file_id)
    assert output_bytes
    with zipfile.ZipFile(BytesIO(output_bytes)) as package:
        document_xml = package.read("word/document.xml").decode("utf-8")
        comments_xml = package.read("word/comments.xml").decode("utf-8")
        content_types = package.read("[Content_Types].xml").decode("utf-8")

    assert "SOP-WM-KYC-ONBOARD-003" in document_xml
    assert "CLIENT ONBOARDING" in document_xml
    assert "<w:del " in document_xml
    assert "<w:ins " in document_xml
    assert "<w:delText>not older than 6 months</w:delText>" in document_xml
    assert "<w:t>not older than 3 months</w:t>" in document_xml
    assert "comments.xml" in content_types
    assert "Why this was suggested" in comments_xml
    assert "The RCM requires current and verifiable address evidence." in comments_xml
    assert "RCM.xlsx" in comments_xml
    assert store.get_output_content("case-1", result.diagram_png_file_id)[:8] == b"\x89PNG\r\n\x1a\n"
    assert store.get_output_content("case-1", result.diagram_pdf_file_id)[:4] == b"%PDF"


def test_generate_outputs_places_addition_near_anchor_and_hides_internal_anchor_ids(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from utils.sop_processing import output_generator

    store = _memory_store()
    monkeypatch.setattr(output_generator, "get_store", lambda: store)
    monkeypatch.setattr(output_generator, "call_llm", _fake_llm_rewrite_unavailable)

    suggestion = {
        "suggestion_id": "add-1",
        "suggestion_type": "process_improvement",
        "severity": "high",
        "title": "Add required control procedure",
        "detail": "The supporting control inventory contains this required control, but the primary SOP does not describe the related procedure.",
        "original_text": None,
        "proposed_text": (
            "Add a procedure step for control CTRL-22 owned by Operations. "
            "The step should describe reconciling entitlement changes before completion "
            "and identify retained evidence: reconciliation log."
        ),
        "review_status": "accepted",
        "source_references": [
            {
                "document_id": "primary-procedure",
                "filename": "Operations_SOP.docx",
                "anchor_id": "proc-anchor-1",
            },
            {
                "document_id": "supporting-control-file",
                "filename": "Control_Inventory.xlsx",
                "sheet_name": "Controls",
                "row_index": 12,
            },
        ],
    }

    result = output_generator.generate_outputs(
        case_id="case-1",
        original_sop_bytes=_process_docx_bytes(),
        suggestions=[suggestion],
        process_steps=[
            {
                "step_id": "PS-1",
                "actor": "Relationship Manager",
                "action": "The Relationship Manager collects and verifies client identity documentation using approved onboarding methods.",
                "anchor_id": "proc-anchor-1",
            }
        ],
        corpus_map={"general_relationships": [{"relationship_type": "control_to_procedure"}]},
        style_profile={"body_font": "Arial"},
        pipeline_id="case-1:stage2",
        stage1_cost=None,
    )

    assert result.status == "success"
    output_bytes = store.get_output_content("case-1", result.docx_file_id or "")
    assert output_bytes
    with zipfile.ZipFile(BytesIO(output_bytes)) as package:
        document_xml = package.read("word/document.xml").decode("utf-8")
        comments_xml = package.read("word/comments.xml").decode("utf-8")

    first_step_index = document_xml.index("The Relationship Manager collects")
    insertion_index = document_xml.index("The Operations team reconciles entitlement changes")
    next_step_index = document_xml.index("The Compliance Officer screens")
    assert first_step_index < insertion_index < next_step_index
    assert "Suggested addition" not in document_xml
    assert "For control CTRL-22" not in document_xml
    assert "**Owner" not in document_xml
    assert "**Procedure" not in document_xml
    assert "proc-anchor-1" not in comments_xml
    assert "anchor" not in comments_xml.lower()
    assert "Operations_SOP.docx" in comments_xml
    assert "Control_Inventory.xlsx - Controls row 12" in comments_xml


def test_generate_outputs_places_non_control_source_addition_near_target_anchor(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from utils.sop_processing import output_generator

    store = _memory_store()
    monkeypatch.setattr(output_generator, "get_store", lambda: store)
    monkeypatch.setattr(output_generator, "call_llm", _fake_llm_rewrite_unavailable)

    result = output_generator.generate_outputs(
        case_id="case-1",
        original_sop_bytes=_process_docx_bytes(),
        suggestions=[
            {
                "suggestion_id": "incident-1",
                "suggestion_type": "process_improvement",
                "severity": "medium",
                "title": "Clarify incident escalation timing",
                "detail": "The incident log shows escalation timing is needed for repeatable handling.",
                "target_anchor_id": "incident-anchor-1",
                "proposed_text": (
                    "Add an escalation step owned by Technology Operations. "
                    "The step should describe notifying the incident manager within 30 minutes "
                    "and identify retained evidence: incident ticket."
                ),
                "review_status": "accepted",
                "source_references": [
                    {
                        "document_id": "incident-evidence",
                        "filename": "Incident_Log.pdf",
                        "page": 4,
                    }
                ],
            }
        ],
        process_steps=[
            {
                "step_id": "INC-1",
                "action": "The Service Desk logs incident tickets before triage.",
                "anchor_id": "incident-anchor-1",
            }
        ],
        corpus_map={"general_relationships": [{"relationship_type": "incident_to_procedure"}]},
        style_profile={"body_font": "Arial"},
        pipeline_id="case-1:stage2",
        stage1_cost=None,
    )

    assert result.status == "success"
    output_bytes = store.get_output_content("case-1", result.docx_file_id or "")
    assert output_bytes
    with zipfile.ZipFile(BytesIO(output_bytes)) as package:
        document_xml = package.read("word/document.xml").decode("utf-8")
        comments_xml = package.read("word/comments.xml").decode("utf-8")

    source_step_index = document_xml.index("The Service Desk logs incident tickets")
    insertion_index = document_xml.index("The Technology Operations team notifies")
    next_step_index = document_xml.index("The Incident Manager reviews major incidents")
    assert source_step_index < insertion_index < next_step_index
    assert "control" not in document_xml[ insertion_index : insertion_index + 250 ].lower()
    assert "Incident_Log.pdf - page 4" in comments_xml


def test_generate_outputs_resolves_source_filenames_from_case_document_tags(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from utils.sop_processing import output_generator

    store = _memory_store()
    monkeypatch.setattr(output_generator, "get_store", lambda: store)
    monkeypatch.setattr(output_generator, "call_llm", _fake_llm_rewrite_unavailable)

    result = output_generator.generate_outputs(
        case_id="case-1",
        original_sop_bytes=_process_docx_bytes(),
        suggestions=[
            {
                "suggestion_id": "source-name-1",
                "suggestion_type": "process_improvement",
                "severity": "medium",
                "title": "Add required procedure detail",
                "detail": "A supporting document identifies a procedure detail that should be reflected in the SOP.",
                "target_anchor_id": "proc-anchor-1",
                "proposed_text": (
                    "Add a procedure step owned by Operations. "
                    "The step should describe reviewing exceptions before closure "
                    "and identify retained evidence: exception log."
                ),
                "review_status": "accepted",
                "source_references": [
                    {"document_id": "procedure-file-id", "anchor_id": "proc-anchor-1"},
                    {
                        "document_id": "support-file-id",
                        "sheet_name": "Exceptions",
                        "row_index": 7,
                    },
                ],
            }
        ],
        process_steps=[
            {
                "step_id": "PS-1",
                "action": "The Relationship Manager collects and verifies client identity documentation using approved onboarding methods.",
                "anchor_id": "proc-anchor-1",
            }
        ],
        corpus_map={"general_relationships": [{"relationship_type": "support_to_procedure"}]},
        style_profile={"body_font": "Arial"},
        pipeline_id="case-1:stage2",
        stage1_cost=None,
        source_documents=[
            {"file_id": "procedure-file-id", "filename": "Primary_Procedure.docx"},
            {"file_id": "support-file-id", "filename": "Exception_Register.xlsx"},
        ],
    )

    assert result.status == "success"
    output_bytes = store.get_output_content("case-1", result.docx_file_id or "")
    assert output_bytes
    with zipfile.ZipFile(BytesIO(output_bytes)) as package:
        comments_xml = package.read("word/comments.xml").decode("utf-8")

    assert "Primary_Procedure.docx" in comments_xml
    assert "Exception_Register.xlsx - Exceptions row 7" in comments_xml
    assert "procedure-file-id" not in comments_xml
    assert "support-file-id" not in comments_xml


def test_generate_outputs_uses_primary_procedure_filename_in_swimlane_metadata(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from utils.sop_processing import output_generator

    store = _memory_store()
    monkeypatch.setattr(output_generator, "get_store", lambda: store)
    monkeypatch.setattr(output_generator, "call_llm", _fake_llm)

    result = output_generator.generate_outputs(
        case_id="case-1",
        original_sop_bytes=_docx_bytes(),
        suggestions=[_suggestion(status="accepted")],
        process_steps=[
            {
                "step_id": "PS-1",
                "actor": "Operations",
                "action": "Operations reviews the request.",
            }
        ],
        corpus_map={},
        style_profile={"body_font": "Arial"},
        pipeline_id="case-1:stage2",
        stage1_cost=None,
        source_documents=[
            {"file_id": "support-file-id", "filename": "Exception_Register.xlsx", "tag": "risk_data"},
            {"file_id": "procedure-file-id", "filename": "Primary_Procedure.docx", "tag": "procedure"},
        ],
    )

    assert result.swimlane_spec is not None
    assert result.swimlane_spec.document_name == "Primary_Procedure.docx"
    assert result.swimlane_spec.process_owner == "Business owner"


def test_generate_outputs_applies_holistic_edit_targets_to_distinct_document_areas(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from utils.sop_processing import output_generator

    store = _memory_store()
    monkeypatch.setattr(output_generator, "get_store", lambda: store)
    monkeypatch.setattr(output_generator, "call_llm", _fake_llm_rewrite_unavailable)

    result = output_generator.generate_outputs(
        case_id="case-1",
        original_sop_bytes=_docx_with_responsibility_and_procedure_sections(),
        suggestions=[
            {
                "suggestion_id": "bundle-1",
                "suggestion_type": "mapping_gap",
                "severity": "high",
                "title": "Add monitoring coverage",
                "detail": "The supporting service report identifies monitoring that is not fully described.",
                "proposed_text": "Coordinate the relevant SOP updates.",
                "review_status": "accepted",
                "source_references": [
                    {"document_id": "service-report", "filename": "Service_Report.pdf"}
                ],
                "edit_targets": [
                    {
                        "target_id": "bundle-1:procedure",
                        "target_type": "procedure_step",
                        "title": "Procedure update",
                        "target_anchor_id": "proc-anchor-1",
                        "target_text": "The Service Desk logs incident tickets before triage.",
                        "proposed_text": (
                            "Add a procedure step owned by Technology Operations. "
                            "The step should describe reviewing service availability alerts daily "
                            "and identify retained evidence: monitoring review log."
                        ),
                    },
                    {
                        "target_id": "bundle-1:responsibility",
                        "target_type": "role_responsibility",
                        "title": "Responsibility update",
                        "target_anchor_id": "role-anchor-1",
                        "target_text": (
                            "Technology Operations monitors incidents and escalates critical events."
                        ),
                        "proposed_text": (
                            "Technology Operations is responsible for reviewing service availability "
                            "alerts daily. Retained evidence includes monitoring review log."
                        ),
                    },
                ],
            }
        ],
        process_steps=[],
        corpus_map={"general_relationships": [{"relationship_type": "source_to_procedure"}]},
        style_profile={"body_font": "Arial"},
        pipeline_id="case-1:stage2",
        stage1_cost=None,
    )

    assert result.status == "success"
    output_bytes = store.get_output_content("case-1", result.docx_file_id or "")
    assert output_bytes
    with zipfile.ZipFile(BytesIO(output_bytes)) as package:
        document_xml = package.read("word/document.xml").decode("utf-8")
        comments_xml = package.read("word/comments.xml").decode("utf-8")

    responsibility_index = document_xml.index("Technology Operations monitors incidents")
    responsibility_insert_index = document_xml.index(
        "Reviews service availability alerts daily."
    )
    procedure_index = document_xml.index("The Service Desk logs incident tickets before triage.")
    procedure_insert_index = document_xml.index(
        "The Technology Operations team reviews service availability alerts daily."
    )
    incident_manager_index = document_xml.index("The Incident Manager reviews major incidents daily.")

    assert responsibility_index < responsibility_insert_index < procedure_index
    assert procedure_index < procedure_insert_index < incident_manager_index
    assert comments_xml.count("<w:comment ") == 2
    assert "Procedure update" in comments_xml
    assert "Responsibility update" in comments_xml
    assert "Retained evidence includes monitoring review log" not in document_xml[
        responsibility_insert_index:procedure_index
    ]


def test_generate_outputs_formats_markdown_table_targets_as_word_tables(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from utils.sop_processing import output_generator

    store = _memory_store()
    monkeypatch.setattr(output_generator, "get_store", lambda: store)
    monkeypatch.setattr(output_generator, "call_llm", _fake_llm_rewrite_unavailable)

    result = output_generator.generate_outputs(
        case_id="case-1",
        original_sop_bytes=_docx_with_responsibility_and_procedure_sections(),
        suggestions=[
            {
                "suggestion_id": "raci-1",
                "suggestion_type": "process_improvement",
                "severity": "high",
                "title": "Add an accountability matrix for multi-role activities",
                "detail": "The procedure has multiple roles but no accountability matrix.",
                "review_status": "accepted",
                "source_references": [{"document_id": "sop", "filename": "SOP.docx"}],
                "edit_targets": [
                    {
                        "target_id": "raci-1:matrix",
                        "target_type": "raci_matrix",
                        "title": "RACI update",
                        "target_text": "Technology Operations monitors incidents and escalates critical events.",
                        "proposed_text": (
                            "Add an accountability matrix for material activities:\n\n"
                            "| Activity / decision | Responsible | Accountable | Consulted | Informed |\n"
                            "| --- | --- | --- | --- | --- |\n"
                            "| Incident triage | SOC analyst | Incident Manager | Legal / Compliance | CISO |\n\n"
                            "Populate the matrix from the uploaded procedure and supporting evidence."
                        ),
                    }
                ],
            }
        ],
        process_steps=[],
        corpus_map={"general_relationships": [{"relationship_type": "source_to_procedure"}]},
        style_profile={"body_font": "Arial"},
        pipeline_id="case-1:stage2",
        stage1_cost=None,
    )

    assert result.status == "success"
    output_bytes = store.get_output_content("case-1", result.docx_file_id or "")
    assert output_bytes
    rendered = Document(BytesIO(output_bytes))
    assert len(rendered.tables) >= 1
    table_text = "\n".join(cell.text for row in rendered.tables[-1].rows for cell in row.cells)
    assert "Activity / decision" in table_text
    assert "Responsible" in table_text
    assert "Incident triage" in table_text

    with zipfile.ZipFile(BytesIO(output_bytes)) as package:
        document_xml = package.read("word/document.xml").decode("utf-8")
    assert "| Activity / decision |" not in document_xml
    assert "| --- | --- |" not in document_xml


def test_generate_outputs_uses_best_process_step_when_target_anchor_is_missing(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from utils.sop_processing import output_generator

    store = _memory_store()
    monkeypatch.setattr(output_generator, "get_store", lambda: store)
    monkeypatch.setattr(output_generator, "call_llm", _fake_llm_rewrite_unavailable)

    result = output_generator.generate_outputs(
        case_id="case-1",
        original_sop_bytes=_process_docx_bytes(),
        suggestions=[
            {
                "suggestion_id": "best-step-1",
                "suggestion_type": "process_improvement",
                "severity": "medium",
                "title": "Clarify major incident review evidence",
                "detail": "The source document requires evidence for major incident review.",
                "proposed_text": (
                    "Add a procedure step owned by Incident Management. "
                    "The step should describe reviewing major incident tickets daily "
                    "and identify retained evidence: review log."
                ),
                "review_status": "accepted",
                "source_references": [
                    {"document_id": "incident-log", "filename": "Major_Incident_Log.csv"}
                ],
            }
        ],
        process_steps=[
            {
                "step_id": "INC-1",
                "action": "The Service Desk logs incident tickets before triage.",
                "anchor_id": "incident-anchor-1",
            },
            {
                "step_id": "INC-2",
                "action": "The Incident Manager reviews major incidents daily.",
                "anchor_id": "incident-anchor-2",
            },
        ],
        corpus_map={"general_relationships": [{"relationship_type": "evidence_to_procedure"}]},
        style_profile={"body_font": "Arial"},
        pipeline_id="case-1:stage2",
        stage1_cost=None,
    )

    assert result.status == "success"
    output_bytes = store.get_output_content("case-1", result.docx_file_id or "")
    assert output_bytes
    with zipfile.ZipFile(BytesIO(output_bytes)) as package:
        document_xml = package.read("word/document.xml").decode("utf-8")

    incident_manager_index = document_xml.index("The Incident Manager reviews major incidents daily.")
    insertion_index = document_xml.index("The Incident Management team reviews major incident tickets daily.")
    problem_manager_index = document_xml.index("The Problem Manager tracks recurring incident themes.")
    assert incident_manager_index < insertion_index < problem_manager_index


def test_generate_outputs_sanitizes_markdown_rewrite_into_sop_prose(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from utils.sop_processing import output_generator

    store = _memory_store()
    monkeypatch.setattr(output_generator, "get_store", lambda: store)
    monkeypatch.setattr(output_generator, "call_llm", _fake_llm_markdown_rewrite)

    result = output_generator.generate_outputs(
        case_id="case-1",
        original_sop_bytes=_process_docx_bytes(),
        suggestions=[
            {
                "suggestion_id": "markdown-1",
                "suggestion_type": "process_improvement",
                "severity": "medium",
                "title": "Clarify evidence review",
                "detail": "The supporting evidence requires clearer review wording.",
                "target_anchor_id": "incident-anchor-1",
                "proposed_text": "Review exception tickets before closure.",
                "review_status": "accepted",
                "source_references": [
                    {"document_id": "evidence-file", "filename": "Exception_Evidence.docx"}
                ],
            }
        ],
        process_steps=[
            {
                "step_id": "INC-1",
                "action": "The Service Desk logs incident tickets before triage.",
                "anchor_id": "incident-anchor-1",
            }
        ],
        corpus_map={"general_relationships": [{"relationship_type": "evidence_to_procedure"}]},
        style_profile={"body_font": "Arial"},
        pipeline_id="case-1:stage2",
        stage1_cost=None,
    )

    assert result.status == "success"
    output_bytes = store.get_output_content("case-1", result.docx_file_id or "")
    assert output_bytes
    with zipfile.ZipFile(BytesIO(output_bytes)) as package:
        document_xml = package.read("word/document.xml").decode("utf-8")

    assert "###" not in document_xml
    assert "Procedure:" not in document_xml
    assert "Owner:" not in document_xml
    assert "Retained Evidence:" not in document_xml
    assert "Evidence review. Review exception tickets before closure." in document_xml
    assert "The accountable owner is Operations." in document_xml
    assert "Retained evidence includes Exception ticket; Review log." in document_xml


def test_generate_outputs_compacts_numbered_subprocedure_additions(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from utils.sop_processing import output_generator

    store = _memory_store()

    def fake_llm(prompt: str, schema_name: str, **_kwargs: Any) -> LLMResult:
        if schema_name == "swimlane_extraction":
            return LLMResult(status="success", output={"title": "Invalid", "lanes": [], "steps": []})
        raise AssertionError(f"Unexpected schema_name {schema_name}")

    monkeypatch.setattr(output_generator, "get_store", lambda: store)
    monkeypatch.setattr(output_generator, "call_llm", fake_llm)
    monkeypatch.setenv("DOCUMENT_UPLIFT_STAGE2_REWRITE_LIMIT", "0")

    result = output_generator.generate_outputs(
        case_id="case-1",
        original_sop_bytes=_process_docx_bytes(),
        suggestions=[
            {
                "suggestion_id": "ltifr-1",
                "suggestion_type": "process_improvement",
                "severity": "high",
                "title": "Add LTIFR deviation monitoring",
                "detail": "The supporting deviation log shows LTIFR threshold breaches.",
                "proposed_text": (
                    "4.2. Monitoring and Remediation of LTIFR Deviations "
                    "4.2.1 Identification and Thresholds LTIFR deviations are identified monthly. "
                    "4.2.2 Escalation Protocol The Health and Safety team escalates all breaches. "
                    "4.2.3 Remediation and Evidencing The team owns remediation through closure."
                ),
                "target_anchor_id": "incident-anchor",
                "target_text": "The Incident Manager reviews major incidents daily.",
                "review_status": "accepted",
                "source_references": [{"filename": "deviation_log.xlsx", "row_index": 4}],
            }
        ],
        process_steps=[],
        corpus_map={"general_relationships": [{"relationship_type": "support_to_procedure"}]},
        style_profile={"body_font": "Arial"},
        pipeline_id="case-1:stage2",
        stage1_cost=None,
    )

    assert result.status == "success"
    output_bytes = store.get_output_content("case-1", result.docx_file_id or "")
    with zipfile.ZipFile(BytesIO(output_bytes)) as package:
        document_xml = package.read("word/document.xml").decode("utf-8")

    assert "LTIFR" in document_xml
    assert "4.2.1" not in document_xml
    assert "4.2.2" not in document_xml
    assert "Escalation Protocol" not in document_xml


def test_generate_outputs_does_not_anchor_additions_to_short_header_fragments(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from utils.sop_processing import output_generator

    store = _memory_store()
    monkeypatch.setattr(output_generator, "get_store", lambda: store)
    monkeypatch.setattr(output_generator, "call_llm", _fake_llm_rewrite_unavailable)

    result = output_generator.generate_outputs(
        case_id="case-1",
        original_sop_bytes=_docx_with_short_header_and_matching_body(),
        suggestions=[
            {
                "suggestion_id": "short-header-1",
                "suggestion_type": "process_improvement",
                "severity": "medium",
                "title": "Clarify KYC refresh evidence",
                "detail": "The supporting source requires the date of review to be retained.",
                "proposed_text": (
                    "Add a procedure step owned by Investment Advisor. "
                    "The step should describe reviewing KYC information and recording the date of review "
                    "and identify retained evidence: KYC refresh record."
                ),
                "review_status": "accepted",
                "source_references": [
                    {"document_id": "kyc-source", "filename": "KYC_Source.xlsx"}
                ],
            }
        ],
        process_steps=[],
        corpus_map={"general_relationships": [{"relationship_type": "source_to_procedure"}]},
        style_profile={"body_font": "Arial"},
        pipeline_id="case-1:stage2",
        stage1_cost=None,
    )

    assert result.status == "success"
    output_bytes = store.get_output_content("case-1", result.docx_file_id or "")
    assert output_bytes
    with zipfile.ZipFile(BytesIO(output_bytes)) as package:
        document_xml = package.read("word/document.xml").decode("utf-8")

    root = ET.fromstring(document_xml)
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    paragraphs = []
    for paragraph in root.findall(".//w:p", ns):
        text = "".join((node.text or "") for node in paragraph.findall(".//w:t", ns))
        if text:
            paragraphs.append(text)
    date_index = paragraphs.index("Date")
    body_index = paragraphs.index("The Investment Advisor reviews KYC information annually.")
    insertion_index = paragraphs.index(
        "The Investment Advisor reviews KYC information and recording the date of review. Retained evidence includes KYC refresh record."
    )
    author_index = paragraphs.index("Author")
    assert date_index < body_index < insertion_index < author_index


def test_generate_outputs_applies_centimeter_margins_from_style_profile(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from utils.sop_processing import output_generator

    store = _memory_store()
    monkeypatch.setattr(output_generator, "get_store", lambda: store)
    monkeypatch.setattr(output_generator, "call_llm", _fake_llm)

    result = output_generator.generate_outputs(
        case_id="case-1",
        original_sop_bytes=_docx_bytes(),
        suggestions=[_suggestion()],
        process_steps=[],
        corpus_map={"risk_to_control_map": [{"risk_id": "R-1"}]},
        style_profile={
            "body_font": "Arial",
            "margin_top_cm": 1.25,
            "margin_bottom_cm": 1.5,
            "margin_left_cm": 1.75,
            "margin_right_cm": 2.0,
        },
        pipeline_id="case-1:stage2",
        stage1_cost=None,
    )

    assert result.status == "success"
    output_bytes = store.get_output_content("case-1", result.docx_file_id or "")
    assert output_bytes
    rendered = Document(BytesIO(output_bytes))
    section = rendered.sections[0]
    assert round(section.top_margin.cm, 2) == 1.25
    assert round(section.bottom_margin.cm, 2) == 1.5
    assert round(section.left_margin.cm, 2) == 1.75
    assert round(section.right_margin.cm, 2) == 2.0


def test_generate_outputs_falls_back_to_process_steps_for_diagram(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from utils.sop_processing import output_generator

    store = _memory_store()
    monkeypatch.setattr(output_generator, "get_store", lambda: store)
    monkeypatch.setattr(output_generator, "call_llm", _fake_llm_without_swimlane)

    result = output_generator.generate_outputs(
        case_id="case-1",
        original_sop_bytes=_docx_bytes(),
        suggestions=[_suggestion()],
        process_steps=[
            {
                "step_id": "PS-01",
                "actor": "Relationship Manager",
                "action": "Collect KYC documents",
            },
            {
                "step_id": "PS-02",
                "actor": "Compliance Officer",
                "action": "Approve exception",
            },
        ],
        corpus_map={"risk_to_control_map": [{"risk_id": "R-1"}]},
        style_profile={"primary_colour_hex": "#00338D"},
        pipeline_id="case-1:stage2",
        stage1_cost=None,
    )

    assert result.status == "success"
    assert result.swimlane_spec is not None
    assert [lane.label for lane in result.swimlane_spec.lanes] == [
        "Relationship Manager",
        "Compliance Officer",
    ]
    assert store.get_output_content("case-1", result.diagram_png_file_id or "")[:8] == b"\x89PNG\r\n\x1a\n"
    assert store.get_output_content("case-1", result.diagram_pdf_file_id or "")[:4] == b"%PDF"


def test_generate_outputs_ignores_rejected_suggestions(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from utils.sop_processing import output_generator

    store = _memory_store()
    monkeypatch.setattr(output_generator, "get_store", lambda: store)
    monkeypatch.setattr(output_generator, "call_llm", _fake_llm_swimlane_from_as_is_sop)

    result = output_generator.generate_outputs(
        case_id="case-1",
        original_sop_bytes=_docx_bytes(),
        suggestions=[_suggestion(status="rejected")],
        process_steps=[],
        corpus_map=None,
        style_profile=None,
        pipeline_id="case-1:stage2",
        stage1_cost=None,
    )

    assert result.sections_rewritten == 0
    assert result.swimlane_spec is not None
    assert result.swimlane_spec.title == "As-is SOP swimlane"
    output_bytes = store.get_output_content("case-1", result.docx_file_id or "")
    assert output_bytes
    with zipfile.ZipFile(BytesIO(output_bytes)) as package:
        document_xml = package.read("word/document.xml").decode("utf-8")
    assert "<w:del " not in document_xml
    assert "<w:ins " not in document_xml


def test_generate_outputs_caps_stage2_rewrite_calls_for_large_batches(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from utils.sop_processing import output_generator

    store = _memory_store()
    rewrite_prompts: list[str] = []

    def fake_llm(prompt: str, schema_name: str, **_kwargs: Any) -> LLMResult:
        if schema_name == "sop_section_rewrite":
            rewrite_prompts.append(prompt)
            return LLMResult(status="success", output={"rewritten_text": "Rewritten accepted change."})
        if schema_name == "swimlane_extraction":
            return LLMResult(status="success", output={"title": "Invalid", "lanes": [], "steps": []})
        raise AssertionError(f"Unexpected schema_name {schema_name}")

    monkeypatch.setattr(output_generator, "get_store", lambda: store)
    monkeypatch.setattr(output_generator, "call_llm", fake_llm)
    monkeypatch.setenv("DOCUMENT_UPLIFT_STAGE2_REWRITE_LIMIT", "2")

    suggestions = [
        {
            "suggestion_id": f"sug-{index}",
            "suggestion_type": "process_improvement",
            "severity": "medium",
            "title": f"Suggestion {index}",
            "detail": "The source document shows a required process update.",
            "original_text": "",
            "proposed_text": f"Apply accepted process update {index}.",
            "review_status": "accepted",
            "source_references": [{"filename": "evidence.xlsx", "row_index": index}],
        }
        for index in range(4)
    ]

    result = output_generator.generate_outputs(
        case_id="case-1",
        original_sop_bytes=_process_docx_bytes(),
        suggestions=suggestions,
        process_steps=[],
        corpus_map={"general_relationships": [{"relationship_type": "support_to_procedure"}]},
        style_profile={"body_font": "Arial"},
        pipeline_id="case-1:stage2",
        stage1_cost=None,
    )

    assert result.status == "success"
    assert len(rewrite_prompts) == 2


def test_generate_outputs_merges_same_anchor_procedure_additions(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from utils.sop_processing import output_generator

    store = _memory_store()

    def fake_llm(prompt: str, schema_name: str, **_kwargs: Any) -> LLMResult:
        if schema_name == "swimlane_extraction":
            return LLMResult(status="success", output={"title": "Invalid", "lanes": [], "steps": []})
        raise AssertionError(f"Unexpected schema_name {schema_name}")

    monkeypatch.setattr(output_generator, "get_store", lambda: store)
    monkeypatch.setattr(output_generator, "call_llm", fake_llm)
    monkeypatch.setenv("DOCUMENT_UPLIFT_STAGE2_REWRITE_LIMIT", "0")

    target_text = "The Service Desk logs incident tickets before triage."
    suggestions = [
        {
            "suggestion_id": "merge-1",
            "suggestion_type": "process_improvement",
            "severity": "medium",
            "title": "Add evidence capture",
            "detail": "The support file requires incident evidence capture.",
            "target_type": "procedure_step",
            "target_anchor_id": "incident-anchor",
            "target_text": target_text,
            "proposed_text": "The Service Desk captures incident evidence before triage.",
            "review_status": "accepted",
            "source_references": [{"filename": "evidence.xlsx", "row_index": 1}],
        },
        {
            "suggestion_id": "merge-2",
            "suggestion_type": "process_improvement",
            "severity": "medium",
            "title": "Add ticket validation",
            "detail": "The support file requires ticket validation.",
            "target_type": "procedure_step",
            "target_anchor_id": "incident-anchor",
            "target_text": target_text,
            "proposed_text": "The Service Desk validates ticket completeness before assignment.",
            "review_status": "accepted",
            "source_references": [{"filename": "evidence.xlsx", "row_index": 2}],
        },
    ]

    result = output_generator.generate_outputs(
        case_id="case-1",
        original_sop_bytes=_process_docx_bytes(),
        suggestions=suggestions,
        process_steps=[],
        corpus_map={"general_relationships": [{"relationship_type": "support_to_procedure"}]},
        style_profile={"body_font": "Arial"},
        pipeline_id="case-1:stage2",
        stage1_cost=None,
    )

    assert result.status == "success"
    output_bytes = store.get_output_content("case-1", result.docx_file_id or "")
    with zipfile.ZipFile(BytesIO(output_bytes)) as package:
        document_xml = package.read("word/document.xml").decode("utf-8")
        comments_xml = package.read("word/comments.xml").decode("utf-8")

    assert "captures incident evidence" in document_xml
    assert "validates ticket completeness" in document_xml
    assert document_xml.count("<w:ins ") == 1
    assert comments_xml.count("<w:comment ") == 1


def test_generate_outputs_falls_back_to_standalone_docx_without_original_docx(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from utils.sop_processing import output_generator

    store = _memory_store()
    monkeypatch.setattr(output_generator, "get_store", lambda: store)
    monkeypatch.setattr(output_generator, "call_llm", _fake_llm)

    result = output_generator.generate_outputs(
        case_id="case-1",
        original_sop_bytes=None,
        suggestions=[_suggestion()],
        process_steps=[],
        corpus_map=None,
        style_profile=None,
        pipeline_id="case-1:stage2",
        stage1_cost=None,
    )

    assert result.status == "success"
    assert result.warnings
    output = store.get_case("case-1")["outputs"][0]
    assert output["output_mode"] == "standalone"
    output_bytes = store.get_output_content("case-1", result.docx_file_id or "")
    assert output_bytes
    with zipfile.ZipFile(BytesIO(output_bytes)) as package:
        document_xml = package.read("word/document.xml").decode("utf-8")
    assert "Clarify recency requirement for address proof" in document_xml
    assert "not older than 3 months" in document_xml
