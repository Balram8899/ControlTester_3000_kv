from __future__ import annotations

import importlib
from io import BytesIO
from types import ModuleType
from typing import Any

import pytest
from openpyxl import Workbook

from utils.services.schemas import LLMResult


HEADERS = ["Risk ID", "Control ID", "Owner", "Frequency", "Evidence"]


def load_excel_pipeline() -> ModuleType:
    try:
        return importlib.import_module("utils.services.excel_pipeline")
    except ModuleNotFoundError as exc:
        pytest.fail(f"utils.services.excel_pipeline is missing: {exc}")


def workbook_bytes(workbook: Workbook) -> bytes:
    payload = BytesIO()
    workbook.save(payload)
    return payload.getvalue()


def append_headers(sheet: Any, row: int = 1) -> None:
    for index, header in enumerate(HEADERS, start=1):
        sheet.cell(row=row, column=index, value=header)


def append_rcm_rows(
    sheet: Any,
    count: int,
    start_row: int = 2,
    owner: str | None = "Operations",
) -> None:
    for offset in range(count):
        row = start_row + offset
        sheet.cell(row=row, column=1, value=f"R-{offset + 1}")
        sheet.cell(row=row, column=2, value=f"C-{offset + 1}")
        sheet.cell(row=row, column=3, value=owner)
        sheet.cell(row=row, column=4, value="Quarterly")
        sheet.cell(row=row, column=5, value="Access review evidence")


def install_schema_detector(
    monkeypatch: pytest.MonkeyPatch,
    excel_pipeline: ModuleType,
) -> list[str]:
    calls: list[str] = []

    def fake_call_llm(
        prompt: str,
        schema_name: str,
        response_schema: dict,
        pipeline_id: str,
        budget_remaining: int,
    ) -> LLMResult:
        calls.append(schema_name)
        headers = [
            line[2:].strip()
            for line in prompt.splitlines()
            if line.startswith("- ")
        ]
        role_map = {
            "risk id": "risk_id",
            "control id": "control_id",
            "owner": "owner",
            "frequency": "frequency",
            "evidence": "evidence_artifact",
        }
        return LLMResult(
            status="success",
            output={
                "columns": [
                    {
                        "column_name": header,
                        "classified_as": role_map.get(header.lower(), "other"),
                        "confidence": 0.99,
                    }
                    for header in headers
                ]
            },
        )

    monkeypatch.setattr(excel_pipeline, "call_llm", fake_call_llm)
    return calls


def test_500_row_rcm_full_coverage(monkeypatch: pytest.MonkeyPatch) -> None:
    excel_pipeline = load_excel_pipeline()
    install_schema_detector(monkeypatch, excel_pipeline)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "RCM"
    append_headers(sheet)
    append_rcm_rows(sheet, 500)

    result = excel_pipeline.process_excel(
        workbook_bytes(workbook),
        "rcm.xlsx",
        "file-1",
        "pipe-1",
        budget_remaining=1,
    )

    assert result.status == "success"
    assert result.total_rows_assessed == 500


def test_ownership_gap_detection(monkeypatch: pytest.MonkeyPatch) -> None:
    excel_pipeline = load_excel_pipeline()
    install_schema_detector(monkeypatch, excel_pipeline)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "RCM"
    append_headers(sheet)
    append_rcm_rows(sheet, 200, owner=None)

    result = excel_pipeline.process_excel(
        workbook_bytes(workbook),
        "rcm.xlsx",
        "file-1",
        "pipe-1",
        budget_remaining=1,
    )

    gaps = [gap for sheet_result in result.sheets for gap in sheet_result.gaps]
    assert len([gap for gap in gaps if gap.gap_type == "ownership_gap"]) == 200


def test_merged_cells_propagated(monkeypatch: pytest.MonkeyPatch) -> None:
    excel_pipeline = load_excel_pipeline()
    install_schema_detector(monkeypatch, excel_pipeline)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "RCM"
    append_headers(sheet)
    append_rcm_rows(sheet, 5, owner=None)
    sheet["C2"] = "Operations"
    sheet.merge_cells("C2:C6")

    result = excel_pipeline.process_excel(
        workbook_bytes(workbook),
        "rcm.xlsx",
        "file-1",
        "pipe-1",
        budget_remaining=1,
    )

    gaps = [gap for sheet_result in result.sheets for gap in sheet_result.gaps]
    assert [gap for gap in gaps if gap.gap_type == "ownership_gap"] == []


def test_header_in_row_2_detected(monkeypatch: pytest.MonkeyPatch) -> None:
    excel_pipeline = load_excel_pipeline()
    install_schema_detector(monkeypatch, excel_pipeline)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "RCM"
    sheet["A1"] = "Risk and Control Matrix"
    append_headers(sheet, row=2)
    append_rcm_rows(sheet, 5, start_row=3)

    result = excel_pipeline.process_excel(
        workbook_bytes(workbook),
        "rcm.xlsx",
        "file-1",
        "pipe-1",
        budget_remaining=1,
    )

    assert result.sheets[0].row_count == 5
    assert result.sheets[0].schema[0].column_name == "Risk ID"


def test_merged_title_row_does_not_mask_row_2_headers(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    excel_pipeline = load_excel_pipeline()
    install_schema_detector(monkeypatch, excel_pipeline)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "RCM"
    sheet["A1"] = "Wealth Management Risk & Controls Matrix"
    sheet.merge_cells("A1:E1")
    append_headers(sheet, row=2)
    append_rcm_rows(sheet, 5, start_row=3)

    result = excel_pipeline.process_excel(
        workbook_bytes(workbook),
        "rcm.xlsx",
        "file-1",
        "pipe-1",
        budget_remaining=1,
    )

    assert result.status == "success"
    assert result.sheets[0].row_count == 5
    assert result.sheets[0].schema[0].column_name == "Risk ID"
    assert result.corpus_map is not None
    assert result.corpus_map.risk_to_control_map[0]["control_id"] == "C-1"


def test_multi_sheet_workbook(monkeypatch: pytest.MonkeyPatch) -> None:
    excel_pipeline = load_excel_pipeline()
    install_schema_detector(monkeypatch, excel_pipeline)
    workbook = Workbook()
    for index in range(3):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.title = f"RCM {index + 1}"
        append_headers(sheet)
        append_rcm_rows(sheet, 5)

    result = excel_pipeline.process_excel(
        workbook_bytes(workbook),
        "multi.xlsx",
        "file-1",
        "pipe-1",
        budget_remaining=3,
    )

    assert result.status == "success"
    assert len(result.sheets) == 3


def test_empty_sheet_skipped(monkeypatch: pytest.MonkeyPatch) -> None:
    excel_pipeline = load_excel_pipeline()
    install_schema_detector(monkeypatch, excel_pipeline)
    workbook = Workbook()
    workbook.active.title = "Empty"
    sheet = workbook.create_sheet("RCM")
    append_headers(sheet)
    append_rcm_rows(sheet, 5)

    result = excel_pipeline.process_excel(
        workbook_bytes(workbook),
        "with_empty.xlsx",
        "file-1",
        "pipe-1",
        budget_remaining=1,
    )

    assert result.status == "success"
    assert [sheet_result.sheet_name for sheet_result in result.sheets] == ["RCM"]


def test_budget_zero_returns_partial(monkeypatch: pytest.MonkeyPatch) -> None:
    excel_pipeline = load_excel_pipeline()
    calls = install_schema_detector(monkeypatch, excel_pipeline)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "RCM"
    append_headers(sheet)
    append_rcm_rows(sheet, 5)

    result = excel_pipeline.process_excel(
        workbook_bytes(workbook),
        "rcm.xlsx",
        "file-1",
        "pipe-1",
        budget_remaining=0,
    )

    assert result.status == "partial"
    assert calls == []
    assert result.sheets[0].status == "skipped_budget"


def test_corpus_map_populated(monkeypatch: pytest.MonkeyPatch) -> None:
    excel_pipeline = load_excel_pipeline()
    install_schema_detector(monkeypatch, excel_pipeline)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "RCM"
    append_headers(sheet)
    append_rcm_rows(sheet, 5)

    result = excel_pipeline.process_excel(
        workbook_bytes(workbook),
        "rcm.xlsx",
        "file-1",
        "pipe-1",
        budget_remaining=1,
    )

    assert result.corpus_map is not None
    assert result.corpus_map.risk_to_control_map[0]["risk_id"] == "R-1"
    assert result.corpus_map.evidence_to_control_map[0]["control_id"] == "C-1"


def test_row_gaps_convert_to_capped_suggestions(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    excel_pipeline = load_excel_pipeline()
    install_schema_detector(monkeypatch, excel_pipeline)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "RCM"
    append_headers(sheet)
    append_rcm_rows(sheet, 20, owner=None)

    result = excel_pipeline.process_excel(
        workbook_bytes(workbook),
        "rcm.xlsx",
        "file-1",
        "pipe-1",
        budget_remaining=1,
    )

    assert len(result.suggestions) == 12
    assert {suggestion.suggestion_type for suggestion in result.suggestions} == {
        "ownership_gap"
    }
    assert all(suggestion.review_status == "pending" for suggestion in result.suggestions)
    assert result.suggestions[0].source_references[0].sheet_name == "RCM"
    assert result.suggestions[0].source_references[0].row_index == 2


def test_generic_structured_records_preserve_raw_attributes_and_emit_facts(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    excel_pipeline = load_excel_pipeline()
    install_schema_detector(monkeypatch, excel_pipeline)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Operational Evidence"
    headers = ["Process Activity", "Processor", "Evidence", "Ticket Status"]
    for column, header in enumerate(headers, start=1):
        sheet.cell(row=1, column=column, value=header)
    for row in range(2, 7):
        sheet.cell(row=row, column=1, value="Review exception queue")
        sheet.cell(row=row, column=2, value="Operations" if row > 2 else None)
        sheet.cell(row=row, column=3, value="Queue review log")
        sheet.cell(row=row, column=4, value="Open")

    result = excel_pipeline.process_excel(
        workbook_bytes(workbook),
        "ops.xlsx",
        "file-ops",
        "pipe-1",
        budget_remaining=1,
    )

    first_record = result.sheets[0].records[0]
    assert first_record["_raw_attributes"]["Processor"] == ""
    assert first_record["_semantic_roles"]["Processor"]["role"] == "owner"
    assert result.facts[0].raw_attributes["Process Activity"] == "Review exception queue"
    assert any(suggestion.suggestion_type == "ownership_gap" for suggestion in result.suggestions)


def test_generic_structured_records_do_not_require_risk_or_control_columns(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    excel_pipeline = load_excel_pipeline()
    install_schema_detector(monkeypatch, excel_pipeline)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Payments"
    headers = ["Activity", "Prepared By", "Approved By", "Evidence"]
    for column, header in enumerate(headers, start=1):
        sheet.cell(row=1, column=column, value=header)
    for row in range(2, 7):
        sheet.cell(row=row, column=1, value="Approve payment file")
        sheet.cell(row=row, column=2, value="A. Shah")
        sheet.cell(row=row, column=3, value="A. Shah")
        sheet.cell(row=row, column=4, value="Payment approval report")

    result = excel_pipeline.process_excel(
        workbook_bytes(workbook),
        "payments.xlsx",
        "file-pay",
        "pipe-1",
        budget_remaining=1,
    )

    assert result.status == "success"
    assert result.corpus_map is not None
    assert result.corpus_map.risk_to_control_map == []
    assert any(
        "segregation" in suggestion.detail.lower()
        or "incompatible" in suggestion.detail.lower()
        for suggestion in result.suggestions
    )


def test_deviation_log_rows_emit_generic_issue_signal_suggestions(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    excel_pipeline = load_excel_pipeline()
    install_schema_detector(monkeypatch, excel_pipeline)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Metric Deviations"
    headers = ["Metric", "Target", "Actual", "Deviation", "Status", "Root Cause", "Recommended Action", "Owner"]
    for column, header in enumerate(headers, start=1):
        sheet.cell(row=1, column=column, value=header)
    for row in range(2, 7):
        sheet.cell(row=row, column=1, value="Supplier assessments completed")
        sheet.cell(row=row, column=2, value="100%")
        sheet.cell(row=row, column=3, value="42%")
        sheet.cell(row=row, column=4, value="-58%")
        sheet.cell(row=row, column=5, value="Open")
        sheet.cell(row=row, column=6, value="Questionnaire not aligned to reporting requirements")
        sheet.cell(row=row, column=7, value="Update questionnaire and track remediation")
        sheet.cell(row=row, column=8, value="Procurement")

    result = excel_pipeline.process_excel(
        workbook_bytes(workbook),
        "metric_deviation_log.xlsx",
        "file-dev",
        "pipe-1",
        budget_remaining=1,
    )

    assert result.status == "success"
    issue_suggestions = [
        suggestion
        for suggestion in result.suggestions
        if "deviation" in suggestion.title.lower()
    ]
    assert issue_suggestions
    assert issue_suggestions[0].source_references[0].filename == "metric_deviation_log.xlsx"
    assert issue_suggestions[0].source_references[0].sheet_name == "Metric Deviations"
