from __future__ import annotations

import importlib
from io import BytesIO
from types import ModuleType

import pytest
from docx import Document
from openpyxl import Workbook


def load_conversion() -> ModuleType:
    try:
        return importlib.import_module("utils.services.conversion")
    except ModuleNotFoundError as exc:
        pytest.fail(f"utils.services.conversion is missing: {exc}")


def make_docx_bytes() -> bytes:
    document = Document()
    document.add_heading("Client onboarding", level=1)
    document.add_paragraph(
        "The operations team validates identity documents, records approval "
        "evidence, and confirms quarterly control ownership before access is granted."
    )
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Control ID"
    table.rows[0].cells[1].text = "Frequency"
    table.rows[1].cells[0].text = "C-1"
    table.rows[1].cells[1].text = "Monthly"
    payload = BytesIO()
    document.save(payload)
    return payload.getvalue()


def make_xlsx_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "RCM"
    sheet.append(["Risk ID", "Control ID", "Owner", "Frequency", "Evidence"])
    for index in range(1, 7):
        sheet.append(
            [
                f"R-{index}",
                f"C-{index}",
                "Operations",
                "Quarterly",
                "Access review evidence retained in the GRC repository.",
            ]
        )
    payload = BytesIO()
    workbook.save(payload)
    return payload.getvalue()


def test_docx_conversion_returns_markdown() -> None:
    conversion = load_conversion()

    result = conversion.convert_document(
        make_docx_bytes(),
        "client_onboarding.docx",
        "file-docx",
        "process_doc",
    )

    assert result.status == "success"
    assert result.file_type == "docx"
    assert result.markdown
    assert "# Client onboarding" in result.markdown
    assert "| Control ID | Frequency |" in result.markdown
    assert result.page_count > 0


def test_docx_conversion_interleaves_tables_at_original_position() -> None:
    conversion = load_conversion()
    document = Document()
    document.add_heading("Client onboarding", level=1)
    document.add_paragraph(
        "The maker captures the client profile and attaches onboarding evidence "
        "before the control table is completed for checker review."
    )
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Control ID"
    table.rows[0].cells[1].text = "Evidence"
    table.rows[1].cells[0].text = "C-2"
    table.rows[1].cells[1].text = "KYC checklist"
    document.add_paragraph(
        "The checker records sign-off after reviewing the embedded control table "
        "and retains approval evidence in the case file."
    )
    payload = BytesIO()
    document.save(payload)

    result = conversion.convert_document(
        payload.getvalue(),
        "client_onboarding.docx",
        "file-docx",
        "process_doc",
    )

    assert result.status == "success"
    maker_index = result.markdown.index("The maker captures the client profile")
    table_index = result.markdown.index("| Control ID | Evidence |")
    checker_index = result.markdown.index("The checker records sign-off")
    assert maker_index < table_index < checker_index


def test_docx_procedure_extracts_style_profile() -> None:
    conversion = load_conversion()

    result = conversion.convert_document(
        make_docx_bytes(),
        "procedure.docx",
        "file-procedure",
        "procedure",
    )

    assert result.status == "success"
    assert result.style_profile is not None


def test_docx_policy_extracts_style_profile() -> None:
    conversion = load_conversion()

    result = conversion.convert_document(
        make_docx_bytes(),
        "policy.docx",
        "file-policy",
        "policy",
    )

    assert result.status == "success"
    assert result.style_profile is not None


def test_xlsx_conversion_no_style_profile() -> None:
    conversion = load_conversion()

    result = conversion.convert_document(
        make_xlsx_bytes(),
        "risk_control_matrix.xls",
        "file-xlsx",
        "rcm",
    )

    assert result.status == "success"
    assert result.file_type == "xlsx"
    assert "## RCM" in result.markdown
    assert "| Risk ID | Control ID | Owner | Frequency | Evidence |" in result.markdown
    assert result.style_profile is None


def test_unsupported_extension_returns_failed() -> None:
    conversion = load_conversion()

    result = conversion.convert_document(
        b"not supported",
        "installer.exe",
        "file-unsupported",
        "evidence",
    )

    assert result.status == "failed"
    assert result.error == "Unsupported file type: exe"


def test_corrupt_pdf_returns_partial(monkeypatch: pytest.MonkeyPatch) -> None:
    conversion = load_conversion()

    def fake_convert_pdf(_file_bytes: bytes) -> str:
        return ("Readable text remains available. " * 8) + "word/styles.xml"

    monkeypatch.setattr(conversion, "_convert_pdf", fake_convert_pdf)

    result = conversion.convert_document(
        b"%PDF-1.7",
        "bad.pdf",
        "file-pdf",
        "evidence",
    )

    assert result.status == "partial"
    assert result.file_type == "pdf"
    assert result.looks_corrupt is True
    assert result.markdown
