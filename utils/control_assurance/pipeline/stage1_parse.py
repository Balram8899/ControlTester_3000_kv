from __future__ import annotations

import io
import uuid
from datetime import datetime, timezone

import openpyxl

from utils.control_assurance import ct_db
from utils.control_assurance.celery_app import celery_app
from utils.control_assurance.control_setup import (
    build_control_doc,
    empty_conclusions,
    empty_sampling,
    empty_testing_methods,
    normalize_header,
    normalize_sampling_mode,
    now_iso,
    parse_bool,
    row_to_control_input,
)
from utils.control_assurance.ct_gridfs import download_from_gridfs


_STEP_LABELS = ["A", "B", "C", "D", "E", "F"]
_SAMPLING_MODE_MAP = {
    "walkthrough": "walkthrough",
    "sample": "sample",
    "both": "both",
    "none": "none",
}


def _now() -> str:
    return now_iso()


def _empty_sampling(mode: str) -> dict:
    return empty_sampling(mode)


def _empty_conclusions() -> dict:
    return empty_conclusions()


def _empty_testing_methods() -> dict:
    return empty_testing_methods()


def _parse_bool(value: str) -> bool:
    return parse_bool(value)


def _parse_sampling_mode(value: str) -> str:
    return normalize_sampling_mode(value)


def _is_new_control_input(headers: list[str]) -> bool:
    normalized = {normalize_header(header) for header in headers}
    return "riskstatement" in normalized and "teststeps" in normalized


def _parse_and_create_controls(session_id: str, gridfs_file_id: str) -> None:
    db = ct_db._get_db()
    db.ct_sessions.update_one(
        {"_id": session_id},
        {
            "$set": {
                "stage_checkpoint": {
                    "stage": "parse_template",
                    "step": "downloading",
                    "updated_at": _now(),
                }
            }
        },
    )

    content = download_from_gridfs(gridfs_file_id)
    workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    worksheet = workbook.active
    headers = [str(cell.value or "").strip() for cell in worksheet[1]]

    now = _now()
    controls_created = 0

    if _is_new_control_input(headers):
        for row in worksheet.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue

            doc = build_control_doc(session_id, row_to_control_input(headers, row), now=now)
            db.ct_controls.insert_one(doc)
            controls_created += 1
            db.ct_sessions.update_one(
                {"_id": session_id},
                {
                    "$set": {
                        "stage_checkpoint": {
                            "stage": "parse_template",
                            "step": f"parsed_{controls_created}_controls",
                            "updated_at": _now(),
                        }
                    }
                },
            )
        db.ct_sessions.update_one(
            {"_id": session_id},
            {
                "$set": {
                    "stage": "analysing",
                    "controls_finalized": False,
                    "stage_checkpoint": {
                        "stage": "parse_template",
                        "step": "complete",
                        "controls_found": controls_created,
                        "updated_at": _now(),
                    },
                    "updated_at": _now(),
                }
            },
        )
        return

    for row in worksheet.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue

        padded = list(row) + [""] * max(0, 23 - len(row))
        (
            control_id,
            control_name,
            control_type,
            domain,
            framework_ref,
            inherent_risk,
            owner,
            frequency,
            prior_period,
            walkthrough_raw,
            sampling_mode_raw,
        ) = [str(value or "").strip() for value in padded[:11]]

        test_steps = []
        for index, label in enumerate(_STEP_LABELS):
            col_start = 11 + (index * 2)
            description = str(padded[col_start] or "").strip()
            evidence_required = str(padded[col_start + 1] or "").strip()
            if description:
                test_steps.append(
                    {
                        "step_id": str(uuid.uuid4()),
                        "attribute_id": label,
                        "label": label,
                        "test_attribute": description,
                        "description": description,
                        "evidence_required": evidence_required,
                    }
                )

        doc = {
            "_id": str(uuid.uuid4()),
            "session_id": session_id,
            "control_id": control_id,
            "control_name": control_name,
            "control_type": control_type,
            "domain": domain,
            "framework_reference": framework_ref,
            "inherent_risk_rating": inherent_risk or "Medium",
            "control_owner": owner,
            "frequency": frequency,
            "prior_period_result": prior_period or "N/A",
            "walkthrough_performed": _parse_bool(walkthrough_raw),
            "risk": "",
            "test_steps": test_steps,
            "field_sources": {},
            "controls_finalized": False,
            "finalized_at": None,
            "sampling": _empty_sampling(_parse_sampling_mode(sampling_mode_raw)),
            "evidence_files": [],
            "sample_results": [],
            "todi_results": {},
            "exceptions": [],
            "conclusions": _empty_conclusions(),
            "testing_methods": _empty_testing_methods(),
            "workbook_output_id": None,
            "status": "pending",
            "created_at": now,
            "updated_at": now,
        }
        db.ct_controls.insert_one(doc)
        controls_created += 1
        db.ct_sessions.update_one(
            {"_id": session_id},
            {
                "$set": {
                    "stage_checkpoint": {
                        "stage": "parse_template",
                        "step": f"parsed_{controls_created}_controls",
                        "updated_at": _now(),
                    }
                }
            },
        )

    db.ct_sessions.update_one(
        {"_id": session_id},
        {
            "$set": {
                "stage": "analysing",
                "controls_finalized": False,
                "stage_checkpoint": {
                    "stage": "parse_template",
                    "step": "complete",
                    "controls_found": controls_created,
                    "updated_at": _now(),
                },
                "updated_at": _now(),
            }
        },
    )


@celery_app.task(
    name="ct.parse_template",
    queue="ct_pipeline",
    acks_late=True,
    reject_on_worker_lost=True,
)
def parse_template(session_id: str, gridfs_file_id: str) -> dict:
    try:
        _parse_and_create_controls(session_id, gridfs_file_id)
        from utils.control_assurance.pipeline.stage2_review import llm_review

        task = llm_review.apply_async(args=[session_id], queue="ct_pipeline")
        ct_db._get_db().ct_sessions.update_one(
            {"_id": session_id},
            {"$set": {"celery_task_id": task.id, "updated_at": _now()}},
        )
        return {"status": "ok", "session_id": session_id, "next_task": task.id}
    except Exception as exc:
        ct_db._get_db().ct_sessions.update_one(
            {"_id": session_id},
            {
                "$set": {
                    "stage": "failed",
                    "stage_checkpoint": {"error": str(exc), "updated_at": _now()},
                }
            },
        )
        raise
