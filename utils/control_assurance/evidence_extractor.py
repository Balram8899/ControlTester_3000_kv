from __future__ import annotations

import csv
import io
import zipfile


def extract_file_content(content: bytes, file_type: str, filename: str) -> tuple[str, dict]:
    if file_type == "pdf":
        return _extract_pdf(content)
    if file_type == "excel":
        return _extract_excel(content)
    if file_type == "csv":
        return _extract_csv(content)
    if file_type == "docx":
        return _extract_docx(content)
    if file_type == "image":
        return _extract_image(content)
    if file_type == "txt":
        return content.decode("utf-8", errors="replace")[:8000], {}
    if file_type == "zip":
        return _extract_zip(content, filename)
    return "", {}


def _extract_pdf(content: bytes) -> tuple[str, dict]:
    try:
        import fitz

        doc = fitz.open(stream=content, filetype="pdf")
        text = "\n".join(page.get_text() for page in doc)
        return text[:8000], doc.metadata or {}
    except Exception:
        return "", {}


def _extract_excel(content: bytes) -> tuple[str, dict]:
    try:
        import openpyxl

        workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        worksheet = workbook.active
        rows = ["\t".join(str(value or "") for value in row) for row in worksheet.iter_rows(values_only=True)]
        return "\n".join(rows)[:8000], {}
    except Exception:
        return "", {}


def _extract_csv(content: bytes) -> tuple[str, dict]:
    try:
        text = content.decode("utf-8", errors="replace")
        rows = ["\t".join(row) for row in csv.reader(text.splitlines())]
        return "\n".join(rows)[:8000], {}
    except Exception:
        return "", {}


def _extract_docx(content: bytes) -> tuple[str, dict]:
    try:
        from docx import Document

        doc = Document(io.BytesIO(content))
        return "\n".join(paragraph.text for paragraph in doc.paragraphs)[:8000], {}
    except Exception:
        return "", {}


def _extract_image(content: bytes) -> tuple[str, dict]:
    try:
        from PIL import Image
        import pytesseract

        image = Image.open(io.BytesIO(content))
        return pytesseract.image_to_string(image)[:8000], dict(image.getexif() or {})
    except Exception:
        return "", {}


def _extract_zip(content: bytes, filename: str) -> tuple[str, dict]:
    texts = []
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as zf:
            for name in zf.namelist()[:20]:
                inner = zf.read(name)
                file_type = _ext_to_type(name.rsplit(".", 1)[-1].lower() if "." in name else "")
                text, _ = extract_file_content(inner, file_type, name)
                texts.append(f"[{name}]\n{text}")
    except Exception:
        return "", {"filename": filename}
    return "\n\n".join(texts)[:8000], {"filename": filename}


def _ext_to_type(ext: str) -> str:
    if ext in {"jpg", "jpeg", "png", "gif", "bmp", "tiff", "tif", "webp"}:
        return "image"
    if ext == "pdf":
        return "pdf"
    if ext in {"xlsx", "xls"}:
        return "excel"
    if ext == "csv":
        return "csv"
    if ext == "docx":
        return "docx"
    return "txt"
