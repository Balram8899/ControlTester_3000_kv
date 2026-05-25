from __future__ import annotations

import io
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image, ImageDraw, UnidentifiedImageError


REPO_ROOT = Path(__file__).resolve().parents[2]
WORKPAPER_TEMPLATE = REPO_ROOT / "utils" / "Testing sheet.xlsx"


TESTING_SHEET = "Test of controls"
OVERRIDE_SHEET = "Auditor override"
TESTING_GRID_DEFAULT_ROWS = 16
RED = "FF0000"
RED_SIDE = Side(style="thick", color=RED)
RED_BORDER = Border(left=RED_SIDE, right=RED_SIDE, top=RED_SIDE, bottom=RED_SIDE)


def _load_base_workbook() -> Workbook:
    if WORKPAPER_TEMPLATE.exists():
        try:
            return openpyxl.load_workbook(WORKPAPER_TEMPLATE)
        except Exception:
            pass

    workbook = Workbook()
    workbook.active.title = TESTING_SHEET
    workbook.create_sheet(OVERRIDE_SHEET)
    return workbook


def _sheet(workbook: Workbook, name: str):
    if name in workbook.sheetnames:
        return workbook[name]
    return workbook.create_sheet(name)


def _safe_cell_text(value: object, limit: int = 5000) -> str:
    if value is None:
        return ""
    return ILLEGAL_CHARACTERS_RE.sub("", str(value))[:limit]


def _period_text(session: dict) -> str:
    period = session.get("testing_period", {})
    start = period.get("from", "")
    end = period.get("to", "")
    if start and end:
        return f"{start} to {end}"
    return start or end


def _yes_no(value: object) -> str:
    return "Y" if bool(value) else "N"


def _control_description(control: dict) -> str:
    return (
        control.get("control_description")
        or control.get("description")
        or control.get("risk")
        or control.get("control_name", "")
    )


def _operating_effectiveness_testing(control: dict) -> bool:
    conclusions = control.get("conclusions", {})
    sampling = control.get("sampling", {})
    return bool(
        control.get("sample_results")
        or conclusions.get("oe")
        or sampling.get("mode") in {"sample", "both", "full"}
    )


def _testing_method_value(control: dict, key: str) -> str:
    return _yes_no(control.get("testing_methods", {}).get(key, False))


def _tickmark(value: object) -> str:
    normalized = str(value or "").strip()
    if not normalized:
        return ""
    upper = normalized.upper()
    if upper in {"PASS", "PASSED", "OK", "Y", "YES", "TRUE"}:
        return "√"
    return normalized


def _sample_notes(sample: dict, exceptions_by_sample: dict[int, list[dict]]) -> str:
    notes = [
        _safe_cell_text(step.get("notes", ""))
        for step in sample.get("step_results", [])
        if step.get("notes")
    ]
    for exception in exceptions_by_sample.get(int(sample.get("sample_num") or 0), []):
        ref = exception.get("ref", "")
        description = exception.get("description", "")
        if ref or description:
            notes.append(f"{ref}: {description}".strip(": "))
    return _safe_cell_text("; ".join(note for note in notes if note))


def _exception_summary(control: dict) -> str:
    exceptions = []
    for exception in control.get("exceptions", []):
        ref = exception.get("ref", "")
        sample_num = exception.get("sample_num", "")
        description = exception.get("description", "")
        root_cause = exception.get("root_cause", "")
        issue_ref = exception.get("issues_log_ref", "")
        parts = [
            f"{ref} sample {sample_num}".strip(),
            description,
            f"Root cause: {root_cause}" if root_cause else "",
            f"Issue: {issue_ref}" if issue_ref else "",
        ]
        exceptions.append(" - ".join(part for part in parts if part))
    return "; ".join(exceptions)


def _conclusion_description(control: dict) -> str:
    conclusions = control.get("conclusions", {})
    parts = [
        conclusions.get("rationale", ""),
        conclusions.get("testing_summary", ""),
        conclusions.get("d_and_i_statement", ""),
        _exception_summary(control),
    ]
    issue_refs = conclusions.get("issues_log_refs", [])
    if issue_refs:
        parts.append(f"Issue log refs: {', '.join(issue_refs)}")
    return _safe_cell_text("\n".join(part for part in parts if part), limit=12000)


def _copy_row_style(worksheet, source_row: int, target_row: int, max_column: int) -> None:
    for column in range(1, max_column + 1):
        source = worksheet.cell(source_row, column)
        target = worksheet.cell(target_row, column)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        if source.alignment:
            target.alignment = copy(source.alignment)


def _find_row(worksheet, text: str, column: int = 1) -> int | None:
    target = text.strip().lower()
    for row in range(1, worksheet.max_row + 1):
        value = worksheet.cell(row, column).value
        if isinstance(value, str) and value.strip().lower() == target:
            return row
    return None


def _find_row_contains(worksheet, text: str, column: int = 1) -> int | None:
    target = text.strip().lower()
    for row in range(1, worksheet.max_row + 1):
        value = worksheet.cell(row, column).value
        if isinstance(value, str) and target in value.strip().lower():
            return row
    return None


def _procedure_header_row(worksheet) -> int:
    return _find_row(worksheet, "Test step/ Testing Attribute", 1) or 32


def _testing_header_row(worksheet) -> int:
    return _find_row(worksheet, "Sample #", 1) or 47


def _tickmark_key_row(worksheet) -> int:
    return _find_row_contains(worksheet, "tick mark key", 1) or 66


def _conclusions_row(worksheet) -> int:
    return _find_row_contains(worksheet, "conclusions", 1) or 75


def _step_label(step: dict, index: int) -> str:
    return str(step.get("attribute_id") or step.get("label") or f"TA-{index:03d}").strip()


def _notes_column(attribute_count: int) -> int:
    return 4 + max(attribute_count, 1)


def _set_title(worksheet, cell: str, value: str) -> None:
    worksheet[cell] = value
    worksheet[cell].font = Font(bold=True, size=14)


def _set_label(worksheet, cell: str, value: str) -> None:
    worksheet[cell] = value
    worksheet[cell].font = Font(bold=True)
    worksheet[cell].fill = PatternFill("solid", fgColor="E7E6E6")


def _sample_step_labels(sample: dict) -> set[str]:
    labels = set()
    for step in sample.get("step_results", []):
        label = step.get("attribute_id") or step.get("label")
        if label:
            labels.add(str(label).upper())
    return labels


def _evidence_matches_sample(evidence: dict, step_labels: set[str]) -> bool:
    mapped = {str(label).upper() for label in evidence.get("mapped_step_labels", [])}
    return not step_labels or not mapped or bool(step_labels & mapped)


def _sample_evidence_records(control: dict, sample: dict) -> list[dict]:
    step_labels = _sample_step_labels(sample)
    records = []
    for evidence in control.get("evidence_files", []):
        if not _evidence_matches_sample(evidence, step_labels):
            continue
        records.append({**evidence, "record_type": "evidence"})
        for support in evidence.get("support_files", []):
            records.append(
                {
                    **support,
                    "record_type": "support",
                    "identified_value": support.get("comments") or support.get("support_type", ""),
                    "mapped_step_labels": evidence.get("mapped_step_labels", []),
                }
            )
    return records


def _annotation_boxes(evidence: dict, step_labels: set[str]) -> list[tuple[int, int, int, int]]:
    boxes = []
    for region in evidence.get("annotation_regions", []):
        label = str(region.get("step_label", "")).upper()
        if step_labels and label and label not in step_labels:
            continue
        if region.get("annotation_type") not in {"bbox", "image"}:
            continue
        bbox = region.get("bbox") or {}
        try:
            x = max(0, int(bbox.get("x", 0)))
            y = max(0, int(bbox.get("y", 0)))
            w = max(1, int(bbox.get("w", 0)))
            h = max(1, int(bbox.get("h", 0)))
        except (TypeError, ValueError):
            continue
        boxes.append((x, y, x + w, y + h))
    return boxes


def _annotated_image_stream(content: bytes, evidence: dict, step_labels: set[str]) -> io.BytesIO | None:
    if not content:
        return None
    try:
        image = Image.open(io.BytesIO(content)).convert("RGB")
    except (UnidentifiedImageError, OSError):
        return None

    draw = ImageDraw.Draw(image)
    boxes = _annotation_boxes(evidence, step_labels)
    if boxes:
        for box in boxes:
            draw.rectangle(box, outline=(255, 0, 0), width=4)
    else:
        draw.rectangle((0, 0, image.width - 1, image.height - 1), outline=(255, 0, 0), width=4)

    stream = io.BytesIO()
    image.save(stream, format="PNG")
    stream.seek(0)
    return stream


def _resize_for_sheet(image: XLImage, max_width: int = 720) -> None:
    if image.width <= max_width:
        return
    ratio = max_width / image.width
    image.width = int(image.width * ratio)
    image.height = int(image.height * ratio)


def _decode_preview(content: bytes, file_type: str) -> str:
    if not content:
        return ""
    if file_type in {"txt", "csv", "text", "conf"}:
        return _safe_cell_text(content.decode("utf-8", errors="replace"), limit=12000)
    if file_type == "excel":
        try:
            workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
            worksheet = workbook.active
            rows = []
            for row in worksheet.iter_rows(max_row=20, max_col=8, values_only=True):
                rows.append(" | ".join(_safe_cell_text(value, limit=200) for value in row if value is not None))
            return _safe_cell_text("\n".join(row for row in rows if row), limit=12000)
        except Exception:
            return "[excel preview unavailable]"
    return f"[{file_type or 'binary'} file attached; preview omitted]"


def _write_preview_block(worksheet, start_row: int, value: str, reviewed_value: str) -> int:
    worksheet.merge_cells(start_row=start_row, start_column=1, end_row=start_row + 4, end_column=10)
    cell = worksheet.cell(start_row, 1, value=_safe_cell_text(value, limit=12000))
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    if reviewed_value:
        cell.border = RED_BORDER
    return start_row + 6


def _write_sample_evidence_sheet(
    worksheet,
    sample: dict,
    sample_index: int,
    control: dict,
    evidence_blobs: dict[str, bytes],
) -> None:
    worksheet.sheet_view.showGridLines = False
    for column in range(1, 12):
        worksheet.column_dimensions[openpyxl.utils.get_column_letter(column)].width = 18

    sample_num = sample.get("sample_num") or sample_index
    step_labels = _sample_step_labels(sample)
    _set_title(worksheet, "A1", f"Sample {sample_num} Evidence")
    _set_label(worksheet, "A3", "Application")
    worksheet["B3"] = _safe_cell_text(sample.get("application", ""))
    _set_label(worksheet, "A4", "Item Reference")
    worksheet["B4"] = _safe_cell_text(sample.get("item_reference", ""))
    _set_label(worksheet, "A5", "Steps Tested")
    worksheet["B5"] = ", ".join(sorted(step_labels))

    row = 8
    records = _sample_evidence_records(control, sample)
    if not records:
        worksheet["A8"] = "No mapped evidence was available for this sample."
        return

    for record in records:
        gridfs_id = record.get("gridfs_id", "")
        file_type = record.get("file_type", "")
        content = evidence_blobs.get(gridfs_id, b"")
        reviewed_value = record.get("identified_value") or record.get("comments") or record.get("supporting_value") or ""

        _set_label(worksheet, f"A{row}", _safe_cell_text(record.get("filename", "Evidence")))
        worksheet[f"B{row}"] = _safe_cell_text(file_type)
        _set_label(worksheet, f"A{row + 1}", "Reviewed value")
        worksheet[f"B{row + 1}"] = _safe_cell_text(reviewed_value, limit=12000)
        worksheet[f"B{row + 1}"].border = RED_BORDER
        worksheet[f"B{row + 1}"].alignment = Alignment(wrap_text=True, vertical="top")

        if file_type == "image":
            image_stream = _annotated_image_stream(content, record, step_labels)
            if image_stream:
                image = XLImage(image_stream)
                _resize_for_sheet(image)
                worksheet.add_image(image, f"A{row + 3}")
                row += max(12, int(image.height / 18) + 6)
                continue

        row = _write_preview_block(worksheet, row + 3, _decode_preview(content, file_type), reviewed_value)


def _sample_sheet_title(sample: dict, index: int) -> str:
    value = sample.get("sample_num") or index
    title = f"Sample {value}"
    cleaned = "".join(ch for ch in title if ch not in "[]:*?/\\")
    return cleaned[:31] or f"Sample {index}"


def _write_sample_evidence_sheets(workbook: Workbook, control: dict, evidence_blobs: dict[str, bytes]) -> None:
    for offset, sample in enumerate(control.get("sample_results", []), start=1):
        title = _sample_sheet_title(sample, offset)
        if title in workbook.sheetnames:
            del workbook[title]
        worksheet = workbook.create_sheet(title, index=offset)
        _write_sample_evidence_sheet(worksheet, sample, offset, control, evidence_blobs)


def _ensure_sample_rows(worksheet, start_row: int, sample_count: int) -> None:
    extra_rows = max(0, sample_count - TESTING_GRID_DEFAULT_ROWS)
    if not extra_rows:
        return

    insert_at = start_row + TESTING_GRID_DEFAULT_ROWS
    worksheet.insert_rows(insert_at, extra_rows)
    for row in range(insert_at, insert_at + extra_rows):
        _copy_row_style(worksheet, insert_at - 1, row, 13)


def _ensure_procedure_rows(worksheet, attribute_count: int) -> int:
    header_row = _procedure_header_row(worksheet)
    first_row = header_row + 1
    testing_section_row = _find_row_contains(worksheet, "testing (add or subtract", 1) or (first_row + 11)
    existing_rows = max(0, testing_section_row - first_row - 2)
    extra_rows = max(0, attribute_count - existing_rows)
    if extra_rows:
        insert_at = testing_section_row
        worksheet.insert_rows(insert_at, extra_rows)
        for row in range(insert_at, insert_at + extra_rows):
            _copy_row_style(worksheet, first_row + max(existing_rows - 1, 0), row, worksheet.max_column)
            try:
                worksheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
                worksheet.merge_cells(start_row=row, start_column=3, end_row=row, end_column=15)
            except ValueError:
                pass
    return first_row


def _write_dynamic_test_steps(worksheet, steps: list[dict]) -> None:
    first_row = _ensure_procedure_rows(worksheet, len(steps))
    existing_last_row = (_find_row_contains(worksheet, "testing (add or subtract", 1) or first_row + len(steps) + 2) - 3
    for row in range(first_row, max(existing_last_row, first_row + len(steps) - 1) + 1):
        worksheet[f"A{row}"] = None
        worksheet[f"C{row}"] = None

    for index, step in enumerate(steps, start=1):
        row = first_row + index - 1
        worksheet[f"A{row}"] = _step_label(step, index)
        procedure = step.get("description", "")
        attribute = step.get("test_attribute", "")
        evidence = step.get("evidence_required", "")
        parts = [
            _safe_cell_text(procedure, limit=12000),
            f"Attribute: {_safe_cell_text(attribute, limit=12000)}" if attribute else "",
            f"Evidence: {_safe_cell_text(evidence, limit=12000)}" if evidence else "",
        ]
        worksheet[f"C{row}"] = "\n".join(part for part in parts if part)
        worksheet[f"C{row}"].alignment = Alignment(wrap_text=True, vertical="top")


def _write_testing_grid_header(worksheet, labels: list[str]) -> tuple[int, int]:
    header_row = _testing_header_row(worksheet)
    sample_start_row = header_row + 1
    max_header_col = max(_notes_column(len(labels)), 13)
    for column in range(4, max_header_col + 1):
        worksheet.cell(header_row, column).value = None

    for offset, label in enumerate(labels, start=4):
        worksheet.cell(header_row, offset).value = label
    notes_col = _notes_column(len(labels))
    worksheet.cell(header_row, notes_col).value = "Notes (Define any sample exceptions)"
    return header_row, sample_start_row


def _write_testing_sheet(worksheet, session: dict, control: dict) -> None:
    sign_off = session.get("sign_off", {})
    preparer = sign_off.get("preparer", {}) if isinstance(sign_off, dict) else {}
    conclusions = control.get("conclusions", {})
    sampling = control.get("sampling", {})

    worksheet["C4"] = _safe_cell_text(f"{control.get('control_id', '')} - {control.get('control_name', '')}".strip(" -"))
    worksheet["C5"] = _period_text(session)
    worksheet["C6"] = _safe_cell_text(preparer.get("name", ""))
    worksheet["C7"] = _safe_cell_text(session.get("entity", ""))

    worksheet["C12"] = _safe_cell_text(control.get("control_id", ""))
    worksheet["C13"] = _safe_cell_text(control.get("control_name", ""))
    worksheet["C14"] = _safe_cell_text(_control_description(control), limit=12000)
    worksheet["C15"] = _safe_cell_text(control.get("control_type", ""))
    worksheet["M12"] = _yes_no(control.get("walkthrough_performed", False))
    worksheet["M13"] = _yes_no(_operating_effectiveness_testing(control))
    worksheet["M14"] = _safe_cell_text(conclusions.get("d_and_i") or "")
    worksheet["M15"] = _safe_cell_text(conclusions.get("oe") or "")

    worksheet["C21"] = _testing_method_value(control, "inquiry")
    worksheet["C22"] = _testing_method_value(control, "observation")
    worksheet["C23"] = _testing_method_value(control, "inspection")
    worksheet["C24"] = _testing_method_value(control, "reperformance")
    worksheet["M21"] = _safe_cell_text(sampling.get("population_description", ""), limit=12000)
    worksheet["M22"] = sampling.get("adjusted_population_count") or sampling.get("population_count") or ""
    worksheet["M23"] = sampling.get("selected_size") or len(control.get("sample_results", [])) or ""
    worksheet["M24"] = _safe_cell_text(sampling.get("selection_strategy") or sampling.get("llm_suggested_strategy") or "")
    worksheet["M25"] = _safe_cell_text(sampling.get("additional_context", ""), limit=12000)

    steps = control.get("test_steps", [])
    labels = [_step_label(step, index) for index, step in enumerate(steps, start=1)]
    _write_dynamic_test_steps(worksheet, steps)
    _, sample_start_row = _write_testing_grid_header(worksheet, labels)

    samples = control.get("sample_results", [])
    _ensure_sample_rows(worksheet, sample_start_row, len(samples))
    exceptions_by_sample: dict[int, list[dict]] = {}
    for exception in control.get("exceptions", []):
        sample_num = exception.get("sample_num")
        if sample_num is not None:
            exceptions_by_sample.setdefault(int(sample_num), []).append(exception)

    for offset, sample in enumerate(samples):
        row = sample_start_row + offset
        worksheet[f"A{row}"] = sample.get("sample_num") or offset + 1
        worksheet[f"B{row}"] = _safe_cell_text(sample.get("application", ""))
        worksheet[f"C{row}"] = _safe_cell_text(sample.get("item_reference", ""))
        step_result_by_label = {
            str(step.get("attribute_id") or step.get("label") or "").upper(): step
            for step in sample.get("step_results", [])
        }
        for index, label in enumerate(labels, start=4):
            step_result = step_result_by_label.get(label.upper())
            if step_result:
                worksheet.cell(row, index).value = _tickmark(step_result.get("tickmark", ""))
        notes_col = _notes_column(len(labels))
        worksheet.cell(row, notes_col).value = _sample_notes(sample, exceptions_by_sample)

    conclusion_row = (_conclusions_row(worksheet) or 75) + 3
    worksheet[f"C{conclusion_row}"] = "Yes" if conclusions.get("deficiencies_noted") else "No"
    worksheet[f"C{conclusion_row + 1}"] = _safe_cell_text(conclusions.get("d_and_i") or "")
    worksheet[f"C{conclusion_row + 2}"] = _safe_cell_text(conclusions.get("oe") or "")
    worksheet[f"C{conclusion_row + 3}"] = _conclusion_description(control)


def _override_action(entry: dict) -> str:
    category = entry.get("override_type") or entry.get("scope") or ""
    detail = entry.get("field") or entry.get("evidence_filename") or entry.get("override_value") or ""
    reason = entry.get("reason") or ""
    return _safe_cell_text(
        " | ".join(
            str(part)
            for part in [
                entry.get("control_id", ""),
                category,
                detail,
                reason,
            ]
            if part
        ),
        limit=12000,
    )


def _write_auditor_override(worksheet, session: dict) -> None:
    for index, entry in enumerate(session.get("override_log", []), start=1):
        row = 7 + index
        if row > 12:
            worksheet.insert_rows(row)
            _copy_row_style(worksheet, row - 1, row, 7)
        worksheet[f"D{row}"] = index
        worksheet[f"E{row}"] = _safe_cell_text(
            entry.get("user") or entry.get("updated_by") or entry.get("control_id", "")
        )
        worksheet[f"F{row}"] = _safe_cell_text(entry.get("timestamp") or entry.get("updated_at", ""))
        worksheet[f"G{row}"] = _override_action(entry)
        worksheet[f"G{row}"].alignment = Alignment(wrap_text=True, vertical="top")


def build_control_workbook(session: dict, control: dict, evidence_blobs: dict[str, bytes] | None = None) -> bytes:
    workbook = _load_base_workbook()

    for sheet_name in list(workbook.sheetnames):
        if sheet_name not in {TESTING_SHEET, OVERRIDE_SHEET}:
            del workbook[sheet_name]

    blobs = evidence_blobs or {}

    _write_testing_sheet(_sheet(workbook, TESTING_SHEET), session, control)
    _write_sample_evidence_sheets(workbook, control, blobs)
    _write_auditor_override(_sheet(workbook, OVERRIDE_SHEET), session)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
