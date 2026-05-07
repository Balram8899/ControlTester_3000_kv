from __future__ import annotations

from html import escape
from io import BytesIO
from math import ceil
from typing import Any, Iterator

from utils.services.schemas import ConversionResult, DocumentTag, FileType, StyleProfile


def convert_document(
    file_bytes: bytes,
    filename: str,
    file_id: str,
    tag: DocumentTag,
) -> ConversionResult:
    suffix = _file_suffix(filename)
    try:
        if suffix in {"xlsx", "xls"}:
            markdown = _convert_excel(file_bytes)
            file_type: FileType = "xlsx"
            style_profile = None
        elif suffix == "docx":
            markdown, style_profile = _convert_docx(
                file_bytes,
                extract_style=tag in {"procedure", "policy"},
            )
            file_type = "docx"
        elif suffix == "pdf":
            markdown = _convert_pdf(file_bytes)
            file_type = "pdf"
            style_profile = None
        elif suffix in {"txt", "md"}:
            markdown = file_bytes.decode("utf-8", errors="replace")
            file_type = "txt"
            style_profile = None
        else:
            unsupported = suffix or "unknown"
            return ConversionResult(
                status="failed",
                error=f"Unsupported file type: {unsupported}",
            )
    except Exception as exc:
        return ConversionResult(status="failed", error=str(exc))

    markdown = markdown.strip()
    corrupt = looks_corrupt_markdown(markdown)
    if not markdown or len(markdown) < 100:
        return ConversionResult(
            status="failed",
            error="Conversion produced empty output",
            looks_corrupt=True,
        )

    return ConversionResult(
        status="partial" if corrupt else "success",
        file_id=file_id,
        filename=filename,
        file_type=file_type,
        tag=tag,
        markdown=markdown,
        page_count=_estimate_page_count(markdown),
        looks_corrupt=corrupt,
        style_profile=style_profile,
    )


def looks_corrupt_markdown(markdown: str) -> bool:
    text = markdown.strip()
    if not text or len(text) < 100:
        return True
    sample = text[:2000]
    non_printable = sum(
        1 for char in sample if ord(char) < 32 and char not in "\n\r\t"
    )
    replacement_count = sample.count("\ufffd") + sample.count("ï¿½")
    non_printable_ratio = non_printable / max(len(sample), 1)
    garbage_patterns = (
        "word/styles.xml",
        "PK\x03\x04",
        "%%EOF",
        "xref",
    )
    return (
        any(pattern in sample for pattern in garbage_patterns)
        or replacement_count > 20
        or non_printable_ratio > 0.05
    )


def _file_suffix(filename: str) -> str:
    return filename.lower().rsplit(".", 1)[-1] if "." in filename else ""


def _estimate_page_count(markdown: str) -> int:
    return max(1, ceil(len(markdown) / 3000))


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).replace("\n", " ").strip()


def _row_to_markdown(cells: list[str]) -> str:
    return "| " + " | ".join(escape(cell) for cell in cells) + " |"


def _docx_paragraph_to_markdown(paragraph: Any) -> str:
    text = paragraph.text.strip()
    if not text:
        return ""
    style = (paragraph.style.name or "").lower() if paragraph.style else ""
    if "heading" in style:
        level = next((char for char in style if char.isdigit()), "2")
        return f"{'#' * int(level)} {text}"
    return text


def _docx_table_to_markdown(table: Any) -> str:
    rows = [[_cell_text(cell.text) for cell in row.cells] for row in table.rows]
    rows = [row for row in rows if any(cell for cell in row)]
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    padded_rows = [row + [""] * (width - len(row)) for row in rows]
    table_lines = [
        _row_to_markdown(padded_rows[0]),
        _row_to_markdown(["---"] * width),
        *[_row_to_markdown(row) for row in padded_rows[1:]],
    ]
    return "\n".join(table_lines)


def _iter_docx_body_blocks(document: Any) -> Iterator[Any]:
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def _convert_excel(file_bytes: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    sections: list[str] = []
    for sheet in workbook.worksheets:
        rows = [[_cell_text(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
        rows = [row for row in rows if any(cell for cell in row)]
        if not rows:
            continue
        width = max(len(row) for row in rows)
        padded_rows = [row + [""] * (width - len(row)) for row in rows]
        section = [
            f"## {sheet.title}",
            _row_to_markdown(padded_rows[0]),
            _row_to_markdown(["---"] * width),
        ]
        section.extend(_row_to_markdown(row) for row in padded_rows[1:])
        sections.append("\n".join(section))
    return "\n\n".join(sections).strip()


def _convert_docx(file_bytes: bytes, extract_style: bool) -> tuple[str, StyleProfile | None]:
    from docx import Document

    document = Document(BytesIO(file_bytes))
    blocks: list[str] = []
    for block in _iter_docx_body_blocks(document):
        if hasattr(block, "rows"):
            markdown = _docx_table_to_markdown(block)
        else:
            markdown = _docx_paragraph_to_markdown(block)
        if markdown:
            blocks.append(markdown)
    style_profile = _extract_style_profile(document) if extract_style else None
    return "\n\n".join(blocks).strip(), style_profile


def _extract_style_profile(document: object) -> StyleProfile:
    sections = getattr(document, "sections", [])
    first_section = sections[0] if sections else None
    return StyleProfile(
        heading1_font=_style_font_name(document, "Heading 1"),
        heading2_font=_style_font_name(document, "Heading 2"),
        body_font=_style_font_name(document, "Normal"),
        primary_colour_hex=_style_font_colour(document, "Heading 1"),
        secondary_colour_hex=_style_font_colour(document, "Heading 2"),
        margin_top_cm=_length_to_cm(getattr(first_section, "top_margin", None)),
        margin_bottom_cm=_length_to_cm(getattr(first_section, "bottom_margin", None)),
        margin_left_cm=_length_to_cm(getattr(first_section, "left_margin", None)),
        margin_right_cm=_length_to_cm(getattr(first_section, "right_margin", None)),
        table_border_style="single" if getattr(document, "tables", []) else None,
    )


def _style_font_name(document: object, style_name: str) -> str | None:
    try:
        style = document.styles[style_name]
    except KeyError:
        return None
    return getattr(style.font, "name", None)


def _style_font_colour(document: object, style_name: str) -> str | None:
    try:
        style = document.styles[style_name]
    except KeyError:
        return None
    rgb = getattr(style.font.color, "rgb", None)
    return f"#{rgb}" if rgb else None


def _length_to_cm(length: object | None) -> float | None:
    if length is None:
        return None
    return round(float(length.cm), 2)


def _convert_pdf(file_bytes: bytes) -> str:
    import pdfplumber

    pages: list[str] = []
    with pdfplumber.open(BytesIO(file_bytes)) as pdf:
        for index, page in enumerate(pdf.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                pages.append(f"## Page {index}\n\n{text}")
    return "\n\n".join(pages).strip()
