from __future__ import annotations

from io import BytesIO
from typing import Any

from utils.services.generic_findings import (
    detect_seed_findings,
    detect_document_issue_signals,
    facts_from_sheet_result,
    findings_to_suggestions,
)
from utils.services.llm_orchestrator import call_llm
from utils.services.semantic_roles import raw_record_with_roles
from utils.services.schemas import (
    ColumnClassification,
    CorpusMapContribution,
    DocumentFact,
    ExcelPipelineResult,
    RowGap,
    SheetResult,
    SourceReference,
    Suggestion,
)


GAP_TYPE_MAP = {
    "owner": "ownership_gap",
    "frequency": "frequency_gap",
    "evidence_artifact": "evidence_gap",
}

COLUMN_ROLES = {
    "control_id",
    "risk_id",
    "description",
    "owner",
    "frequency",
    "evidence_artifact",
    "system",
    "status",
    "other",
}

HEADER_HINTS = {
    "control",
    "risk",
    "owner",
    "frequency",
    "evidence",
    "description",
    "status",
    "rating",
    "type",
    "name",
    "id",
    "title",
    "category",
    "likelihood",
    "severity",
    "residual",
    "linked",
    "mitigation",
    "process",
    "area",
}


def process_excel(
    file_bytes: bytes,
    filename: str,
    file_id: str,
    pipeline_id: str,
    budget_remaining: int,
) -> ExcelPipelineResult:
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(file_bytes), read_only=False, data_only=True)
    except Exception as exc:
        return ExcelPipelineResult(
            status="failed",
            error=f"openpyxl failed to open file: {exc}",
        )

    sheets: list[SheetResult] = []
    risk_to_control_map: list[dict] = []
    evidence_to_control_map: list[dict] = []
    suggestions: list[Suggestion] = []
    facts: list[DocumentFact] = []
    total_rows_assessed = 0
    status = "success"
    remaining_budget = budget_remaining

    for worksheet in workbook.worksheets:
        try:
            sheet_result, remaining_budget, corpus_map = _process_sheet(
                worksheet=worksheet,
                filename=filename,
                file_id=file_id,
                pipeline_id=pipeline_id,
                budget_remaining=remaining_budget,
            )
            if sheet_result is None:
                continue
            sheets.append(sheet_result)
            if sheet_result.status == "complete":
                total_rows_assessed += sheet_result.row_count
            if sheet_result.status in {"failed", "skipped_budget"}:
                status = "partial"
            if corpus_map:
                risk_to_control_map.extend(corpus_map.risk_to_control_map)
                evidence_to_control_map.extend(corpus_map.evidence_to_control_map)
            if sheet_result.status == "complete":
                suggestions.extend(
                    _gap_suggestions(sheet_result.gaps, file_id, filename, sheet_result.sheet_name)
                )
                sheet_facts = facts_from_sheet_result(
                    file_id=file_id,
                    filename=filename,
                    sheet=sheet_result,
                    document_role="control_inventory" if corpus_map and corpus_map.risk_to_control_map else "evidence_test_result",
                )
                facts.extend(sheet_facts)
                suggestions.extend(
                    findings_to_suggestions(detect_document_issue_signals(sheet_facts))
                )
                if not (corpus_map and corpus_map.risk_to_control_map):
                    suggestions.extend(findings_to_suggestions(detect_seed_findings(sheet_facts)))
        except Exception as exc:
            status = "partial"
            sheets.append(
                SheetResult(
                    sheet_name=worksheet.title,
                    row_count=0,
                    schema=[],
                    gaps=[],
                    records=[],
                    status="failed",
                    error=str(exc),
                )
            )

    total_gaps_found = sum(len(sheet.gaps) for sheet in sheets)
    return ExcelPipelineResult(
        status=status,
        file_id=file_id,
        sheets=sheets,
        total_rows_assessed=total_rows_assessed,
        total_gaps_found=total_gaps_found,
        corpus_map=CorpusMapContribution(
            risk_to_control_map=risk_to_control_map,
            sop_to_control_map=[],
            evidence_to_control_map=evidence_to_control_map,
        ),
        suggestions=suggestions,
        facts=facts,
    )


def _process_sheet(
    worksheet: Any,
    filename: str,
    file_id: str,
    pipeline_id: str,
    budget_remaining: int,
) -> tuple[SheetResult | None, int, CorpusMapContribution | None]:
    cell_values = _normalised_cell_values(worksheet)
    max_row = worksheet.max_row or 0
    max_col = worksheet.max_column or 0
    header_row = _detect_header_row(cell_values, max_row, max_col)
    headers = _row_text(cell_values, header_row, max_col)
    data_rows = _data_row_numbers(cell_values, header_row + 1, max_row, max_col)
    row_count = len(data_rows)

    if row_count == 0:
        return None, budget_remaining, None
    if row_count < 5:
        return (
            SheetResult(
                sheet_name=worksheet.title,
                row_count=row_count,
                schema=[],
                gaps=[],
                records=[],
                status="too_small",
            ),
            budget_remaining,
            None,
        )
    if budget_remaining <= 0:
        return (
            SheetResult(
                sheet_name=worksheet.title,
                row_count=row_count,
                schema=[],
                gaps=[],
                records=[],
                status="skipped_budget",
            ),
            budget_remaining,
            None,
        )

    llm_result = call_llm(
        prompt=_schema_detection_prompt(headers),
        schema_name="schema_detection",
        response_schema=_schema_response_schema(),
        pipeline_id=pipeline_id,
        budget_remaining=budget_remaining,
    )
    remaining_budget = budget_remaining - 1
    if llm_result.status != "success" or llm_result.output is None:
        return (
            SheetResult(
                sheet_name=worksheet.title,
                row_count=row_count,
                schema=[],
                gaps=[],
                records=[],
                status="failed",
                error=llm_result.error or "schema detection failed",
            ),
            remaining_budget,
            None,
        )

    schema = _column_classifications(headers, llm_result.output)
    records, gaps, corpus_map = _scan_rows(
        cell_values=cell_values,
        row_numbers=data_rows,
        schema=schema,
        filename=filename,
        file_id=file_id,
        sheet_name=worksheet.title,
    )
    return (
        SheetResult(
            sheet_name=worksheet.title,
            row_count=row_count,
            schema=schema,
            gaps=gaps,
            records=records,
        ),
        remaining_budget,
        corpus_map,
    )


def _normalised_cell_values(worksheet: Any) -> dict[tuple[int, int], Any]:
    cell_values: dict[tuple[int, int], Any] = {}
    for row in worksheet.iter_rows():
        for cell in row:
            cell_values[(cell.row, cell.column)] = cell.value

    for merge_range in worksheet.merged_cells.ranges:
        top_left = cell_values.get((merge_range.min_row, merge_range.min_col))
        for row in range(merge_range.min_row, merge_range.max_row + 1):
            for col in range(merge_range.min_col, merge_range.max_col + 1):
                cell_values[(row, col)] = top_left
    return cell_values


def _detect_header_row(
    cell_values: dict[tuple[int, int], Any],
    max_row: int,
    max_col: int,
) -> int:
    best_row = 1
    best_score = float("-inf")
    for row in range(1, min(max_row, 10) + 1):
        values = [_cell_text(cell_values.get((row, col))) for col in range(1, max_col + 1)]
        non_empty = [value for value in values if value]
        if not non_empty:
            continue
        unique_values = {value.casefold() for value in non_empty}
        duplicate_count = len(non_empty) - len(unique_values)
        hint_score = sum(_header_hint_score(value) for value in unique_values)
        score = (hint_score * 10) + (len(unique_values) * 2) + len(non_empty) - (duplicate_count * 4)
        if score > best_score:
            best_score = score
            best_row = row
    return best_row


def _header_hint_score(text: str) -> int:
    normalized = " ".join(
        text.casefold()
        .replace("/", " ")
        .replace("-", " ")
        .replace("_", " ")
        .replace("(", " ")
        .replace(")", " ")
        .split()
    )
    words = set(normalized.split())
    phrase_bonus = 0
    for phrase in ("control id", "risk id", "risk title", "control owner", "linked controls"):
        if phrase in normalized:
            phrase_bonus += 2
    return phrase_bonus + len(words.intersection(HEADER_HINTS))


def _filled_count(
    cell_values: dict[tuple[int, int], Any],
    row: int,
    max_col: int,
) -> int:
    return sum(1 for col in range(1, max_col + 1) if _cell_text(cell_values.get((row, col))))


def _row_text(
    cell_values: dict[tuple[int, int], Any],
    row: int,
    max_col: int,
) -> list[str]:
    headers: list[str] = []
    for col in range(1, max_col + 1):
        header = _cell_text(cell_values.get((row, col)))
        headers.append(header or f"Column {col}")
    return headers


def _data_row_numbers(
    cell_values: dict[tuple[int, int], Any],
    start_row: int,
    max_row: int,
    max_col: int,
) -> list[int]:
    return [
        row
        for row in range(start_row, max_row + 1)
        if any(
            _cell_text(cell_values.get((row, col)))
            for col in range(1, max_col + 1)
        )
    ]


def _schema_detection_prompt(headers: list[str]) -> str:
    headers_text = "\n".join(f"- {header}" for header in headers)
    return f"Classify these Excel column headers for schema_detection:\n{headers_text}"


def _schema_response_schema() -> dict:
    return {
        "type": "object",
        "properties": {
            "columns": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "column_name": {"type": "string"},
                        "classified_as": {"type": "string"},
                        "confidence": {"type": "number"},
                    },
                },
            }
        },
        "required": ["columns"],
    }


def _column_classifications(
    headers: list[str],
    output: dict,
) -> list[ColumnClassification]:
    raw_columns = output.get("columns", [])
    by_name = {
        str(column.get("column_name", "")): column
        for column in raw_columns
        if isinstance(column, dict)
    }
    classifications: list[ColumnClassification] = []
    for index, header in enumerate(headers, start=1):
        raw = by_name.get(header, {})
        role = str(raw.get("classified_as", "other"))
        if role not in COLUMN_ROLES:
            role = "other"
        if role == "other":
            role = _header_role_hint(header)
        classifications.append(
            ColumnClassification(
                column_name=header,
                column_index=int(raw.get("column_index", index)),
                classified_as=role,
                confidence=float(raw.get("confidence", 0.0)),
            )
        )
    return classifications


def _header_role_hint(header: str) -> str:
    normalized = " ".join(
        header.casefold()
        .replace("/", " ")
        .replace("-", " ")
        .replace("_", " ")
        .replace("(", " ")
        .replace(")", " ")
        .split()
    )
    if "control id" in normalized or normalized in {"control", "control reference"}:
        return "control_id"
    if "risk id" in normalized or "risk event id" in normalized:
        return "risk_id"
    if "owner" in normalized:
        return "owner"
    if "frequency" in normalized or "periodicity" in normalized:
        return "frequency"
    if "evidence" in normalized or "artifact" in normalized or "artefact" in normalized:
        return "evidence_artifact"
    if "description" in normalized or "control name" in normalized or "risk title" in normalized:
        return "description"
    if "status" in normalized:
        return "status"
    return "other"


def _scan_rows(
    cell_values: dict[tuple[int, int], Any],
    row_numbers: list[int],
    schema: list[ColumnClassification],
    filename: str,
    file_id: str,
    sheet_name: str,
) -> tuple[list[dict], list[RowGap], CorpusMapContribution]:
    records: list[dict] = []
    gaps: list[RowGap] = []
    risk_to_control_map: list[dict] = []
    evidence_to_control_map: list[dict] = []

    for row_number in row_numbers:
        record: dict[str, str] = {}
        raw_record: dict[str, str] = {}
        for column in schema:
            value = _cell_text(cell_values.get((row_number, column.column_index)))
            raw_record[column.column_name] = value
            if column.classified_as != "other":
                record[column.classified_as] = value
            gap_type = GAP_TYPE_MAP.get(column.classified_as)
            if gap_type and not value:
                gaps.append(
                    RowGap(
                        row_index=row_number,
                        missing_columns=[column.column_name],
                        gap_type=gap_type,
                    )
                )
        record.update(raw_record_with_roles(raw_record))
        record["_source"] = {
            "file_id": file_id,
            "sheet_name": sheet_name,
            "row_index": row_number,
        }
        records.append(record)
        control_id = record.get("control_id", "")
        risk_id = record.get("risk_id", "")
        owner = record.get("owner", "")
        evidence_ref = record.get("evidence_artifact", "")
        if risk_id and control_id:
            risk_to_control_map.append(
                {
                    "risk_id": risk_id,
                    "control_id": control_id,
                    "owner": owner,
                    "description": record.get("description", ""),
                    "frequency": record.get("frequency", ""),
                    "evidence_ref": evidence_ref,
                    "filename": filename,
                    "file_id": file_id,
                    "sheet": sheet_name,
                    "row": row_number,
                }
            )
        if evidence_ref and control_id:
            evidence_to_control_map.append(
                {
                    "evidence_ref": evidence_ref,
                    "control_id": control_id,
                    "filename": filename,
                    "file_id": file_id,
                    "sheet": sheet_name,
                    "row": row_number,
                }
            )

    return (
        records,
        gaps,
        CorpusMapContribution(
            risk_to_control_map=risk_to_control_map,
            sop_to_control_map=[],
            evidence_to_control_map=evidence_to_control_map,
        ),
    )


def _gap_suggestions(
    gaps: list[RowGap],
    file_id: str,
    filename: str,
    sheet_name: str,
) -> list[Suggestion]:
    suggestions: list[Suggestion] = []
    counts_by_type: dict[str, int] = {}
    for gap in gaps:
        if gap.gap_type not in {"ownership_gap", "frequency_gap", "evidence_gap"}:
            continue
        count = counts_by_type.get(gap.gap_type, 0)
        if count >= 12:
            continue
        counts_by_type[gap.gap_type] = count + 1
        missing = ", ".join(gap.missing_columns)
        suggestions.append(
            Suggestion(
                suggestion_id=(
                    f"{file_id}_{_slug(sheet_name)}_{gap.gap_type}_"
                    f"{gap.row_index}_{count + 1}"
                ),
                suggestion_type=gap.gap_type,
                severity="medium",
                title=_gap_title(gap.gap_type),
                detail=(
                    f"Sheet '{sheet_name}' row {gap.row_index} is missing "
                    f"required column value(s): {missing}."
                ),
                source_references=[
                    SourceReference(
                        document_id=file_id,
                        filename=filename,
                        sheet_name=sheet_name,
                        row_index=gap.row_index,
                    )
                ],
            )
        )
    return suggestions


def _gap_title(gap_type: str) -> str:
    return {
        "ownership_gap": "Ownership gap in RCM row",
        "frequency_gap": "Frequency gap in RCM row",
        "evidence_gap": "Evidence gap in RCM row",
    }.get(gap_type, "RCM row gap")


def _slug(value: str) -> str:
    return "".join(char.lower() if char.isalnum() else "_" for char in value).strip("_")


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()
