from __future__ import annotations

import zipfile
from dataclasses import dataclass
from datetime import datetime, timezone
from io import BytesIO
import os
import re
from typing import Any
from xml.sax.saxutils import escape

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.table import Table
from docx.text.paragraph import Paragraph

from utils.services.llm_orchestrator import call_llm, get_batch_size_chars
from utils.services.schemas import OutputGenerationResult, SwimlaneSpec
from utils.sop_processing.case_store import DocumentUpliftCaseStore
from utils.sop_processing.content_sanitizer import sanitize_chunk
from utils.sop_processing.prompts import (
    sop_section_rewrite_prompt,
    swimlane_extraction_prompt,
)
from utils.sop_processing import diagram_renderer


DIAGRAM_RENDERER_READY = True
AUTHOR = "TRACE Document Uplift"
DOCX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
DEFAULT_LANE_PALETTE = ["#1E49E2", "#098E7E", "#7213EA", "#EAAA00", "#00B8F5"]

REWRITE_SCHEMA: dict[str, Any] = {
    "type": "object",
    "properties": {"rewritten_text": {"type": "string"}},
}

SWIMLANE_SCHEMA: dict[str, Any] = {
    "type": "object",
    "properties": {
        "title": {"type": "string"},
        "process_owner": {"type": "string"},
        "document_name": {"type": "string"},
        "lanes": {"type": "array"},
        "steps": {"type": "array"},
    },
}

_store: DocumentUpliftCaseStore | None = None


@dataclass
class PreparedChange:
    suggestion_id: str
    title: str
    comment_title: str
    target_label: str
    target_type: str
    detail: str
    original_text: str
    raw_proposed_text: str
    rewritten_text: str
    source_references: list[dict[str, Any]]
    review_status: str
    placement_text: str = ""
    fallback_placement_text: str = ""
    target_anchor_id: str = ""


def get_store() -> DocumentUpliftCaseStore:
    global _store
    if _store is None:
        _store = DocumentUpliftCaseStore()
    return _store


def generate_outputs(
    case_id: str,
    original_sop_bytes: bytes | None,
    suggestions: list[dict[str, Any]],
    process_steps: list[dict[str, Any]],
    corpus_map: dict[str, Any] | None,
    style_profile: dict[str, Any] | None,
    pipeline_id: str,
    stage1_cost: dict[str, Any] | None,
    source_documents: list[dict[str, Any]] | None = None,
) -> OutputGenerationResult:
    accepted_changes = _prepare_changes(
        suggestions=suggestions,
        process_steps=process_steps,
        source_documents=source_documents,
        style_profile=style_profile,
        pipeline_id=pipeline_id,
    )
    warnings: list[str] = []
    if original_sop_bytes:
        output_mode = "track_changes"
        docx_bytes, uplifted_text, sections_rewritten = _build_track_changes_docx(
            original_sop_bytes,
            accepted_changes,
            style_profile or {},
        )
    else:
        output_mode = "standalone"
        warnings.append(
            "Track-changes output unavailable because the source procedure is not a DOCX; generated a standalone DOCX."
        )
        docx_bytes, uplifted_text, sections_rewritten = _build_standalone_docx(
            accepted_changes,
            style_profile or {},
        )

    case_store = get_store()
    docx_output = case_store.save_output_content(
        case_id,
        {
            "output_id": f"{case_id}-uplifted-docx",
            "output_type": "docx",
            "filename": f"{case_id}-uplifted-sop.docx",
            "content_type": DOCX_CONTENT_TYPE,
            "output_mode": output_mode,
        },
        docx_bytes,
    )
    swimlane_spec = _extract_swimlane_spec(
        uplifted_text=uplifted_text,
        style_profile=style_profile or {},
        pipeline_id=pipeline_id,
        process_steps=process_steps,
        document_name=_primary_procedure_document_name(source_documents or []),
        enabled=bool(original_sop_bytes or accepted_changes or process_steps or corpus_map),
    )
    diagram_png_file_id, diagram_pdf_file_id = _save_diagram_outputs(
        case_id,
        swimlane_spec,
        case_store,
    )
    return OutputGenerationResult(
        status="success",
        case_id=case_id,
        docx_file_id=docx_output["output_id"],
        diagram_png_file_id=diagram_png_file_id,
        diagram_pdf_file_id=diagram_pdf_file_id,
        sections_rewritten=sections_rewritten,
        swimlane_spec=swimlane_spec,
        stage2_cost=None,
        warnings=warnings,
    )


def _save_diagram_outputs(
    case_id: str,
    swimlane_spec: SwimlaneSpec | None,
    case_store: DocumentUpliftCaseStore,
) -> tuple[str | None, str | None]:
    if not swimlane_spec:
        return None, None
    png_bytes, pdf_bytes = diagram_renderer.render(swimlane_spec)
    png_output = case_store.save_output_content(
        case_id,
        {
            "output_id": f"{case_id}-swimlane-png",
            "output_type": "png_diagram",
            "filename": f"{case_id}-swimlane.png",
            "content_type": "image/png",
            "output_mode": "standalone",
        },
        png_bytes,
    )
    pdf_output = case_store.save_output_content(
        case_id,
        {
            "output_id": f"{case_id}-swimlane-pdf",
            "output_type": "pdf_diagram",
            "filename": f"{case_id}-swimlane.pdf",
            "content_type": "application/pdf",
            "output_mode": "standalone",
        },
        pdf_bytes,
    )
    return png_output["output_id"], pdf_output["output_id"]


def _prepare_changes(
    suggestions: list[dict[str, Any]],
    process_steps: list[dict[str, Any]] | None,
    source_documents: list[dict[str, Any]] | None,
    style_profile: dict[str, Any] | None,
    pipeline_id: str,
) -> list[PreparedChange]:
    prepared: list[PreparedChange] = []
    process_steps = process_steps or []
    source_name_by_id = _source_name_by_id(source_documents or [])
    rewrite_limit = _stage2_rewrite_limit()
    rewrite_calls_used = 0

    def use_rewrite_budget() -> bool:
        nonlocal rewrite_calls_used
        if rewrite_limit < 0:
            return True
        if rewrite_calls_used >= rewrite_limit:
            return False
        rewrite_calls_used += 1
        return True

    for suggestion in suggestions:
        if suggestion.get("review_status", "pending") not in {"accepted", "edited"}:
            continue
        edit_targets = [
            dict(item)
            for item in suggestion.get("edit_targets", [])
            if isinstance(item, dict)
        ]
        if edit_targets:
            for target in edit_targets:
                target_status = str(
                    target.get("review_status")
                    or suggestion.get("review_status", "accepted")
                )
                if target_status not in {"accepted", "edited"}:
                    continue
                target_source_references = [
                    dict(item)
                    for item in (
                        target.get("source_references")
                        or suggestion.get("source_references", [])
                    )
                    if isinstance(item, dict)
                ]
                target_source_references = _with_resolved_source_names(
                    target_source_references,
                    source_name_by_id,
                )
                proposed_text = str(
                    target.get("edited_proposed_text")
                    or target.get("proposed_text")
                    or ""
                ).strip()
                if not proposed_text:
                    continue
                target_payload = {
                    **suggestion,
                    **target,
                    "parent_title": str(suggestion.get("title") or ""),
                    "suggestion_id": str(
                        target.get("target_id")
                        or f"{suggestion.get('suggestion_id')}:target"
                    ),
                    "title": str(
                        target.get("title")
                        or _target_type_title(target.get("target_type"))
                        or suggestion.get("title")
                        or "Document uplift suggestion"
                    ),
                    "detail": str(target.get("detail") or suggestion.get("detail") or ""),
                    "source_references": target_source_references,
                    "review_status": target_status,
                }
                prepared.append(
                    _prepared_change_from_payload(
                        suggestion=target_payload,
                        proposed_text=proposed_text,
                        original_text=str(target.get("original_text") or "").strip(),
                        source_references=target_source_references,
                        process_steps=process_steps,
                        style_profile=style_profile or {},
                        pipeline_id=pipeline_id,
                        rewrite_enabled=use_rewrite_budget(),
                    )
                )
            continue
        proposed_text = str(
            suggestion.get("edited_proposed_text")
            or suggestion.get("proposed_text")
            or ""
        ).strip()
        if not proposed_text:
            continue
        original_text = str(suggestion.get("original_text") or "").strip()
        source_references = [
            dict(item)
            for item in suggestion.get("source_references", [])
            if isinstance(item, dict)
        ]
        source_references = _with_resolved_source_names(source_references, source_name_by_id)
        prepared.append(
            _prepared_change_from_payload(
                suggestion=suggestion,
                proposed_text=proposed_text,
                original_text=original_text,
                source_references=source_references,
                process_steps=process_steps,
                style_profile=style_profile or {},
                pipeline_id=pipeline_id,
                rewrite_enabled=use_rewrite_budget(),
            )
        )
    return prepared


def _prepared_change_from_payload(
    suggestion: dict[str, Any],
    proposed_text: str,
    original_text: str,
    source_references: list[dict[str, Any]],
    process_steps: list[dict[str, Any]],
    style_profile: dict[str, Any],
    pipeline_id: str,
    rewrite_enabled: bool = True,
) -> PreparedChange:
    rewritten_text = proposed_text
    if rewrite_enabled:
        rewritten_text = _rewrite_text(
            original_text=original_text,
            suggestion=suggestion,
            proposed_text=proposed_text,
            style_profile=style_profile,
            pipeline_id=pipeline_id,
        )
    rewritten_text = _clean_rewritten_text(rewritten_text or proposed_text)
    target_type = str(suggestion.get("target_type") or "")
    if not original_text:
        rewritten_text = _compact_additive_subprocedure_text(rewritten_text, suggestion)
    if target_type == "role_responsibility":
        rewritten_text = _concise_responsibility_text(rewritten_text or proposed_text)
    parent_title = str(suggestion.get("parent_title") or "").strip()
    target_label = str(suggestion.get("title") or "").strip()
    return PreparedChange(
        suggestion_id=str(suggestion.get("suggestion_id") or ""),
        title=target_label or "Document uplift suggestion",
        comment_title=parent_title or target_label or "Document uplift suggestion",
        target_label=target_label if parent_title and target_label != parent_title else "",
        target_type=target_type,
        detail=str(suggestion.get("detail") or ""),
        original_text=original_text,
        raw_proposed_text=proposed_text,
        rewritten_text=rewritten_text,
        source_references=source_references,
        review_status=str(suggestion.get("review_status", "accepted")),
        placement_text=_placement_text_for_suggestion(
            suggestion=suggestion,
            source_references=source_references,
            process_steps=process_steps,
        ),
        fallback_placement_text=_fallback_placement_text_for_suggestion(
            source_references=source_references,
            process_steps=process_steps,
        ),
        target_anchor_id=str(suggestion.get("target_anchor_id") or "").strip(),
    )


def _target_type_title(target_type: Any) -> str:
    labels = {
        "procedure_step": "Procedure update",
        "role_responsibility": "Responsibility update",
        "raci_matrix": "RACI update",
        "evidence_requirement": "Evidence update",
        "monitoring_reporting": "Monitoring update",
        "document_metadata": "Document metadata update",
    }
    return labels.get(str(target_type or ""), "")


def _stage2_rewrite_limit() -> int:
    raw_value = os.getenv("DOCUMENT_UPLIFT_STAGE2_REWRITE_LIMIT", "50").strip()
    try:
        return int(raw_value)
    except ValueError:
        return 50


def _source_name_by_id(source_documents: list[dict[str, Any]]) -> dict[str, str]:
    names: dict[str, str] = {}
    for document in source_documents:
        if not isinstance(document, dict):
            continue
        filename = str(document.get("filename") or document.get("name") or "").strip()
        if not filename:
            continue
        for key in ("file_id", "document_id", "id", "gridfs_file_id"):
            value = str(document.get(key) or "").strip()
            if value:
                names[value] = filename
    return names


def _primary_procedure_document_name(source_documents: list[dict[str, Any]]) -> str:
    candidates = [document for document in source_documents if isinstance(document, dict)]
    preferred_tags = {"procedure", "process_doc", "policy"}
    for document in candidates:
        if str(document.get("tag") or "").strip().lower() in preferred_tags:
            filename = str(document.get("filename") or document.get("name") or "").strip()
            if filename:
                return filename
    for document in candidates:
        filename = str(document.get("filename") or document.get("name") or "").strip()
        if filename.lower().endswith((".docx", ".doc", ".pdf")):
            return filename
    for document in candidates:
        filename = str(document.get("filename") or document.get("name") or "").strip()
        if filename:
            return filename
    return ""


def _with_resolved_source_names(
    source_references: list[dict[str, Any]],
    source_name_by_id: dict[str, str],
) -> list[dict[str, Any]]:
    resolved: list[dict[str, Any]] = []
    for ref in source_references:
        item = dict(ref)
        if not item.get("filename"):
            for key in ("document_id", "file_id"):
                source_id = str(item.get(key) or "").strip()
                if source_id and source_id in source_name_by_id:
                    item["filename"] = source_name_by_id[source_id]
                    break
        resolved.append(item)
    return resolved


def _rewrite_text(
    original_text: str,
    suggestion: dict[str, Any],
    proposed_text: str,
    style_profile: dict[str, Any],
    pipeline_id: str,
) -> str:
    original_for_prompt = original_text or proposed_text
    prompt = sop_section_rewrite_prompt(
        sanitize_chunk(
            original_for_prompt,
            file_id=str(suggestion.get("suggestion_id") or "stage2"),
            anchor_id=_anchor_id(suggestion),
            max_chars=get_batch_size_chars(),
        ),
        [suggestion],
        _style_notes(style_profile),
    )
    result = call_llm(
        prompt=prompt,
        schema_name="sop_section_rewrite",
        response_schema=REWRITE_SCHEMA,
        pipeline_id=pipeline_id,
        budget_remaining=1,
    )
    if result.status != "success" or not result.output:
        return proposed_text
    rewritten = result.output.get("rewritten_text")
    return str(rewritten).strip() if rewritten else proposed_text


def _placement_text_for_suggestion(
    suggestion: dict[str, Any],
    source_references: list[dict[str, Any]],
    process_steps: list[dict[str, Any]],
) -> str:
    for key in ("target_text", "placement_text"):
        value = str(suggestion.get(key) or "").strip()
        if value:
            return value
    target_anchor_id = str(suggestion.get("target_anchor_id") or "").strip()
    anchor_ids = [target_anchor_id] if target_anchor_id else []
    for anchor_id in anchor_ids:
        for step in process_steps:
            if str(step.get("anchor_id") or "").strip() != anchor_id:
                continue
            placement_text = _process_step_text(step)
            if placement_text:
                return placement_text
    best_process_step = _best_process_step_text_for_suggestion(suggestion, process_steps)
    if best_process_step:
        return best_process_step
    for ref in source_references:
        for key in ("excerpt", "section_heading", "section_id"):
            value = str(ref.get(key) or "").strip()
            if value:
                return value
    return str(suggestion.get("original_text") or "").strip()


def _fallback_placement_text_for_suggestion(
    source_references: list[dict[str, Any]],
    process_steps: list[dict[str, Any]],
) -> str:
    anchor_ids = [
        str(ref.get("anchor_id") or "").strip()
        for ref in source_references
        if str(ref.get("anchor_id") or "").strip()
    ]
    for anchor_id in anchor_ids:
        for step in process_steps:
            if str(step.get("anchor_id") or "").strip() != anchor_id:
                continue
            placement_text = _process_step_text(step)
            if placement_text:
                return placement_text
    return ""


def _best_process_step_text_for_suggestion(
    suggestion: dict[str, Any],
    process_steps: list[dict[str, Any]],
) -> str:
    query_text = " ".join(
        str(suggestion.get(key) or "")
        for key in (
            "title",
            "detail",
            "original_text",
            "edited_proposed_text",
            "proposed_text",
        )
    )
    query_tokens = _match_tokens(query_text)
    if len(query_tokens) < 3:
        return ""
    best_text = ""
    best_score = 0.0
    for step in process_steps:
        step_text = " ".join(
            str(step.get(key) or "")
            for key in (
                "actor",
                "action",
                "text",
                "description",
                "procedure_step",
                "system",
                "evidence",
                "section",
                "section_heading",
                "sop_section",
            )
        )
        step_tokens = _match_tokens(step_text)
        if not step_tokens:
            continue
        overlap = query_tokens.intersection(step_tokens)
        score = len(overlap) / max(1, min(len(query_tokens), len(step_tokens)))
        if len(overlap) >= 2 and score > best_score:
            best_text = _process_step_text(step)
            best_score = score
    return best_text if best_score >= 0.20 else ""


def _match_tokens(text: str) -> set[str]:
    stop_words = {
        "accountable",
        "activity",
        "addition",
        "against",
        "before",
        "closure",
        "control",
        "describe",
        "document",
        "evidence",
        "identified",
        "includes",
        "manager",
        "owner",
        "procedure",
        "process",
        "required",
        "requires",
        "retained",
        "review",
        "reviewing",
        "shall",
        "should",
        "source",
        "step",
        "supporting",
        "team",
        "that",
        "the",
        "this",
        "within",
    }
    return {
        token
        for token in _normalize_match_text(text).split()
        if len(token) > 2 and token not in stop_words
    }


def _process_step_text(step: dict[str, Any]) -> str:
    for key in (
        "original_text",
        "action",
        "text",
        "description",
        "procedure_step",
        "section",
        "section_heading",
        "sop_section",
    ):
        value = str(step.get(key) or "").strip()
        if value:
            return value
    return ""


def _clean_rewritten_text(text: str) -> str:
    cleaned = str(text or "").strip()
    cleaned = cleaned.replace("\r\n", "\n").replace("\r", "\n")
    cleaned = re.sub(r"(?m)^\s*#{1,6}\s*", "", cleaned)
    cleaned = re.sub(r"^\s*Suggested addition:\s*", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\*\*\s*([^*:\n]{1,40})\s*:\s*\*\*", r"\1:", cleaned)
    cleaned = cleaned.replace("**", "")
    cleaned = _naturalize_instructional_addition(cleaned)
    cleaned = _naturalize_label_sections(cleaned)
    cleaned = _format_inline_list(cleaned)
    cleaned = " ".join(cleaned.split())
    cleaned = _strip_inline_numbered_heading(cleaned)
    cleaned = _fix_common_casing(cleaned)
    return cleaned


def _compact_additive_subprocedure_text(text: str, suggestion: dict[str, Any]) -> str:
    cleaned = str(text or "").strip()
    if not _looks_like_embedded_subprocedure(cleaned):
        return cleaned
    focus = _suggestion_focus(suggestion)
    if focus:
        return _ensure_sentence(
            "The accountable owner documents SOP-specific handling for "
            f"{_lower_first_unless_acronym(focus)}, including escalation, remediation, "
            "and retained evidence"
        )
    flattened = _remove_subprocedure_markers(cleaned)
    sentence = _first_sentence_fragment(flattened)
    if not sentence:
        return cleaned
    return _ensure_sentence(sentence[:280].rsplit(" ", 1)[0] if len(sentence) > 280 else sentence)


def _looks_like_embedded_subprocedure(text: str) -> bool:
    cleaned = str(text or "")
    # Multi-level numbered headings: "4.2.1 Heading", "4.2.2 Heading"
    if len(re.findall(r"\b\d+(?:\.\d+){1,}\.?\s+[A-Z]", cleaned)) >= 2:
        return True
    # Single-level numbered headings: "1. Heading", "2. Heading", "3. Heading"
    if len(re.findall(r"(?<!\d)\b\d{1,2}\.\s+[A-Z][a-z]", cleaned)) >= 3:
        return True
    # Multiple capitalised phrase headings followed by body text (generic section pattern)
    section_starts = re.findall(
        r"(?:^|\s)([A-Z][A-Za-z ]{4,40})\s+(?:[A-Z][a-z]|\d)",
        cleaned,
    )
    if len(section_starts) >= 3 and len(cleaned) > 300:
        return True
    return False


def _suggestion_focus(suggestion: dict[str, Any]) -> str:
    title = _sentence_fragment(str(suggestion.get("title") or ""))
    title = re.sub(r"^(?:add|include|insert|reflect)\s+", "", title, flags=re.IGNORECASE)
    title = re.sub(r"^(?:a|an|the)\s+", "", title, flags=re.IGNORECASE)
    if title:
        return title
    detail = _sentence_fragment(str(suggestion.get("detail") or ""))
    return detail[:180].rsplit(" ", 1)[0] if len(detail) > 180 else detail


def _remove_subprocedure_markers(text: str) -> str:
    cleaned = re.sub(r"\b\d+(?:\.\d+){1,}\.?\s+", " ", str(text or ""))
    cleaned = re.sub(r"(?<!\d)\b\d{1,2}\.\s+(?=[A-Z])", " ", cleaned)
    return " ".join(cleaned.split())


def _strip_inline_numbered_heading(text: str) -> str:
    """Remove a copied numbered heading when it was merged into body prose."""

    cleaned = str(text or "").strip()
    colon_heading = re.match(
        r"^\d+(?:\.\d+)*\.?\s+[A-Z][A-Za-z0-9&/(), -]{3,100}:\s+"
        r"(?=[A-Z][A-Za-z/&() -]{2,100}\s+"
        r"(?:must|shall|will|should|is|are|reviews?|records?|performs?|approves?|"
        r"documents?|maintains?|monitors?|tests?|coordinates?|escalates?|verifies?)\b)",
        cleaned,
    )
    if colon_heading:
        return cleaned[colon_heading.end() :].strip()
    match = re.match(
        r"^\d+(?:\.\d+)*\.?\s+[A-Z][A-Za-z0-9&/(), -]{3,90}?\s+"
        r"(?=(The|This|Each|All|Any|If|Where|When|For|In)\b)",
        cleaned,
    )
    if match:
        return cleaned[match.end() :].strip()
    process_heading = re.match(
        r"^\d+(?:\.\d+)*\.?\s+.+?\bProcess\s+(?=[A-Z][A-Za-z/&() -]{2,80}\s+"
        r"(must|shall|will|should|is|are|reviews?|records?|performs?|approves?|documents?|maintains?|monitors?|tests?|coordinates?|escalates?)\b)",
        cleaned,
    )
    if process_heading:
        return cleaned[process_heading.end() :].strip()
    return cleaned


def _fix_common_casing(text: str) -> str:
    return re.sub(r"\bkYC\b", "KYC", str(text or ""))


def _concise_responsibility_text(text: str) -> str:
    cleaned = _clean_rewritten_text(text)
    cleaned = re.split(
        r"\bRetained evidence includes\b|\bRetained Evidence\b",
        cleaned,
        maxsplit=1,
        flags=re.IGNORECASE,
    )[0]
    cleaned = _sentence_fragment(cleaned)
    cleaned = re.sub(
        r"^(?:The\s+)?(?P<role>.+?)\s+is responsible for ensuring that\s+",
        "",
        cleaned,
        flags=re.IGNORECASE,
    )
    cleaned = re.sub(
        r"^(?:The\s+)?(?P<role>.+?)\s+is responsible for\s+",
        "",
        cleaned,
        flags=re.IGNORECASE,
    )
    cleaned = re.sub(r"^(?:to\s+)?ensure(?:s)?\s+that\s+", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"^that\s+", "", cleaned, flags=re.IGNORECASE)
    cleaned = _first_sentence_fragment(cleaned)
    cleaned = _active_responsibility_text(cleaned)
    if not re.match(
        r"^(Owns|Maintains|Reviews|Approves|Coordinates|Performs|Monitors|Tests|Escalates|"
        r"Documents|Records|Screens|Validates|Verifies|Files|Assesses|Updates|Completes)\b",
        cleaned,
        flags=re.IGNORECASE,
    ):
        cleaned = f"Owns {_lower_first(cleaned)}"
    cleaned = _ensure_sentence(cleaned)
    if len(cleaned) <= 220:
        return cleaned
    return _shorten_responsibility_sentence(cleaned, max_length=220)


def _shorten_responsibility_sentence(text: str, max_length: int) -> str:
    cleaned = _ensure_sentence(text)
    if len(cleaned) <= max_length:
        return cleaned
    limit = max_length - 3
    prefix = cleaned[:limit]
    clause_breaks = [
        match.start()
        for match in re.finditer(r"[,;]", prefix)
        if match.start() >= int(max_length * 0.55)
    ]
    if clause_breaks:
        shortened = prefix[: clause_breaks[-1]]
    else:
        shortened = prefix.rsplit(" ", 1)[0]
    return _ensure_sentence(shortened.rstrip(" ,;"))


def _first_sentence_fragment(text: str) -> str:
    cleaned = " ".join(str(text or "").strip().split())
    if not cleaned:
        return ""
    match = re.search(r"(?<=[.!?])\s+(?=[A-Z])", cleaned)
    if match:
        cleaned = cleaned[: match.start()]
    return _sentence_fragment(cleaned)


def _active_responsibility_text(text: str) -> str:
    cleaned = _sentence_fragment(text)
    if not cleaned:
        return ""

    passive = _passive_responsibility_text(cleaned)
    if passive:
        return passive

    received = _received_responsibility_text(cleaned)
    if received:
        return received

    actor_action = _actor_action_responsibility_text(cleaned)
    if actor_action:
        return actor_action

    first, _, remainder = cleaned.partition(" ")
    gerund_replacements = {
        "approving": "Approves",
        "coordinating": "Coordinates",
        "documenting": "Documents",
        "escalating": "Escalates",
        "maintaining": "Maintains",
        "monitoring": "Monitors",
        "performing": "Performs",
        "reviewing": "Reviews",
        "testing": "Tests",
    }
    active_replacements = {
        "approve": "Approves",
        "approves": "Approves",
        "coordinate": "Coordinates",
        "coordinates": "Coordinates",
        "document": "Documents",
        "documents": "Documents",
        "escalate": "Escalates",
        "escalates": "Escalates",
        "maintain": "Maintains",
        "maintains": "Maintains",
        "monitor": "Monitors",
        "monitors": "Monitors",
        "perform": "Performs",
        "performs": "Performs",
        "review": "Reviews",
        "reviews": "Reviews",
        "test": "Tests",
        "tests": "Tests",
        "update": "Updates",
        "updates": "Updates",
    }
    replacement = gerund_replacements.get(first.casefold()) or active_replacements.get(
        first.casefold()
    )
    if replacement:
        return " ".join(part for part in (replacement, remainder) if part)

    cleaned = re.sub(r"\bmust\s+", "", cleaned, flags=re.IGNORECASE)
    if re.search(r"\b(shall|should|will|is|are|may|must)\b", cleaned, flags=re.IGNORECASE):
        return f"Ensures that {_lower_first(cleaned)}"
    return cleaned


def _passive_responsibility_text(text: str) -> str:
    passive_verbs = (
        "reviewed and updated",
        "re-screened",
        "screened",
        "stored",
        "retained",
        "documented",
        "recorded",
        "submitted",
        "reviewed",
        "approved",
        "monitored",
        "tested",
        "reconciled",
        "validated",
        "verified",
        "escalated",
        "assessed",
        "filed",
        "maintained",
        "completed",
        "updated",
    )
    aux = r"(?:must\s+be|shall\s+be|should\s+be|will\s+be|is|are|was|were|be|being|been)"
    pattern = re.compile(
        rf"^(?:all\s+)?(?P<object>.+?)\s+{aux}\s+"
        rf"(?P<verb>{'|'.join(re.escape(verb) for verb in passive_verbs)})\b"
        r"(?P<rest>.*)$",
        flags=re.IGNORECASE,
    )
    match = pattern.match(text)
    if not match:
        return ""
    object_text = _responsibility_object(match.group("object"))
    rest = _responsibility_rest(match.group("rest"))
    verb = match.group("verb").casefold()
    if verb == "reviewed and updated":
        return f"Reviews and updates {object_text}{rest}"
    if verb == "reviewed":
        return f"Reviews {object_text}{rest}"
    if verb == "updated":
        return f"Updates {object_text}{rest}"
    if verb in {"stored", "maintained"}:
        return f"Maintains {object_text}{rest}"
    if verb == "retained":
        return f"Maintains retention of {object_text}{rest}"
    if verb in {"documented", "recorded"}:
        return f"Documents {object_text}{rest}"
    if verb == "submitted":
        return f"Coordinates submission of {object_text}{rest}"
    if verb == "approved":
        return f"Approves {object_text}{rest}"
    if verb == "re-screened":
        return f"Performs re-screening of {object_text}{rest}"
    if verb == "screened":
        return f"Screens {object_text}{rest}"
    if verb == "monitored":
        return f"Monitors {object_text}{rest}"
    if verb == "tested":
        return f"Tests {object_text}{rest}"
    if verb == "reconciled":
        return f"Reconciles {object_text}{rest}"
    if verb == "validated":
        return f"Validates {object_text}{rest}"
    if verb == "verified":
        return f"Verifies {object_text}{rest}"
    if verb == "escalated":
        return f"Coordinates escalation of {object_text}{rest}"
    if verb == "assessed":
        return f"Assesses {object_text}{rest}"
    if verb == "filed":
        return f"Files {object_text}{rest}"
    if verb == "completed":
        return f"Completes {object_text}{rest}"
    return ""


def _received_responsibility_text(text: str) -> str:
    match = re.match(
        r"^(?:all\s+)?(?P<object>.+?)\s+(?:receives?|undergo(?:es)?)\s+(?P<activity>.+)$",
        text,
        flags=re.IGNORECASE,
    )
    if not match:
        return ""
    object_text = _responsibility_object(match.group("object"))
    activity = re.split(
        r"\bincludes\b|\bincluding\b",
        _sentence_fragment(match.group("activity")),
        maxsplit=1,
        flags=re.IGNORECASE,
    )[0].strip(" ,;")
    if not activity:
        return ""
    return f"Performs {activity} for {object_text}"


def _actor_action_responsibility_text(text: str) -> str:
    verb_map = {
        "approve": "Approves",
        "approves": "Approves",
        "assess": "Assesses",
        "assesses": "Assesses",
        "complete": "Completes",
        "completes": "Completes",
        "coordinate": "Coordinates",
        "coordinates": "Coordinates",
        "document": "Documents",
        "documents": "Documents",
        "escalate": "Escalates",
        "escalates": "Escalates",
        "file": "Files",
        "files": "Files",
        "maintain": "Maintains",
        "maintains": "Maintains",
        "monitor": "Monitors",
        "monitors": "Monitors",
        "perform": "Performs",
        "performs": "Performs",
        "record": "Records",
        "records": "Records",
        "review": "Reviews",
        "reviews": "Reviews",
        "screen": "Screens",
        "screens": "Screens",
        "test": "Tests",
        "tests": "Tests",
        "update": "Updates",
        "updates": "Updates",
        "validate": "Validates",
        "validates": "Validates",
        "verify": "Verifies",
        "verifies": "Verifies",
    }
    verb_pattern = "|".join(re.escape(verb) for verb in sorted(verb_map, key=len, reverse=True))
    match = re.match(
        rf"^(?:the\s+)?(?P<actor>[A-Z][A-Za-z0-9&/()., -]{{2,110}}?)\s+"
        rf"(?:(?:must|shall|should|will)\s+)?(?P<verb>{verb_pattern})\b(?P<rest>.*)$",
        text,
        flags=re.IGNORECASE,
    )
    if not match:
        return ""
    rest = _sentence_fragment(match.group("rest"))
    if not rest:
        return ""
    return f"{verb_map[match.group('verb').casefold()]}{_responsibility_rest(rest)}"


def _responsibility_object(text: str) -> str:
    cleaned = _sentence_fragment(text)
    cleaned = re.sub(r"^(?:all|the)\s+", "", cleaned, flags=re.IGNORECASE)
    return _lower_first_unless_acronym(cleaned)


def _responsibility_rest(text: str) -> str:
    cleaned = _sentence_fragment(text)
    return f" {cleaned}" if cleaned else ""


def _lower_first_unless_acronym(text: str) -> str:
    cleaned = _sentence_fragment(text)
    if not cleaned or re.match(r"^[A-Z]{2,}(?:\b|/)", cleaned):
        return cleaned
    return cleaned[:1].lower() + cleaned[1:]


def _naturalize_instructional_addition(text: str) -> str:
    pattern = re.compile(
        r"^Add (?:a|an) (?P<step_label>.+?) owned by (?P<owner>.+?)\. "
        r"The step should describe (?P<body>.+?) and identify retained evidence:\s*(?P<evidence>.+?)\.?$",
        flags=re.IGNORECASE | re.DOTALL,
    )
    match = pattern.match(text.strip())
    if not match:
        return text
    owner = _sentence_fragment(match.group("owner"))
    body = _sentence_fragment(match.group("body"))
    evidence = _sentence_fragment(match.group("evidence"))
    return f"{_owned_activity_sentence(owner, body)} Retained evidence includes {evidence}."


def _owned_activity_sentence(owner: str, body: str) -> str:
    subject = _owner_subject(owner)
    phrase = _third_person_activity_phrase(body)
    if phrase:
        return f"{subject} {phrase}."
    return f"{subject} documents the procedure."


def _owner_subject(owner: str) -> str:
    cleaned = _sentence_fragment(owner) or "accountable owner"
    lower = cleaned.casefold()
    if lower.startswith(("the ", "a ", "an ")):
        return cleaned[0].upper() + cleaned[1:]
    role_words = {
        "advisor",
        "analyst",
        "approver",
        "committee",
        "department",
        "function",
        "group",
        "lead",
        "manager",
        "officer",
        "operator",
        "owner",
        "team",
        "unit",
    }
    if any(word in lower.split() for word in role_words):
        return f"The {cleaned}"
    return f"The {cleaned} team"


def _third_person_activity_phrase(body: str) -> str:
    cleaned = _sentence_fragment(body)
    if not cleaned:
        return ""
    first, _, remainder = cleaned.partition(" ")
    replacements = {
        "approving": "approves",
        "checking": "checks",
        "collecting": "collects",
        "confirming": "confirms",
        "documenting": "documents",
        "escalating": "escalates",
        "identifying": "identifies",
        "maintaining": "maintains",
        "monitoring": "monitors",
        "notifying": "notifies",
        "obtaining": "obtains",
        "performing": "performs",
        "reconciling": "reconciles",
        "recording": "records",
        "retaining": "retains",
        "reviewing": "reviews",
        "screening": "screens",
        "testing": "tests",
        "validating": "validates",
        "verifying": "verifies",
    }
    replacement = replacements.get(first.casefold())
    if replacement:
        return " ".join(part for part in (replacement, remainder) if part)
    if re.search(r"\b(shall|must|should|will|is|are)\b", cleaned, flags=re.IGNORECASE):
        return f"ensures that {_lower_first(cleaned)}"
    return f"performs the activity to {_lower_first(cleaned)}"


def _lower_first(text: str) -> str:
    text = _sentence_fragment(text)
    return text[:1].lower() + text[1:] if text else text


def _sentence_fragment(value: str) -> str:
    return " ".join(str(value or "").strip().rstrip(".").split())


def _naturalize_label_sections(text: str) -> str:
    matches = list(
        re.finditer(
            r"\b(Owner|Procedure Step|Procedure|Description|Retained Evidence)\s*:",
            text,
            flags=re.IGNORECASE,
        )
    )
    if not matches:
        return text
    sections: dict[str, str] = {}
    title = _format_inline_list(text[: matches[0].start()].strip(" :"))
    for index, match in enumerate(matches):
        start = match.end()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        key = match.group(1).casefold().replace(" ", "_")
        sections[key] = _format_inline_list(text[start:end].strip(" ;"))
    heading = title
    if not heading and sections.get("procedure_step") and (
        sections.get("procedure") or sections.get("description")
    ):
        heading = sections["procedure_step"]
    body = (
        sections.get("procedure")
        or sections.get("description")
        or (sections.get("procedure_step") if sections.get("procedure_step") != heading else "")
    )
    if not body:
        return text
    parts: list[str] = []
    if heading:
        parts.append(_ensure_sentence(heading))
    parts.append(_ensure_sentence(body))
    if sections.get("owner"):
        parts.append(f"The accountable owner is {_sentence_fragment(sections['owner'])}.")
    if sections.get("retained_evidence"):
        parts.append(
            f"Retained evidence includes {_format_inline_list(sections['retained_evidence'])}."
        )
    return " ".join(parts)


def _format_inline_list(value: str) -> str:
    cleaned = str(value or "").replace("\n", " ")
    cleaned = re.sub(r"\s*[*•]\s+", "; ", cleaned)
    cleaned = re.sub(r"\s+-\s+", "; ", cleaned)
    cleaned = re.sub(r"\s*;\s*", "; ", cleaned)
    cleaned = re.sub(r"(;\s*){2,}", "; ", cleaned)
    cleaned = re.sub(r"\.{2,}", ".", cleaned)
    return " ".join(cleaned.strip(" ;").split())


def _ensure_sentence(value: str) -> str:
    cleaned = _sentence_fragment(value)
    if not cleaned:
        return ""
    return cleaned if cleaned.endswith((".", "?", "!")) else f"{cleaned}."


def _build_track_changes_docx(
    original_sop_bytes: bytes,
    changes: list[PreparedChange],
    style_profile: dict[str, Any],
) -> tuple[bytes, str, int]:
    changes = _merge_collocated_additions(changes)
    document = Document(BytesIO(original_sop_bytes))
    _apply_style_profile(document, style_profile)
    base_text = "\n".join(paragraph.text for paragraph in _iter_paragraphs(document))
    comments: list[dict[str, str]] = []
    sections_rewritten = 0
    for change in changes:
        if _apply_change_to_document(document, change, len(comments)):
            comments.append(
                {
                    "id": str(len(comments)),
                    "text": _comment_text(change),
                }
            )
            sections_rewritten += 1

    stream = BytesIO()
    document.save(stream)
    content = stream.getvalue()
    if comments:
        content = _add_comments_and_revision_settings(content, comments)
    return content, _apply_change_text_to_base(base_text, changes), sections_rewritten


def _merge_collocated_additions(changes: list[PreparedChange]) -> list[PreparedChange]:
    grouped: list[tuple[tuple[str, str] | None, list[PreparedChange]]] = []
    group_by_key: dict[tuple[str, str], list[PreparedChange]] = {}
    for change in changes:
        key = _collocated_addition_key(change)
        if key is None:
            grouped.append((None, [change]))
            continue
        if key in group_by_key:
            group_by_key[key].append(change)
            continue
        group = [change]
        group_by_key[key] = group
        grouped.append((key, group))
    merged: list[PreparedChange] = []
    for _key, group in grouped:
        merged.append(_merge_change_group(group) if len(group) > 1 else group[0])
    return merged


def _collocated_addition_key(change: PreparedChange) -> tuple[str, str] | None:
    if change.original_text:
        return None
    if change.target_type not in {"procedure_step", "evidence_requirement", "monitoring_reporting"}:
        return None
    if _markdown_table_block(change.raw_proposed_text):
        return None
    anchor_key = (
        change.target_anchor_id
        or _normalize_match_text(change.placement_text)
        or _normalize_match_text(change.fallback_placement_text)
    )
    if not anchor_key:
        return None
    return (change.target_type, anchor_key)


def _merge_change_group(group: list[PreparedChange]) -> PreparedChange:
    first = group[0]
    return PreparedChange(
        suggestion_id=";".join(_unique_texts([change.suggestion_id for change in group])),
        title=_combined_title(group),
        comment_title=_combined_comment_title(group),
        target_label=first.target_label,
        target_type=first.target_type,
        detail=" ".join(_ensure_sentence(text) for text in _unique_texts([change.detail for change in group])),
        original_text="",
        raw_proposed_text="\n".join(_unique_texts([change.raw_proposed_text for change in group])),
        rewritten_text=" ".join(
            _ensure_sentence(text)
            for text in _unique_texts([change.rewritten_text for change in group])
        ),
        source_references=_unique_source_references(
            ref for change in group for ref in change.source_references
        ),
        review_status=first.review_status,
        placement_text=first.placement_text,
        fallback_placement_text=first.fallback_placement_text,
        target_anchor_id=first.target_anchor_id,
    )


def _combined_title(group: list[PreparedChange]) -> str:
    titles = _unique_texts([change.title for change in group])
    if len(titles) == 1:
        return titles[0]
    return f"Combined {group[0].target_type.replace('_', ' ')} updates"


def _combined_comment_title(group: list[PreparedChange]) -> str:
    titles = _unique_texts([change.comment_title for change in group])
    if len(titles) == 1:
        return titles[0]
    return f"Combined {group[0].target_type.replace('_', ' ')} updates"


def _unique_texts(values: list[str]) -> list[str]:
    seen: set[str] = set()
    unique: list[str] = []
    for value in values:
        cleaned = str(value or "").strip()
        if not cleaned:
            continue
        key = _normalize_match_text(cleaned)
        if key in seen:
            continue
        seen.add(key)
        unique.append(cleaned)
    return unique


def _unique_source_references(refs: Any) -> list[dict[str, Any]]:
    seen: set[tuple[tuple[str, str], ...]] = set()
    unique: list[dict[str, Any]] = []
    for ref in refs:
        if not isinstance(ref, dict):
            continue
        key = tuple(sorted((str(k), str(v)) for k, v in ref.items() if v is not None))
        if key in seen:
            continue
        seen.add(key)
        unique.append(dict(ref))
    return unique


def _build_standalone_docx(
    changes: list[PreparedChange],
    style_profile: dict[str, Any],
) -> tuple[bytes, str, int]:
    document = Document()
    _apply_style_profile(document, style_profile)
    document.add_heading("Document Uplift Output", level=1)
    document.add_paragraph(
        "Track-changes output is unavailable because the primary procedure was not supplied as a DOCX file."
    )
    for change in changes:
        document.add_heading(change.title, level=2)
        document.add_paragraph(change.rewritten_text)
        document.add_paragraph(f"Why this was suggested: {change.detail}")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue(), _document_text(document, changes), len(changes)


def _apply_style_profile(
    document: DocxDocument,
    style_profile: dict[str, Any],
) -> None:
    body_font = style_profile.get("body_font")
    if body_font and "Normal" in document.styles:
        document.styles["Normal"].font.name = str(body_font)
    for section in document.sections:
        if style_profile.get("margin_top_cm") is not None:
            section.top_margin = Cm(float(style_profile["margin_top_cm"]))
        if style_profile.get("margin_bottom_cm") is not None:
            section.bottom_margin = Cm(float(style_profile["margin_bottom_cm"]))
        if style_profile.get("margin_left_cm") is not None:
            section.left_margin = Cm(float(style_profile["margin_left_cm"]))
        if style_profile.get("margin_right_cm") is not None:
            section.right_margin = Cm(float(style_profile["margin_right_cm"]))


def _apply_change_to_document(
    document: DocxDocument,
    change: PreparedChange,
    comment_id: int,
) -> bool:
    if not change.original_text:
        if _apply_markdown_table_change(document, change, comment_id):
            return True
        if change.target_type == "role_responsibility" and _apply_role_responsibility_change(
            document,
            change,
            comment_id,
        ):
            return True
        paragraph = _paragraph_for_addition(document, change)
        _clear_paragraph_runs(paragraph)
        _append_revision_pair(
            paragraph,
            deleted_text="",
            inserted_text=change.rewritten_text,
            comment_id=comment_id,
        )
        return True

    for paragraph in _iter_paragraphs(document):
        if not _can_anchor_paragraph(paragraph, change):
            continue
        paragraph_text = paragraph.text
        if change.original_text not in paragraph_text:
            continue
        before, _, after = paragraph_text.partition(change.original_text)
        _clear_paragraph_runs(paragraph)
        if before:
            paragraph.add_run(before)
        _append_revision_pair(
            paragraph,
            deleted_text=change.original_text,
            inserted_text=change.rewritten_text,
            comment_id=comment_id,
        )
        if after:
            paragraph.add_run(after)
        return True

    fuzzy_paragraph = _find_paragraph_by_text(document, change.original_text, change)
    if fuzzy_paragraph:
        paragraph_text = fuzzy_paragraph.text
        _clear_paragraph_runs(fuzzy_paragraph)
        _append_revision_pair(
            fuzzy_paragraph,
            deleted_text=paragraph_text,
            inserted_text=change.rewritten_text,
            comment_id=comment_id,
        )
        return True

    paragraph = _paragraph_for_addition(document, change)
    _append_revision_pair(
        paragraph,
        deleted_text="",
        inserted_text=change.rewritten_text,
        comment_id=comment_id,
    )
    return True


def _apply_markdown_table_change(
    document: DocxDocument,
    change: PreparedChange,
    comment_id: int,
) -> bool:
    table_block = _markdown_table_block(change.raw_proposed_text)
    if not table_block:
        return False
    intro, rows, outro = table_block
    if len(rows) < 2:
        return False
    paragraph = _paragraph_for_addition(document, change)
    _clear_paragraph_runs(paragraph)
    _append_revision_pair(
        paragraph,
        deleted_text="",
        inserted_text=_table_context_text(intro, fallback=change.title),
        comment_id=comment_id,
    )
    table = _insert_table_after_paragraph(paragraph, rows)
    outro_text = _table_context_text(outro, fallback="")
    if outro_text:
        outro_paragraph = _insert_paragraph_after_table(table)
        _append_revision_pair(
            outro_paragraph,
            deleted_text="",
            inserted_text=outro_text,
            comment_id=comment_id,
        )
    return True


def _apply_role_responsibility_change(
    document: DocxDocument,
    change: PreparedChange,
    comment_id: int,
) -> bool:
    target_cell = _find_role_responsibility_cell(document, change)
    if target_cell is not None:
        paragraph = target_cell.paragraphs[-1] if target_cell.paragraphs else target_cell.add_paragraph()
        if paragraph.text and not paragraph.text.endswith((" ", "\n")):
            paragraph.add_run(" ")
        _append_revision_pair(
            paragraph,
            deleted_text="",
            inserted_text=change.rewritten_text,
            comment_id=comment_id,
        )
        return True
    paragraph = _find_role_responsibility_paragraph(document, change)
    if paragraph is None:
        return False
    if paragraph.text and not paragraph.text.endswith((" ", "\n")):
        paragraph.add_run(" ")
    _append_revision_pair(
        paragraph,
        deleted_text="",
        inserted_text=change.rewritten_text,
        comment_id=comment_id,
    )
    return True


def _find_role_responsibility_cell(
    document: DocxDocument,
    change: PreparedChange,
) -> Any | None:
    role_tokens = ("role", "owner", "team", "function", "party", "actor")
    responsibility_tokens = (
        "responsibility",
        "responsibilities",
        "accountability",
        "accountabilities",
        "duty",
        "duties",
    )
    best_cell: Any | None = None
    best_score = 0.0
    proposal_text = f"{change.raw_proposed_text} {change.rewritten_text}"
    normalized_proposal = _normalize_match_text(proposal_text)
    for table in document.tables:
        if not table.rows:
            continue
        headers = [_normalize_match_text(cell.text) for cell in table.rows[0].cells]
        role_col = _semantic_column_index(headers, role_tokens)
        responsibility_col = _semantic_column_index(headers, responsibility_tokens)
        if role_col is None or responsibility_col is None or role_col == responsibility_col:
            continue
        if not _table_has_descriptive_responsibility_column(table, responsibility_col):
            continue
        for row in table.rows[1:]:
            if len(row.cells) <= max(role_col, responsibility_col):
                continue
            role_text = row.cells[role_col].text
            score = _role_match_score(role_text, normalized_proposal)
            if score > best_score:
                best_score = score
                best_cell = row.cells[responsibility_col]
    return best_cell if best_score >= 0.75 else None


def _table_has_descriptive_responsibility_column(table: Any, col_index: int) -> bool:
    word_counts = []
    for row in table.rows[1:]:
        if len(row.cells) <= col_index:
            continue
        text = row.cells[col_index].text.strip()
        if text:
            word_counts.append(len(text.split()))
    if not word_counts:
        return False
    avg_words = sum(word_counts) / len(word_counts)
    return avg_words >= 5


def _find_role_responsibility_paragraph(
    document: DocxDocument,
    change: PreparedChange,
) -> Paragraph | None:
    proposal_text = f"{change.raw_proposed_text} {change.rewritten_text}"
    normalized_proposal = _normalize_match_text(proposal_text)
    best_paragraph: Paragraph | None = None
    best_score = 0.0
    for paragraph in _iter_paragraphs(document):
        if _paragraph_is_in_table(paragraph):
            continue
        paragraph_text = paragraph.text.strip()
        if not paragraph_text or ":" not in paragraph_text:
            continue
        role_text, responsibility_text = paragraph_text.split(":", 1)
        if len(role_text.split()) > 9 or len(responsibility_text.split()) < 3:
            continue
        score = _role_match_score(role_text, normalized_proposal)
        if score > best_score:
            best_score = score
            best_paragraph = paragraph
    return best_paragraph if best_score >= 0.75 else None


def _paragraph_is_in_table(paragraph: Paragraph) -> bool:
    parent = paragraph._p.getparent()
    while parent is not None:
        if parent.tag == qn("w:tbl"):
            return True
        parent = parent.getparent()
    return False


def _semantic_column_index(headers: list[str], tokens: tuple[str, ...]) -> int | None:
    for index, header in enumerate(headers):
        header_tokens = set(header.split())
        if header_tokens.intersection(tokens):
            return index
    return None


def _role_match_score(role_text: str, normalized_proposal: str) -> float:
    normalized_role = _normalize_match_text(role_text)
    if not normalized_role or not normalized_proposal:
        return 0.0
    if normalized_role in normalized_proposal:
        return 1.0
    role_tokens = {
        token
        for token in normalized_role.split()
        if token
        not in {
            "and",
            "the",
            "team",
            "function",
            "owner",
            "group",
            "department",
        }
    }
    if not role_tokens:
        return 0.0
    proposal_tokens = set(normalized_proposal.split())
    return len(role_tokens.intersection(proposal_tokens)) / len(role_tokens)


def _paragraph_for_addition(document: DocxDocument, change: PreparedChange) -> Paragraph:
    anchor_paragraph = None
    for candidate in (
        change.placement_text,
        change.rewritten_text,
        f"{change.title} {change.detail}",
        change.fallback_placement_text,
    ):
        anchor_paragraph = _find_paragraph_by_text(document, candidate, change)
        if anchor_paragraph:
            break
    if not anchor_paragraph:
        return document.add_paragraph()
    anchor_paragraph = _last_body_paragraph_in_section(document, anchor_paragraph)
    paragraph = _insert_paragraph_after(anchor_paragraph)
    paragraph.style = anchor_paragraph.style
    return paragraph


def _is_heading_paragraph(paragraph: Paragraph) -> bool:
    style_name = ""
    try:
        style_name = paragraph.style.name or ""
    except Exception:
        pass
    return style_name.lower().startswith("heading")


def _last_body_paragraph_in_section(
    document: DocxDocument,
    anchor: Paragraph,
) -> Paragraph:
    if not _is_heading_paragraph(anchor):
        return anchor
    all_paragraphs = list(document.paragraphs)
    try:
        start = all_paragraphs.index(anchor)
    except ValueError:
        return anchor
    last_body = anchor
    for paragraph in all_paragraphs[start + 1 :]:
        if _is_heading_paragraph(paragraph):
            break
        if paragraph.text.strip():
            last_body = paragraph
    return last_body


def _find_paragraph_by_text(
    document: DocxDocument,
    target_text: str,
    change: PreparedChange | None = None,
) -> Paragraph | None:
    normalized_target = _normalize_match_text(target_text)
    if len(normalized_target) < 8:
        return None
    for paragraph in _iter_paragraphs(document):
        if change is not None and not _can_anchor_paragraph(paragraph, change):
            continue
        if change is None and _is_role_name_cell_paragraph(paragraph):
            continue
        normalized_paragraph = _normalize_match_text(paragraph.text)
        if not normalized_paragraph:
            continue
        if normalized_target in normalized_paragraph:
            return paragraph
        paragraph_can_anchor = (
            len(normalized_paragraph) >= 24
            and len(normalized_paragraph.split()) >= 4
        )
        if paragraph_can_anchor and normalized_paragraph in normalized_target:
            return paragraph
    target_tokens = set(normalized_target.split())
    if len(target_tokens) < 4:
        return None
    best_paragraph: Paragraph | None = None
    best_score = 0.0
    for paragraph in _iter_paragraphs(document):
        if change is not None and not _can_anchor_paragraph(paragraph, change):
            continue
        if change is None and _is_role_name_cell_paragraph(paragraph):
            continue
        paragraph_tokens = set(_normalize_match_text(paragraph.text).split())
        if len(paragraph_tokens) < 4:
            continue
        overlap = target_tokens.intersection(paragraph_tokens)
        score = len(overlap) / max(1, min(len(target_tokens), len(paragraph_tokens)))
        if score > best_score:
            best_score = score
            best_paragraph = paragraph
    return best_paragraph if best_score >= 0.60 else None


def _can_anchor_paragraph(paragraph: Paragraph, change: PreparedChange) -> bool:
    if _is_role_name_cell_paragraph(paragraph):
        return False
    if _is_reference_catalog_cell_paragraph(paragraph) and not _is_reference_catalog_update(change):
        return False
    return True


def _is_reference_catalog_update(change: PreparedChange) -> bool:
    target_type = str(change.target_type or "").casefold()
    if target_type in {"document_metadata", "reference_catalog", "reference_document"}:
        return True
    combined = _normalize_match_text(f"{change.title} {change.target_label}")
    return bool(
        combined
        and any(token in combined.split() for token in {"reference", "references", "bibliography"})
    )


def _is_reference_catalog_cell_paragraph(paragraph: Paragraph) -> bool:
    table_info = _paragraph_table_position(paragraph)
    if table_info is None:
        return False
    _cell, _cell_index, row_index, rows = table_info
    if not rows or row_index == 0:
        return False
    header_cells = [child for child in rows[0] if child.tag == qn("w:tc")]
    headers = [_normalize_match_text(_xml_text(header)) for header in header_cells]
    if len(headers) < 2:
        return False

    header_tokens = set(" ".join(headers).split())
    process_like_tokens = {
        "step",
        "steps",
        "activity",
        "activities",
        "action",
        "actions",
        "task",
        "tasks",
        "owner",
        "owners",
        "role",
        "roles",
        "responsibility",
        "responsibilities",
        "frequency",
        "evidence",
        "risk",
        "risks",
        "control",
        "controls",
    }
    if header_tokens.intersection(process_like_tokens):
        return False

    reference_col = _semantic_column_index(
        headers,
        (
            "reference",
            "references",
            "ref",
            "document",
            "documents",
            "id",
            "identifier",
            "code",
            "standard",
            "standards",
            "policy",
            "policies",
            "procedure",
            "procedures",
            "regulation",
            "regulations",
            "source",
            "sources",
            "citation",
            "citations",
        ),
    )
    description_col = _semantic_column_index(
        headers,
        (
            "description",
            "descriptions",
            "title",
            "name",
            "details",
            "detail",
            "summary",
            "purpose",
        ),
    )
    return (
        reference_col is not None
        and description_col is not None
        and reference_col != description_col
    )


def _is_role_name_cell_paragraph(paragraph: Paragraph) -> bool:
    table_info = _paragraph_table_position(paragraph)
    if table_info is None:
        return False
    _cell, cell_index, row_index, rows = table_info
    if not rows or row_index == 0:
        return False
    header_cells = [child for child in rows[0] if child.tag == qn("w:tc")]
    headers = [_normalize_match_text(_xml_text(header)) for header in header_cells]
    role_col = _semantic_column_index(headers, ("role", "owner", "team", "function", "party", "actor"))
    responsibility_col = _semantic_column_index(
        headers,
        (
            "responsibility",
            "responsibilities",
            "accountability",
            "accountabilities",
            "duty",
            "duties",
            "activity",
            "activities",
        ),
    )
    return (
        role_col is not None
        and responsibility_col is not None
        and role_col != responsibility_col
        and cell_index == role_col
    )


def _paragraph_table_position(paragraph: Paragraph) -> tuple[Any, int, int, list[Any]] | None:
    cell = paragraph._p.getparent()
    if cell is None or cell.tag != qn("w:tc"):
        return None
    row = cell.getparent()
    if row is None or row.tag != qn("w:tr"):
        return None
    table = row.getparent()
    if table is None or table.tag != qn("w:tbl"):
        return None
    cells = [child for child in row if child.tag == qn("w:tc")]
    rows = [child for child in table if child.tag == qn("w:tr")]
    try:
        return cell, cells.index(cell), rows.index(row), rows
    except ValueError:
        return None


def _xml_text(element: Any) -> str:
    return "".join(node.text or "" for node in element.iter() if node.tag == qn("w:t"))


def _normalize_match_text(text: str) -> str:
    normalized = re.sub(r"[^a-z0-9]+", " ", str(text or "").casefold())
    return " ".join(normalized.split())


def _insert_paragraph_after(paragraph: Paragraph) -> Paragraph:
    new_paragraph = OxmlElement("w:p")
    paragraph._p.addnext(new_paragraph)
    return Paragraph(new_paragraph, paragraph._parent)


def _insert_paragraph_after_table(table: Table) -> Paragraph:
    new_paragraph = OxmlElement("w:p")
    table._tbl.addnext(new_paragraph)
    return Paragraph(new_paragraph, table._parent)


def _insert_table_after_paragraph(paragraph: Paragraph, rows: list[list[str]]) -> Table:
    column_count = max(len(row) for row in rows)
    document = paragraph.part.document
    table = document.add_table(rows=len(rows), cols=column_count)
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    for row_index, row_values in enumerate(rows):
        for col_index in range(column_count):
            table.cell(row_index, col_index).text = (
                row_values[col_index] if col_index < len(row_values) else ""
            )
    paragraph._p.addnext(table._tbl)
    return table


def _markdown_table_block(text: str) -> tuple[str, list[list[str]], str] | None:
    lines = str(text or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")
    table_start = -1
    table_end = -1
    for index, line in enumerate(lines):
        if _looks_like_markdown_table_row(line):
            table_start = index
            table_end = index
            break
    if table_start < 0:
        return None
    for index in range(table_start + 1, len(lines)):
        if not _looks_like_markdown_table_row(lines[index]):
            break
        table_end = index
    raw_rows = [_markdown_cells(line) for line in lines[table_start : table_end + 1]]
    rows = [row for row in raw_rows if row and not _is_markdown_separator_row(row)]
    if len(rows) < 2:
        return None
    intro = "\n".join(lines[:table_start]).strip()
    outro = "\n".join(lines[table_end + 1 :]).strip()
    return intro, rows, outro


def _looks_like_markdown_table_row(line: str) -> bool:
    stripped = str(line or "").strip()
    return stripped.startswith("|") and stripped.endswith("|") and stripped.count("|") >= 2


def _markdown_cells(line: str) -> list[str]:
    return [
        _clean_markdown_cell(cell)
        for cell in str(line or "").strip().strip("|").split("|")
    ]


def _clean_markdown_cell(text: str) -> str:
    cleaned = str(text or "").strip().replace("**", "")
    return re.sub(r"\s+", " ", cleaned)


def _is_markdown_separator_row(cells: list[str]) -> bool:
    populated = [cell.strip() for cell in cells if cell.strip()]
    return bool(populated) and all(
        re.fullmatch(r":?-{3,}:?", cell) for cell in populated
    )


def _table_context_text(text: str, fallback: str) -> str:
    source = str(text or fallback or "").strip()
    if not source:
        return ""
    cleaned = _clean_rewritten_text(source)
    cleaned = re.sub(r"^Add an?\s+", "", cleaned, flags=re.IGNORECASE)
    return cleaned[:1].upper() + cleaned[1:] if cleaned else ""


def _iter_paragraphs(document: DocxDocument) -> list[Paragraph]:
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    return paragraphs


def _clear_paragraph_runs(paragraph: Paragraph) -> None:
    p = paragraph._p
    for child in list(p):
        if child.tag == qn("w:r"):
            p.remove(child)


def _append_revision_pair(
    paragraph: Paragraph,
    deleted_text: str,
    inserted_text: str,
    comment_id: int,
) -> None:
    if deleted_text:
        paragraph._p.append(_revision_element("del", deleted_text, comment_id * 2 + 1))
    paragraph._p.append(_comment_boundary("commentRangeStart", comment_id))
    paragraph._p.append(_revision_element("ins", inserted_text, comment_id * 2 + 2))
    paragraph._p.append(_comment_boundary("commentRangeEnd", comment_id))
    paragraph._p.append(_comment_reference_run(comment_id))


def _revision_element(kind: str, text: str, revision_id: int) -> OxmlElement:
    element = OxmlElement(f"w:{kind}")
    element.set(qn("w:id"), str(revision_id))
    element.set(qn("w:author"), AUTHOR)
    element.set(qn("w:date"), _utc_revision_date())
    run = OxmlElement("w:r")
    text_element = OxmlElement("w:delText" if kind == "del" else "w:t")
    text_element.text = text
    run.append(text_element)
    element.append(run)
    return element


def _comment_boundary(kind: str, comment_id: int) -> OxmlElement:
    element = OxmlElement(f"w:{kind}")
    element.set(qn("w:id"), str(comment_id))
    return element


def _comment_reference_run(comment_id: int) -> OxmlElement:
    run = OxmlElement("w:r")
    reference = OxmlElement("w:commentReference")
    reference.set(qn("w:id"), str(comment_id))
    run.append(reference)
    return run


def _add_comments_and_revision_settings(
    docx_bytes: bytes,
    comments: list[dict[str, str]],
) -> bytes:
    source = BytesIO(docx_bytes)
    files: dict[str, bytes] = {}
    with zipfile.ZipFile(source, "r") as package:
        for name in package.namelist():
            files[name] = package.read(name)

    files["[Content_Types].xml"] = _with_comments_content_type(
        files["[Content_Types].xml"].decode("utf-8")
    ).encode("utf-8")
    files["word/_rels/document.xml.rels"] = _with_comments_relationship(
        files["word/_rels/document.xml.rels"].decode("utf-8")
    ).encode("utf-8")
    if "word/settings.xml" in files:
        files["word/settings.xml"] = _with_track_revisions(
            files["word/settings.xml"].decode("utf-8")
        ).encode("utf-8")
    files["word/comments.xml"] = _comments_xml(comments).encode("utf-8")

    output = BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as package:
        for name, content in files.items():
            package.writestr(name, content)
    return output.getvalue()


def _with_comments_content_type(xml: str) -> str:
    if "/word/comments.xml" in xml:
        return xml
    override = (
        '<Override PartName="/word/comments.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"/>'
    )
    return xml.replace("</Types>", f"{override}</Types>")


def _with_comments_relationship(xml: str) -> str:
    if "relationships/comments" in xml:
        return xml
    relationship = (
        '<Relationship Id="rIdDocumentUpliftComments" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments" '
        'Target="comments.xml"/>'
    )
    return xml.replace("</Relationships>", f"{relationship}</Relationships>")


def _with_track_revisions(xml: str) -> str:
    if "<w:trackRevisions" in xml:
        return xml
    return xml.replace("</w:settings>", "<w:trackRevisions/></w:settings>")


def _comments_xml(comments: list[dict[str, str]]) -> str:
    body = "".join(
        (
            f'<w:comment w:id="{escape(comment["id"])}" '
            f'w:author="{escape(AUTHOR)}" w:date="{escape(_utc_revision_date())}">'
            f"<w:p><w:r><w:t>{escape(comment['text'])}</w:t></w:r></w:p>"
            "</w:comment>"
        )
        for comment in comments
    )
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"{body}</w:comments>"
    )


def _extract_swimlane_spec(
    uplifted_text: str,
    style_profile: dict[str, Any],
    pipeline_id: str,
    process_steps: list[dict[str, Any]],
    document_name: str,
    enabled: bool,
) -> SwimlaneSpec | None:
    if not enabled:
        return None
    palette = _colour_palette(style_profile)
    prompt = swimlane_extraction_prompt(
        sanitize_chunk(
            uplifted_text or "No uplifted SOP text available.",
            file_id=pipeline_id,
            anchor_id="stage2-swimlane",
            max_chars=get_batch_size_chars(),
        ),
        palette,
    )
    result = call_llm(
        prompt=prompt,
        schema_name="swimlane_extraction",
        response_schema=SWIMLANE_SCHEMA,
        pipeline_id=pipeline_id,
        budget_remaining=1,
    )
    if result.status != "success" or not result.output:
        return _fallback_swimlane_spec(process_steps, palette, document_name=document_name)
    try:
        spec = SwimlaneSpec(**result.output)
        if not spec.lanes or not spec.steps:
            return _fallback_swimlane_spec(process_steps, palette, document_name=document_name)
        return _complete_swimlane_metadata(spec, document_name=document_name)
    except ValueError:
        return _fallback_swimlane_spec(process_steps, palette, document_name=document_name)


def _fallback_swimlane_spec(
    process_steps: list[dict[str, Any]],
    palette: list[str],
    *,
    document_name: str = "",
) -> SwimlaneSpec | None:
    usable_steps = [
        step
        for step in process_steps
        if str(step.get("action") or "").strip()
    ]
    if not usable_steps:
        return None

    actors: list[str] = []
    for step in usable_steps:
        actor = str(step.get("actor") or "Unassigned").strip() or "Unassigned"
        if actor not in actors:
            actors.append(actor)
        if len(actors) >= 6:
            break
    lane_by_actor = {
        actor: f"l{index + 1}"
        for index, actor in enumerate(actors)
    }
    lanes = [
        {
            "lane_id": lane_id,
            "label": actor,
            "colour_hex": palette[index % len(palette)],
        }
        for index, (actor, lane_id) in enumerate(lane_by_actor.items())
    ]

    steps: list[dict[str, Any]] = [
        {
            "step_id": "start",
            "lane_id": lanes[0]["lane_id"],
            "label": "Start",
            "step_type": "start",
            "next_steps": ["step-1"],
            "branch_labels": {},
        }
    ]
    for index, step in enumerate(usable_steps, start=1):
        actor = str(step.get("actor") or "Unassigned").strip() or "Unassigned"
        lane_id = lane_by_actor.get(actor, lanes[-1]["lane_id"])
        next_step = f"step-{index + 1}" if index < len(usable_steps) else "end"
        steps.append(
            {
                "step_id": f"step-{index}",
                "lane_id": lane_id,
                "label": _compact_step_label(str(step.get("action") or "")),
                "step_type": "action",
                "next_steps": [next_step],
                "branch_labels": {},
            }
        )
    steps.append(
        {
            "step_id": "end",
            "lane_id": steps[-1]["lane_id"],
            "label": "End",
            "step_type": "end",
            "next_steps": [],
            "branch_labels": {},
        }
    )
    return SwimlaneSpec(
        title="Document uplift process swimlane",
        process_owner=actors[0] if actors else "",
        document_name=document_name,
        lanes=lanes,
        steps=steps,
    )


def _complete_swimlane_metadata(
    spec: SwimlaneSpec,
    *,
    document_name: str,
) -> SwimlaneSpec:
    if not spec.document_name and document_name:
        spec.document_name = document_name
    if not spec.process_owner and spec.lanes:
        spec.process_owner = spec.lanes[0].label
    return spec


def _compact_step_label(text: str) -> str:
    cleaned = " ".join(text.split())
    if len(cleaned) <= 54:
        return cleaned
    return f"{cleaned[:51].rstrip()}..."


def _colour_palette(style_profile: dict[str, Any]) -> list[str]:
    primary = style_profile.get("primary_colour_hex")
    palette = [str(primary)] if primary else []
    palette.extend(colour for colour in DEFAULT_LANE_PALETTE if colour not in palette)
    return palette


def _document_text(document: DocxDocument, changes: list[PreparedChange]) -> str:
    return _apply_change_text_to_base(
        "\n".join(paragraph.text for paragraph in _iter_paragraphs(document)),
        changes,
    )


def _apply_change_text_to_base(base: str, changes: list[PreparedChange]) -> str:
    for change in changes:
        if change.original_text:
            base = base.replace(change.original_text, change.rewritten_text)
        elif change.rewritten_text not in base:
            base = f"{base}\n{change.rewritten_text}"
    return base


def _comment_text(change: PreparedChange) -> str:
    labels: list[str] = []
    for ref in change.source_references:
        label = _source_reference_label(ref)
        if label and label not in labels:
            labels.append(label)
    source_text = "; ".join(labels)
    target_text = f" Target: {change.target_label};" if change.target_label else ""
    return (
        f"{change.comment_title}.{target_text} Why this was suggested: {change.detail}"
        + (f" Source document(s): {source_text}." if source_text else "")
    )


def _source_reference_label(ref: dict[str, Any]) -> str:
    document_name = str(
        ref.get("filename")
        or ref.get("document_name")
        or ref.get("document_id")
        or ref.get("file_id")
        or ""
    ).strip()
    details: list[str] = []
    if ref.get("page") is not None:
        details.append(f"page {ref['page']}")
    if ref.get("sheet_name"):
        details.append(str(ref["sheet_name"]))
    row = ref.get("row_index") if ref.get("row_index") is not None else ref.get("row_number")
    if row is not None:
        details.append(f"row {row}")
    if ref.get("column_name"):
        details.append(f"column {ref['column_name']}")
    if document_name and details:
        return f"{document_name} - {' '.join(details)}"
    return document_name or "source document"


def _style_notes(style_profile: dict[str, Any]) -> str:
    if not style_profile:
        return "Retain the uploaded SOP's formal tone and document formatting."
    return (
        "Retain the uploaded SOP's formal tone and document formatting. "
        f"Body font: {style_profile.get('body_font') or 'source document default'}. "
        f"Primary colour: {style_profile.get('primary_colour_hex') or 'source document default'}."
    )


def _anchor_id(suggestion: dict[str, Any]) -> str:
    if suggestion.get("target_anchor_id"):
        return str(suggestion["target_anchor_id"])
    for ref in suggestion.get("source_references", []):
        if isinstance(ref, dict) and ref.get("anchor_id"):
            return str(ref["anchor_id"])
    return str(suggestion.get("suggestion_id") or "stage2")


def _utc_revision_date() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()
