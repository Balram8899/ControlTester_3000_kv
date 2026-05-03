from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO
from html import escape


@dataclass
class MarkdownConversion:
    markdown: str
    converter: str
    status: str = "converted"
    fallback_used: bool = False
    warnings: list[str] = field(default_factory=list)


def convert_bytes_to_markdown(content: bytes, filename: str, content_type: str = "") -> MarkdownConversion:
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if suffix in {"md", "txt", "csv"}:
        return MarkdownConversion(markdown=content.decode("utf-8", errors="replace"), converter="plain_text")
    if suffix in {"xlsx", "xlsm", "xltx", "xltm"}:
        try:
            return MarkdownConversion(markdown=_xlsx_to_markdown(content), converter="openpyxl")
        except Exception as exc:
            xlsx_warning = f"openpyxl conversion failed: {exc}"
        else:
            xlsx_warning = ""
    else:
        xlsx_warning = ""
    if suffix == "docx":
        try:
            return MarkdownConversion(markdown=_docx_to_markdown(content), converter="python-docx")
        except Exception as exc:
            docx_warning = f"python-docx conversion failed: {exc}"
        else:
            docx_warning = ""
    else:
        docx_warning = ""
    try:
        from markitdown import MarkItDown

        result = MarkItDown().convert_stream(BytesIO(content))
        text = getattr(result, "text_content", "") or str(result)
        warnings = [warning for warning in [xlsx_warning, docx_warning] if warning]
        return MarkdownConversion(markdown=text, converter="markitdown", warnings=warnings)
    except Exception as exc:
        warnings = [warning for warning in [xlsx_warning, docx_warning] if warning]
        warnings.append(f"MarkItDown unavailable or failed: {exc}")
        if _looks_binary(content):
            return MarkdownConversion(
                markdown="",
                converter="binary_rejected",
                status="failed",
                fallback_used=True,
                warnings=[*warnings, "Binary Office content could not be converted; raw bytes were not decoded as text."],
            )
        return MarkdownConversion(
            markdown=content.decode("utf-8", errors="replace"),
            converter="utf8_fallback",
            fallback_used=True,
            warnings=warnings,
        )


def looks_corrupt_markdown(markdown: str) -> bool:
    if not markdown:
        return False
    sample = markdown[:2000]
    replacement_count = sample.count("\ufffd") + sample.count("�")
    control_count = sum(1 for char in sample if ord(char) < 32 and char not in "\n\r\t")
    return sample.startswith("PK") or "word/styles.xml" in sample or replacement_count > 20 or control_count > 20


def _looks_binary(content: bytes) -> bool:
    sample = content[:2048]
    if not sample:
        return False
    return sample.startswith(b"PK") or b"\x00" in sample or sum(1 for byte in sample if byte < 9 or 13 < byte < 32) > 20


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).replace("\n", " ").strip()


def _row_to_markdown(cells: list[str]) -> str:
    return "| " + " | ".join(escape(cell) for cell in cells) + " |"


def _xlsx_to_markdown(content: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sections: list[str] = []
    for sheet in workbook.worksheets:
        rows = [[_cell_text(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
        rows = [row for row in rows if any(cell for cell in row)]
        if not rows:
            continue
        width = max(len(row) for row in rows)
        padded_rows = [row + [""] * (width - len(row)) for row in rows]
        section = [f"## {sheet.title}", _row_to_markdown(padded_rows[0]), _row_to_markdown(["---"] * width)]
        section.extend(_row_to_markdown(row) for row in padded_rows[1:])
        sections.append("\n".join(section))
    return "\n\n".join(sections).strip()


def _docx_to_markdown(content: bytes) -> str:
    from docx import Document

    document = Document(BytesIO(content))
    blocks: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "").lower() if paragraph.style else ""
        if "heading" in style:
            level = next((char for char in style if char.isdigit()), "2")
            blocks.append(f"{'#' * int(level)} {text}")
        else:
            blocks.append(text)
    for table in document.tables:
        rows = [[_cell_text(cell.text) for cell in row.cells] for row in table.rows]
        rows = [row for row in rows if any(cell for cell in row)]
        if not rows:
            continue
        width = max(len(row) for row in rows)
        padded_rows = [row + [""] * (width - len(row)) for row in rows]
        blocks.append("\n".join([_row_to_markdown(padded_rows[0]), _row_to_markdown(["---"] * width), *[_row_to_markdown(row) for row in padded_rows[1:]]]))
    return "\n\n".join(blocks).strip()
