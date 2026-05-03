from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Any

from docx import Document
from docx.shared import RGBColor


BLUE = RGBColor(30, 73, 226)
GREEN = RGBColor(0, 154, 68)
RED = RGBColor(128, 0, 0)
AMBER = RGBColor(181, 111, 0)
GRAY = RGBColor(100, 116, 139)


def _color_for(suggestion: dict[str, Any]) -> RGBColor:
    if suggestion.get("severity") in {"high", "critical"}:
        return RED
    if suggestion.get("type") in {"evidence_gap", "frequency_gap", "ownership_gap"}:
        return AMBER
    if suggestion.get("status") == "edited":
        return GREEN
    return BLUE


def _accepted_language(suggestion: dict[str, Any]) -> str:
    return str(suggestion.get("user_text") or suggestion.get("suggested_text") or "").strip()


def build_revised_sections(sections: list[dict[str, Any]], suggestions: list[dict[str, Any]]) -> list[dict[str, Any]]:
    applicable = [item for item in suggestions if item.get("status") in {"accepted", "edited"} and _accepted_language(item)]
    suggestions_by_anchor: dict[str, list[dict[str, Any]]] = {}
    for suggestion in applicable:
        suggestions_by_anchor.setdefault(str(suggestion.get("anchor_id") or ""), []).append(suggestion)

    revised: list[dict[str, Any]] = []
    for section in sections:
        anchor_id = str(section.get("anchor_id") or "")
        section_suggestions = suggestions_by_anchor.get(anchor_id, [])
        revised_text_parts = [_accepted_language(item) for item in section_suggestions if _accepted_language(item)]
        revised_text = "\n".join(revised_text_parts).strip() or str(section.get("text") or "")
        revised.append(
            {
                **section,
                "original_text": section.get("text", ""),
                "revised_text": revised_text,
                "applied_suggestions": section_suggestions,
                "applied_suggestion_ids": [item.get("suggestion_id") for item in section_suggestions],
            }
        )
    return revised


def generate_docx(
    case: dict[str, Any],
    sections: list[dict[str, Any]],
    suggestions: list[dict[str, Any]],
) -> bytes:
    doc = Document()
    doc.add_heading(case.get("title") or "SOP Uplift Output", level=0)
    meta = doc.add_paragraph()
    meta.add_run("Process: ").bold = True
    meta.add_run(case.get("process_name", ""))
    generated = doc.add_paragraph()
    generated.add_run("Generated: ").bold = True
    generated.add_run(datetime.utcnow().isoformat(timespec="seconds") + "Z")

    applicable = [item for item in suggestions if item.get("status") in {"accepted", "edited"}]
    rejected = [item for item in suggestions if item.get("status") == "rejected"]

    if not sections:
        sections = [{"anchor_id": "", "heading": "Revised SOP Content", "text": ""}]

    revised_sections = build_revised_sections(sections, suggestions)
    for section in revised_sections:
        doc.add_heading(section.get("heading") or section.get("anchor_id") or "SOP Section", level=1)
        section_applicable = section.get("applied_suggestions", [])
        if section.get("original_text") and section_applicable:
            paragraph = doc.add_paragraph()
            label = paragraph.add_run("Original SOP language: ")
            label.bold = True
            original = paragraph.add_run(str(section.get("original_text") or ""))
            original.font.strike = True
            original.font.color.rgb = GRAY
        elif section.get("original_text"):
            doc.add_paragraph(str(section.get("original_text") or ""))
        for suggestion in section_applicable:
            paragraph = doc.add_paragraph()
            label = "User-edited SOP language: " if suggestion.get("status") == "edited" else "Revised SOP language: "
            paragraph.add_run(label).bold = True
            run = paragraph.add_run(_accepted_language(suggestion))
            run.font.color.rgb = _color_for(suggestion)

    doc.add_heading("Appendix: Rejected Suggestions", level=1)
    if not rejected:
        doc.add_paragraph("None")
    for suggestion in rejected:
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(suggestion.get("title") or suggestion.get("suggestion_id") or "Rejected suggestion")
        run.font.color.rgb = GRAY
        detail = suggestion.get("user_text") or suggestion.get("suggested_text") or suggestion.get("summary") or ""
        if detail:
            paragraph.add_run(f": {detail}")

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
