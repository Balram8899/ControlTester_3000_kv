import logging
import os
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional

import pandas as pd
from docx import Document as DocxDocument
from langchain.schema import Document
from langchain_community.document_loaders import CSVLoader, PyPDFLoader, TextLoader

logger = logging.getLogger(__name__)

SUPPORTED_DOCUMENT_EXTENSIONS = {
    ".pdf",
    ".txt",
    ".md",
    ".docx",
    ".doc",
    ".csv",
    ".xlsx",
    ".xls",
    ".png",
    ".jpg",
    ".jpeg",
}


class DocumentLoadError(ValueError):
    """Raised when a document cannot be converted into usable text."""


def supported_extension_tuple() -> tuple[str, ...]:
    return tuple(sorted(SUPPORTED_DOCUMENT_EXTENSIONS))


def _normalize_docs(
    docs: Iterable[Document],
    filename: str,
    extra_metadata: Optional[Dict[str, Any]] = None,
) -> List[Document]:
    normalized: List[Document] = []
    base_metadata = extra_metadata or {}
    for doc in docs:
        doc.metadata = {
            **getattr(doc, "metadata", {}),
            **base_metadata,
            "source": filename,
        }
        normalized.append(doc)
    return normalized


def _has_meaningful_text(docs: Iterable[Document], min_chars: int = 20) -> bool:
    return any(len((doc.page_content or "").strip()) >= min_chars for doc in docs)


def _load_docx_documents(file_path: str, filename: str) -> List[Document]:
    try:
        from langchain_community.document_loaders import Docx2txtLoader

        docs = Docx2txtLoader(file_path).load()
        if _has_meaningful_text(docs):
            return docs
    except Exception as exc:
        logger.info(f"[DOCLOAD] Docx2txtLoader fallback for {filename}: {exc}")

    doc = DocxDocument(file_path)
    blocks: List[str] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            blocks.append(text)

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))

    combined = "\n".join(blocks).strip()
    if not combined:
        raise DocumentLoadError(
            f"{filename}: No readable text could be extracted from the Word document."
        )
    return [Document(page_content=combined, metadata={"loader": "python-docx"})]


def _load_spreadsheet_documents(file_path: str, filename: str) -> List[Document]:
    excel_file = pd.ExcelFile(file_path)
    docs: List[Document] = []

    for sheet_name in excel_file.sheet_names:
        frame = excel_file.parse(sheet_name=sheet_name, dtype=str).fillna("")
        if frame.empty:
            continue

        rows: List[str] = []
        columns = [str(col).strip() for col in frame.columns]
        for _, row in frame.iterrows():
            pairs = []
            for col in columns:
                val = str(row[col]).strip()
                if val:
                    pairs.append(f"{col}: {val}")
            if pairs:
                rows.append(" | ".join(pairs))

        if rows:
            docs.append(
                Document(
                    page_content="\n".join(rows),
                    metadata={"sheet_name": sheet_name, "loader": "pandas"},
                )
            )

    if not docs:
        raise DocumentLoadError(
            f"{filename}: No readable text could be extracted from the spreadsheet."
        )
    return docs


def _ocr_image_to_document(file_path: str, filename: str) -> List[Document]:
    try:
        import pytesseract
        from PIL import Image
    except Exception as exc:
        raise DocumentLoadError(
            f"{filename}: OCR dependencies are unavailable for image ingestion."
        ) from exc

    try:
        text = pytesseract.image_to_string(Image.open(file_path)).strip()
    except Exception as exc:
        raise DocumentLoadError(f"{filename}: OCR failed for the uploaded image.") from exc

    if not text:
        raise DocumentLoadError(
            f"{filename}: No readable text could be extracted from the image."
        )

    return [Document(page_content=text, metadata={"loader": "pytesseract"})]


def _ocr_pdf_to_documents(file_path: str, filename: str) -> List[Document]:
    try:
        import pytesseract
        from pdf2image import convert_from_path
    except Exception as exc:
        raise DocumentLoadError(
            f"{filename}: No readable text could be extracted from the document. "
            "If this is a scanned PDF, OCR support is required before ingestion."
        ) from exc

    try:
        images = convert_from_path(file_path)
    except Exception as exc:
        raise DocumentLoadError(
            f"{filename}: The PDF could not be rendered for OCR. It may be corrupted or missing OCR system dependencies."
        ) from exc

    docs: List[Document] = []
    for page_num, image in enumerate(images, start=1):
        text = pytesseract.image_to_string(image).strip()
        if text:
            docs.append(
                Document(
                    page_content=text,
                    metadata={"page": page_num, "loader": "ocr"},
                )
            )

    if not docs:
        raise DocumentLoadError(
            f"{filename}: No readable text could be extracted from the document. "
            "If this is a scanned PDF, OCR support is required before ingestion."
        )

    return docs


def load_documents(
    file_path: str,
    filename: Optional[str] = None,
    extra_metadata: Optional[Dict[str, Any]] = None,
) -> List[Document]:
    """Load many common business document formats into normalized langchain Documents."""
    actual_filename = filename or os.path.basename(file_path)
    ext = Path(file_path).suffix.lower()

    try:
        if ext == ".pdf":
            docs = PyPDFLoader(file_path).load()
            if not _has_meaningful_text(docs):
                docs = _ocr_pdf_to_documents(file_path, actual_filename)
        elif ext in {".txt", ".md"}:
            docs = TextLoader(file_path, encoding="utf-8").load()
        elif ext in {".docx", ".doc"}:
            docs = _load_docx_documents(file_path, actual_filename)
        elif ext == ".csv":
            docs = CSVLoader(file_path).load()
        elif ext in {".xlsx", ".xls"}:
            docs = _load_spreadsheet_documents(file_path, actual_filename)
        elif ext in {".png", ".jpg", ".jpeg"}:
            docs = _ocr_image_to_document(file_path, actual_filename)
        else:
            raise DocumentLoadError(
                f"{actual_filename}: Unsupported file type '{ext or '[none]'}'."
            )
    except DocumentLoadError:
        raise
    except UnicodeDecodeError as exc:
        raise DocumentLoadError(
            f"{actual_filename}: Text decoding failed for this file."
        ) from exc
    except Exception as exc:
        raise DocumentLoadError(
            f"{actual_filename}: Failed to load the document ({exc})."
        ) from exc

    docs = _normalize_docs(docs, actual_filename, extra_metadata=extra_metadata)
    if not _has_meaningful_text(docs):
        raise DocumentLoadError(
            f"{actual_filename}: No readable text could be extracted from the document."
        )
    return docs
