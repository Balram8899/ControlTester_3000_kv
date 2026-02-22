"""
Streamlined Workpaper Filler
- Fills ICMP sheet with ALL controls
- Creates new Findings sheet with tabular summary

FIXES APPLIED
=============
BUG 6  control.get("control_name") → field never produced by test_script_parser.
        Replaced with _control_display_name() using control_description.
BUG 7  control.get("control_type") → field never produced by test_script_parser.
        Replaced with _control_type_label() inferred from test_objective.
BUG 8  (companion to audit_analyzer fix)
        Observation / Recommendation / Exceptions were written as raw JSON
        blobs because audit_analyzer previously put the raw LLM JSON string
        into those fields.  Now that audit_analyzer parses the Assessment JSON
        and returns clean prose strings, workpaper_filler just writes them
        directly — no additional stripping needed.
        Added _clean_cell_value() as a safety net to strip any residual JSON
        artefacts before writing to a cell.
"""

import json
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from typing import Any, Dict, List
from datetime import datetime
import logging

logger = logging.getLogger(__name__)


# ─────────────────────────────────────────────────────────────────────────────
# HELPERS
# ─────────────────────────────────────────────────────────────────────────────

def _clean_cell_value(value: Any, max_len: int = 500) -> str:
    """
    Safety net: if a value is still a JSON string (shouldn't happen after the
    audit_analyzer fix, but kept as defence in depth), extract readable text
    from it rather than dumping JSON into the cell.
    """
    if value is None:
        return ""
    text = str(value)

    # If it looks like JSON, try to extract readable fields
    stripped = text.strip()
    if stripped.startswith("{"):
        try:
            obj = json.loads(stripped)
            parts: List[str] = []
            # Prefer rationale narrative fields
            rationale = obj.get("assessment_rationale", {})
            for key in ("Why_it_failed", "Gap_analysis", "Evidence_of_compliance",
                        "Effectiveness_assessment", "Impact"):
                val = rationale.get(key, "").strip()
                if val:
                    parts.append(val)
            # Fall back to control statement
            if not parts:
                cs = obj.get("control_statement", "").strip()
                if cs:
                    parts.append(cs)
            if parts:
                text = "  ".join(parts)
        except (json.JSONDecodeError, AttributeError):
            # Not valid JSON — write as-is but truncate
            pass

    return text[:max_len]


def _clean_exception(exc: Any) -> str:
    """Strip JSON artefacts from a single exception string."""
    return _clean_cell_value(exc, max_len=250)


def _control_display_name(control: Dict[str, Any]) -> str:
    """BUG 6 FIX: control_name field doesn't exist; derive from control_description."""
    desc = control.get("control_description", "")
    if desc and desc not in ("Not specified", "N/A"):
        desc = desc.replace("\n", " ").strip()
        return desc[:60] + ("…" if len(desc) > 60 else "")
    return control.get("control_id", "N/A")


def _control_type_label(control: Dict[str, Any]) -> str:
    """BUG 7 FIX: control_type field doesn't exist; infer from test_objective."""
    obj = control.get("test_objective", "").lower()
    if any(k in obj for k in ("prevent", "restrict", "block", "enforce")):
        return "Preventive"
    if any(k in obj for k in ("detect", "identify", "monitor", "review", "verify")):
        return "Detective"
    if any(k in obj for k in ("correct", "remediat", "recover", "restor")):
        return "Corrective"
    return "Manual"


def safe_write_cell(ws, cell_address: str, value: Any):
    """Write to a cell, unmerging any merged range that covers it first."""
    try:
        cell = ws[cell_address]
        for merged_range in list(ws.merged_cells.ranges):
            if cell.coordinate in merged_range:
                ws.unmerge_cells(str(merged_range))
                break
        ws[cell_address] = value
    except Exception as exc:
        logger.error(f"Failed to write cell {cell_address}: {exc}")


# ─────────────────────────────────────────────────────────────────────────────
# STYLE HELPERS
# ─────────────────────────────────────────────────────────────────────────────

def _apply_result_colour(cell, value: str):
    colours = {
        "PASS":        ("C6EFCE", "006100"),
        "FAIL":        ("FFC7CE", "9C0006"),
        "PARTIAL":     ("FFEB9C", "9C6500"),
        "NO_EVIDENCE": ("F2DCDB", "843C0C"),
    }
    bg, fg = colours.get(str(value).upper(), ("FFFFFF", "000000"))
    cell.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
    cell.font = Font(bold=True, color=fg)


def _apply_impact_colour(cell, value: str):
    colours = {
        "HIGH":   ("FFC7CE", "9C0006"),
        "MEDIUM": ("FFEB9C", "9C6500"),
        "LOW":    ("C6EFCE", "006100"),
    }
    bg, fg = colours.get(str(value).upper(), ("FFFFFF", "000000"))
    cell.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
    cell.font = Font(bold=True, color=fg)


def _section_header(ws, row: int, title: str, bg: str):
    cell = ws.cell(row=row, column=2, value=title)
    cell.font = Font(bold=True, size=11)
    cell.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")


_THIN = Side(style="thin")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


# ─────────────────────────────────────────────────────────────────────────────
# PUBLIC API
# ─────────────────────────────────────────────────────────────────────────────

def fill_workpaper_template(
    template_path: str,
    output_path: str,
    session_data: Dict[str, Any],
    analysis_results: List[Dict[str, Any]],
) -> str:
    logger.info(f"Filling workpaper with {len(analysis_results)} control results")

    wb = openpyxl.load_workbook(template_path)
    controls: List[Dict]     = session_data.get("controls", [])
    result_map: Dict[str, Dict] = {r.get("control_id"): r for r in analysis_results}

    fill_icmp_sheet_all_controls(wb, controls, result_map, session_data)
    create_findings_sheet(wb, controls, result_map, session_data)

    wb.save(output_path)
    logger.info(f"Workpaper saved: {output_path}")
    return output_path


# ─────────────────────────────────────────────────────────────────────────────
# ICMP SHEET
# ─────────────────────────────────────────────────────────────────────────────

def fill_icmp_sheet_all_controls(
    wb,
    controls: List[Dict],
    result_map: Dict[str, Dict],
    session_data: Dict,
):
    if "ICMP" not in wb.sheetnames:
        wb.create_sheet("ICMP")
    ws = wb["ICMP"]

    # Unmerge everything, then clear rows ≥ 6
    for mr in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(mr))
    for row in ws.iter_rows(min_row=6):
        for cell in row:
            try:
                cell.value = None
            except AttributeError:
                pass

    # ── Header banner ─────────────────────────────────────────────────────────
    safe_write_cell(ws, "B1", "ICMP Controls – Comprehensive Testing Results")
    ws["B1"].font = Font(bold=True, size=14)
    safe_write_cell(ws, "B2", f"Testing Date: {datetime.now().strftime('%Y-%m-%d')}")
    safe_write_cell(ws, "B3", f"Total Controls Tested: {len(controls)}")

    passed  = sum(1 for r in result_map.values() if r.get("result") == "PASS")
    failed  = sum(1 for r in result_map.values() if r.get("result") == "FAIL")
    partial = sum(1 for r in result_map.values() if r.get("result") == "PARTIAL")
    rate    = f"{passed / len(controls) * 100:.1f}%" if controls else "N/A"
    safe_write_cell(ws, "B4",
        f"Results: {passed} PASS  |  {failed} FAIL  |  {partial} PARTIAL  |  Pass Rate: {rate}")
    ws["B4"].font = Font(bold=True)

    cur = 7   # current row cursor

    for idx, control in enumerate(controls, 1):
        cid    = control.get("control_id", f"CTRL-{idx}")
        result = result_map.get(cid, {})

        # ── Control header bar ────────────────────────────────────────────────
        hdr = ws.cell(row=cur, column=2, value=f"CONTROL #{idx}: {cid}")
        hdr.font = Font(bold=True, size=12, color="FFFFFF")
        hdr.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        hdr.alignment = Alignment(horizontal="left", vertical="center")
        ws.merge_cells(f"B{cur}:I{cur}")
        ws.row_dimensions[cur].height = 25
        cur += 1

        # ── Control Information ───────────────────────────────────────────────
        _section_header(ws, cur, "CONTROL INFORMATION", "D9E1F2")
        ws.merge_cells(f"B{cur}:I{cur}")
        cur += 1

        info_rows = [
            ("Control ID:",       cid),
            ("Control Name:",     _control_display_name(control)),      # BUG 6 fix
            ("Control Type:",     _control_type_label(control)),         # BUG 7 fix
            ("Frequency:",        control.get("frequency", "N/A")),
            ("Control Owner:",    control.get("control_owner", "N/A")),
            ("Risk Statement:",   control.get("risk_statement", "N/A")[:300]),
            ("Description:",      control.get("control_description", "N/A")[:300]),
            ("Test Objective:",   control.get("test_objective", "N/A")[:300]),
        ]
        for label, value in info_rows:
            lc = ws.cell(row=cur, column=2, value=label)
            lc.font = Font(bold=True)
            lc.alignment = Alignment(horizontal="right", vertical="top")
            vc = ws.cell(row=cur, column=3, value=str(value))
            vc.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            ws.merge_cells(f"C{cur}:I{cur}")
            ws.row_dimensions[cur].height = (
                40 if label in ("Description:", "Test Objective:", "Risk Statement:") else 20
            )
            cur += 1

        # ── Test Procedures ───────────────────────────────────────────────────
        cur += 1
        _section_header(ws, cur, "TEST PROCEDURES", "D9E1F2")
        ws.merge_cells(f"B{cur}:I{cur}")
        cur += 1

        for step_num, step in enumerate(control.get("test_steps", "").split("\n")[:10], 1):
            if step.strip():
                ws.cell(row=cur, column=2, value=f"{step_num}.").font = Font(bold=True)
                sc = ws.cell(row=cur, column=3, value=step.strip()[:200])
                sc.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                ws.merge_cells(f"C{cur}:I{cur}")
                ws.row_dimensions[cur].height = 28
                cur += 1

        # ── Sample & Evidence ─────────────────────────────────────────────────
        cur += 1
        ws.cell(row=cur, column=2, value="Sample Size:").font = Font(bold=True)
        ws.cell(row=cur, column=3, value=control.get("sample_size", "N/A"))
        cur += 1

        ws.cell(row=cur, column=2, value="Evidence Required:").font = Font(bold=True)
        er = ws.cell(row=cur, column=3, value=control.get("evidence_required", "N/A")[:300])
        er.alignment = Alignment(wrap_text=True)
        ws.merge_cells(f"C{cur}:I{cur}")
        ws.row_dimensions[cur].height = 30
        cur += 1

        uploaded = session_data.get("uploaded_files", {})
        ws.cell(row=cur, column=2, value="Evidence Files:").font = Font(bold=True)
        ev_files = ws.cell(row=cur, column=3,
            value=", ".join(list(uploaded.keys())[:5]) if uploaded else "None")
        ev_files.alignment = Alignment(wrap_text=True)
        ws.merge_cells(f"C{cur}:I{cur}")
        ws.row_dimensions[cur].height = 22
        cur += 1

        # ── Test Results ──────────────────────────────────────────────────────
        cur += 1
        rh = ws.cell(row=cur, column=2, value="TEST RESULTS")
        rh.font = Font(bold=True, size=11, color="FFFFFF")
        rh.fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
        ws.merge_cells(f"B{cur}:I{cur}")
        cur += 1

        # Result
        ws.cell(row=cur, column=2, value="Result:").font = Font(bold=True)
        result_val  = result.get("result", "NOT_TESTED")
        result_cell = ws.cell(row=cur, column=3, value=result_val)
        _apply_result_colour(result_cell, result_val)
        ws.row_dimensions[cur].height = 22
        cur += 1

        # Observation — BUG 8 companion: now a clean prose string, not JSON
        ws.cell(row=cur, column=2, value="Observation:").font = Font(bold=True)
        obs_text = _clean_cell_value(result.get("observation", "No observation recorded"), 1500)
        obs_cell = ws.cell(row=cur, column=3, value=obs_text)
        obs_cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws.merge_cells(f"C{cur}:I{cur}")
        # Taller row to show multi-line observation
        line_count = max(obs_text.count("\n") + 1, 1)
        ws.row_dimensions[cur].height = max(35, min(line_count * 15, 120))
        cur += 1

        # Exceptions — BUG 8 companion: now genuine gap text, not JSON fragments
        exceptions = result.get("exceptions", [])
        if exceptions:
            ws.cell(row=cur, column=2, value="Exceptions:").font = Font(bold=True, color="C00000")
            ws.cell(row=cur, column=3, value=f"{len(exceptions)} exception(s)").font = Font(bold=True, color="C00000")
            cur += 1
            for n, exc in enumerate(exceptions[:5], 1):
                exc_text = _clean_exception(exc)
                if not exc_text:
                    continue
                ws.cell(row=cur, column=2, value=f"{n}.").alignment = Alignment(horizontal="right")
                ec = ws.cell(row=cur, column=3, value=exc_text)
                ec.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                ws.merge_cells(f"C{cur}:I{cur}")
                ws.row_dimensions[cur].height = 28
                cur += 1
        else:
            ws.cell(row=cur, column=2, value="Exceptions:").font = Font(bold=True)
            ws.cell(row=cur, column=3, value="None identified.")
            cur += 1

        # Recommendation — BUG 8 companion: clean prose string
        cur += 1
        ws.cell(row=cur, column=2, value="Recommendation:").font = Font(bold=True)
        rec_text = _clean_cell_value(result.get("recommendation", "No recommendation"), 600)
        rc = ws.cell(row=cur, column=3, value=rec_text)
        rc.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws.merge_cells(f"C{cur}:I{cur}")
        ws.row_dimensions[cur].height = max(28, min(rec_text.count(";") * 18 + 20, 80))
        cur += 1

        # Impact — always populated after BUG 4 fix
        ws.cell(row=cur, column=2, value="Impact:").font = Font(bold=True)
        impact_val  = result.get("impact", "MEDIUM")
        impact_cell = ws.cell(row=cur, column=3, value=impact_val)
        _apply_impact_colour(impact_cell, impact_val)
        ws.row_dimensions[cur].height = 22
        cur += 1

        # Spacer between controls
        cur += 2

    # Column widths
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 75
    for col in ["D", "E", "F", "G", "H", "I"]:
        ws.column_dimensions[col].width = 12

    logger.info(f"ICMP sheet filled with {len(controls)} controls")


# ─────────────────────────────────────────────────────────────────────────────
# FINDINGS SHEET
# ─────────────────────────────────────────────────────────────────────────────

def create_findings_sheet(
    wb,
    controls: List[Dict],
    result_map: Dict[str, Dict],
    session_data: Dict,
):
    if "Findings" in wb.sheetnames:
        del wb["Findings"]
    ws = wb.create_sheet("Findings", 0)

    # ── Header ────────────────────────────────────────────────────────────────
    ws["A1"] = "AUDIT FINDINGS SUMMARY"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:L1")
    ws["A1"].alignment = Alignment(horizontal="center")

    ws["A2"] = f"Testing Date: {datetime.now().strftime('%Y-%m-%d')}"
    ws.merge_cells("A2:L2")
    ws["A3"] = f"Total Controls Tested: {len(controls)}"
    ws.merge_cells("A3:L3")

    passed  = sum(1 for r in result_map.values() if r.get("result") == "PASS")
    failed  = sum(1 for r in result_map.values() if r.get("result") == "FAIL")
    partial = sum(1 for r in result_map.values() if r.get("result") == "PARTIAL")
    rate    = f"{passed / len(controls) * 100:.1f}%" if controls else "N/A"
    ws["A4"] = f"Results: {passed} PASS  |  {failed} FAIL  |  {partial} PARTIAL  |  Pass Rate: {rate}"
    ws["A4"].font = Font(bold=True)
    ws.merge_cells("A4:L4")

    # ── Column headers ────────────────────────────────────────────────────────
    headers = ["#", "Control ID", "Control Name / Description", "Type",
               "Frequency", "Result", "Sample Size", "Exceptions",
               "Impact", "Observation", "Recommendation", "Evidence Files"]

    for col_i, hdr in enumerate(headers, 1):
        cell = ws.cell(row=6, column=col_i, value=hdr)
        cell.font      = Font(bold=True, color="FFFFFF", size=11)
        cell.fill      = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border    = _BORDER

    ws.freeze_panes = "A7"

    # ── Data rows ─────────────────────────────────────────────────────────────
    for row_i, control in enumerate(controls, 1):
        cid    = control.get("control_id", f"CTRL-{row_i}")
        result = result_map.get(cid, {})

        # BUG 8 companion: observation & recommendation are now clean prose
        obs_text = _clean_cell_value(result.get("observation", ""), 200)
        rec_text = _clean_cell_value(result.get("recommendation", ""), 200)

        row_data = [
            row_i,
            cid,
            _control_display_name(control),    # BUG 6 fix
            _control_type_label(control),       # BUG 7 fix
            control.get("frequency", "Ongoing"),
            result.get("result", "NOT_TESTED"),
            control.get("sample_size", "N/A"),
            len(result.get("exceptions", [])),
            result.get("impact", "MEDIUM"),     # BUG 4 fix: always set
            obs_text,
            rec_text,
            ", ".join(list(session_data.get("uploaded_files", {}).keys())[:3]),
        ]

        excel_row = 6 + row_i   # header at row 6, data from row 7
        for col_i, value in enumerate(row_data, 1):
            cell = ws.cell(row=excel_row, column=col_i, value=value)
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            cell.border = _BORDER
            if col_i == 6:    # Result
                _apply_result_colour(cell, value)
            if col_i == 9:    # Impact
                _apply_impact_colour(cell, value)

        ws.row_dimensions[excel_row].height = 45

    # ── Column widths ─────────────────────────────────────────────────────────
    widths = {"A": 5, "B": 13, "C": 35, "D": 13, "E": 12,
              "F": 12, "G": 14, "H": 10, "I": 10,
              "J": 40, "K": 40, "L": 28}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    logger.info(f"Findings sheet created with {len(controls)} rows")


# Backward-compatibility alias
fill_workpaper_template_all_controls = fill_workpaper_template