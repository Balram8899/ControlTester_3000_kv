from __future__ import annotations

import io
import json
from unittest.mock import MagicMock, patch

import mongomock
import openpyxl
import pytest


@pytest.fixture()
def mock_db():
    client = mongomock.MongoClient()
    return client["trace_db"]


def test_control_testing_prompt_uses_canonical_schema():
    from utils.control_assurance.prompts.control_testing import build_control_testing_prompt

    session = {
        "title": "Q1 ITGC",
        "entity": "Tech",
        "testing_period": {"from": "2026-01-01", "to": "2026-03-31"},
        "framework": "SOX",
    }
    control = {
        "control_id": "ITGC-001",
        "control_name": "Password Policy",
        "control_type": "Preventive",
        "domain": "Access Mgmt",
        "frequency": "Continuous",
        "inherent_risk_rating": "High",
        "walkthrough_performed": False,
        "test_steps": [{"label": "A", "description": "Inspect policy", "evidence_required": "Screenshot"}],
        "sampling": {"mode": "sample", "selected_size": 25, "selection_strategy": "random"},
        "evidence_files": [
            {
                "filename": "ad_policy.png",
                "file_type": "image",
                "mapped_step_labels": ["A"],
                "identified_value": "Min length: 12",
            }
        ],
    }

    prompt = build_control_testing_prompt(session, control)

    assert isinstance(prompt, str)
    assert "sample_results" in prompt
    assert "sample_num" in prompt
    assert "item_reference" in prompt
    assert '"ref": "X1"' in prompt
    assert "ITGC-001" in prompt


def test_workpaper_narrative_prompt_is_string():
    from utils.control_assurance.prompts.workpaper_narrative import build_workpaper_narrative_prompt

    session = {
        "title": "Q1 ITGC",
        "entity": "Tech",
        "testing_period": {"from": "2026-01-01", "to": "2026-03-31"},
        "framework": "SOX",
    }
    control = {
        "control_id": "ITGC-001",
        "control_name": "Password Policy",
        "test_steps": [{"label": "A", "description": "Inspect policy"}],
        "sample_results": [{"sample_num": 1, "step_results": [{"label": "A", "tickmark": "PASS"}]}],
        "exceptions": [],
        "conclusions": {"d_and_i": "Effective", "oe": "Effective"},
    }

    prompt = build_workpaper_narrative_prompt(session, control)

    assert isinstance(prompt, str)
    assert "testing_summary" in prompt
    assert "d_and_i_statement" in prompt


def test_issue_drafting_prompt_uses_spec_issue_fields():
    from utils.control_assurance.prompts.issue_drafting import build_issue_drafting_prompt

    session = {
        "title": "Q1 ITGC",
        "entity": "Tech",
        "testing_period": {"from": "2026-01-01", "to": "2026-03-31"},
        "framework": "SOX",
    }
    control = {
        "control_id": "ITGC-001",
        "control_name": "Password Policy",
        "control_type": "Preventive",
        "domain": "Access Mgmt",
    }
    exceptions = [
        {
            "ref": "X1",
            "sample_num": 2,
            "description": "Password min length is 8, expected at least 12.",
            "root_cause": "Configuration baseline not enforced.",
            "auditor_disposition": "Issue drafted",
        }
    ]

    prompt = build_issue_drafting_prompt(session, control, exceptions)

    assert isinstance(prompt, str)
    assert "title" in prompt
    assert "summary" in prompt
    assert "detail" in prompt
    assert "root_cause" in prompt
    assert "recommendation" in prompt
    assert "X1" in prompt


def _insert_testing_control(mock_db, session_id: str = "sess-s4-001", control_id: str = "ctrl-s4-001"):
    if not mock_db.ct_sessions.find_one({"_id": session_id}):
        mock_db.ct_sessions.insert_one(
            {
                "_id": session_id,
                "stage": "testing",
                "title": "Q1 ITGC",
                "entity": "Tech",
                "testing_period": {"from": "2026-01-01", "to": "2026-03-31"},
                "framework": "SOX",
                "stage_checkpoint": {"stage": "mapping", "step": "confirmed"},
            }
        )
    mock_db.ct_controls.insert_one(
        {
            "_id": control_id,
            "session_id": session_id,
            "control_id": "ITGC-001",
            "control_name": "Password Policy",
            "control_type": "Preventive",
            "domain": "Access Mgmt",
            "frequency": "Continuous",
            "inherent_risk_rating": "High",
            "walkthrough_performed": False,
            "test_steps": [{"label": "A", "description": "Inspect policy", "evidence_required": "Screenshot"}],
            "sampling": {
                "mode": "sample",
                "selected_size": 5,
                "selection_strategy": "random",
                "selected_items": [
                    {"sample_num": 1, "application": "AD", "item_reference": "POL-001"},
                    {"sample_num": 2, "application": "AD", "item_reference": "POL-002"},
                    {"sample_num": 3, "application": "AD", "item_reference": "POL-003"},
                    {"sample_num": 4, "application": "AD", "item_reference": "POL-004"},
                    {"sample_num": 5, "application": "AD", "item_reference": "POL-005"},
                ],
            },
            "evidence_files": [
                {
                    "filename": "policy.png",
                    "file_type": "image",
                    "mapped_step_labels": ["A"],
                    "identified_value": "Min length: 12",
                    "gridfs_id": "g001",
                }
            ],
            "sample_results": [],
            "exceptions": [],
            "conclusions": {"d_and_i": None, "oe": None, "issues_log_refs": []},
            "testing_methods": {
                "inquiry": False,
                "observation": False,
                "inspection": False,
                "reperformance": False,
            },
            "status": "pending",
        }
    )
    return session_id, control_id


def _sample_result(sample_num: int, tickmark: str = "PASS") -> dict:
    return {
        "sample_num": sample_num,
        "application": "AD",
        "item_reference": f"POL-{sample_num:03d}",
        "step_results": [{"label": "A", "tickmark": tickmark, "notes": "Reviewed policy evidence."}],
    }


def test_run_testing_writes_canonical_results_and_ct_issue(mock_db):
    from utils.control_assurance.pipeline.stage4_testing import _run_testing

    session_id, control_id = _insert_testing_control(mock_db)
    fake_testing_response = json.dumps(
        {
            "todi_results": {
                "design": {"conclusion": "Effective", "rationale": "Control objective is clear."},
                "implementation": {"conclusion": "Effective", "rationale": "Policy is implemented."},
            },
            "sample_results": [
                _sample_result(1),
                _sample_result(2, "X1"),
                _sample_result(3),
                _sample_result(4),
                _sample_result(5),
            ],
            "exceptions": [
                {
                    "ref": "X1",
                    "sample_num": 2,
                    "description": "Password minimum length is 8, expected at least 12.",
                    "root_cause": "Password baseline was not enforced.",
                    "auditor_disposition": "Issue drafted",
                    "issues_log_ref": None,
                }
            ],
            "conclusions": {
                "d_and_i": "Effective",
                "oe": "Effective",
                "deficiencies_noted": True,
                "rationale": "One isolated exception noted.",
                "issues_log_refs": [],
            },
            "testing_methods": {
                "inquiry": False,
                "observation": False,
                "inspection": True,
                "reperformance": False,
            },
        }
    )
    fake_issue_response = json.dumps(
        {
            "issues": [
                {
                    "title": "Password baseline not enforced",
                    "severity": "High",
                    "summary": "One password policy sample did not meet the expected baseline.",
                    "detail": "Sample 2 showed a minimum password length of 8 rather than 12.",
                    "root_cause": "Password baseline was not enforced.",
                    "recommendation": "Update and monitor the password policy baseline.",
                    "issues_log_ref": "CTI-001",
                    "exception_refs": ["X1"],
                }
            ]
        }
    )
    mock_llm = MagicMock()
    mock_llm.invoke.side_effect = [
        MagicMock(content=fake_testing_response),
        MagicMock(content=fake_issue_response),
    ]

    with patch("utils.control_assurance.pipeline.stage4_testing._get_db", return_value=mock_db):
        with patch("utils.control_assurance.pipeline.stage4_testing.get_llm", return_value=mock_llm):
            _run_testing(session_id)

    control = mock_db.ct_controls.find_one({"_id": control_id})
    assert control["status"] == "complete"
    assert control["todi_results"]["design"]["conclusion"] == "Effective"
    assert control["sample_results"][1]["sample_num"] == 2
    assert control["sample_results"][1]["item_reference"] == "POL-002"
    assert control["exceptions"][0]["ref"] == "X1"
    assert control["exceptions"][0]["issues_log_ref"] == "CTI-001"
    assert control["conclusions"]["issues_log_refs"] == ["CTI-001"]

    issues = list(mock_db.ct_issues.find({"session_id": session_id, "control_id": control_id}))
    assert len(issues) == 1
    assert issues[0]["title"] == "Password baseline not enforced"
    assert issues[0]["issues_log_ref"] == "CTI-001"
    assert issues[0]["exception_refs"] == ["X1"]
    assert issues[0]["pushed_to_issues"] is False
    assert issues[0]["issues_module_id"] is None


def test_run_testing_enforces_oe_ineffective_when_exception_rate_exceeds_threshold(mock_db):
    from utils.control_assurance.pipeline.stage4_testing import _run_testing

    session_id, control_id = _insert_testing_control(mock_db, "sess-s4-threshold", "ctrl-s4-threshold")
    fake_testing_response = json.dumps(
        {
            "todi_results": {
                "design": {"conclusion": "Effective", "rationale": "Design is adequate."},
                "implementation": {"conclusion": "Effective", "rationale": "Implemented for most samples."},
            },
            "sample_results": [
                _sample_result(1, "X1"),
                _sample_result(2, "X2"),
                _sample_result(3),
                _sample_result(4),
                _sample_result(5),
            ],
            "exceptions": [
                {
                    "ref": "X1",
                    "sample_num": 1,
                    "description": "Sample 1 failed.",
                    "root_cause": "Baseline not monitored.",
                    "auditor_disposition": "Issue drafted",
                    "issues_log_ref": None,
                },
                {
                    "ref": "X2",
                    "sample_num": 2,
                    "description": "Sample 2 failed.",
                    "root_cause": "Baseline not monitored.",
                    "auditor_disposition": "Issue drafted",
                    "issues_log_ref": None,
                },
            ],
            "conclusions": {
                "d_and_i": "Effective",
                "oe": "Effective",
                "deficiencies_noted": False,
                "rationale": "LLM did not apply the threshold.",
                "issues_log_refs": [],
            },
            "testing_methods": {
                "inquiry": False,
                "observation": False,
                "inspection": True,
                "reperformance": False,
            },
        }
    )
    fake_issue_response = json.dumps(
        {
            "issues": [
                {
                    "title": "Password failures exceed tolerance",
                    "severity": "High",
                    "summary": "Two of five samples failed.",
                    "detail": "The exception rate exceeded the SOX ITGC tolerance.",
                    "root_cause": "Baseline not monitored.",
                    "recommendation": "Implement monitoring and remediation.",
                    "issues_log_ref": "CTI-001",
                    "exception_refs": ["X1", "X2"],
                }
            ]
        }
    )
    mock_llm = MagicMock()
    mock_llm.invoke.side_effect = [
        MagicMock(content=fake_testing_response),
        MagicMock(content=fake_issue_response),
    ]

    with patch("utils.control_assurance.pipeline.stage4_testing._get_db", return_value=mock_db):
        with patch("utils.control_assurance.pipeline.stage4_testing.get_llm", return_value=mock_llm):
            _run_testing(session_id)

    control = mock_db.ct_controls.find_one({"_id": control_id})
    assert control["conclusions"]["oe"] == "Ineffective"
    assert control["conclusions"]["deficiencies_noted"] is True
    assert "20%" in control["conclusions"]["rationale"]


def test_w_tickmark_absent_from_sample_results_rows(mock_db):
    from utils.control_assurance.pipeline.stage4_testing import _run_testing

    session_id, control_id = _insert_testing_control(mock_db, "sess-s4-w", "ctrl-s4-w")
    fake_testing_response = json.dumps(
        {
            "todi_results": {
                "design": {"conclusion": "Effective", "rationale": "Walkthrough evidence reviewed."},
                "implementation": {"conclusion": "Effective", "rationale": "Implementation confirmed."},
            },
            "sample_results": [_sample_result(1, "W")],
            "exceptions": [],
            "conclusions": {
                "d_and_i": "Effective",
                "oe": "Effective",
                "deficiencies_noted": False,
                "rationale": "No exceptions noted.",
                "issues_log_refs": [],
            },
            "testing_methods": {
                "inquiry": False,
                "observation": False,
                "inspection": True,
                "reperformance": False,
            },
        }
    )
    mock_llm = MagicMock()
    mock_llm.invoke.return_value = MagicMock(content=fake_testing_response)

    with patch("utils.control_assurance.pipeline.stage4_testing._get_db", return_value=mock_db):
        with patch("utils.control_assurance.pipeline.stage4_testing.get_llm", return_value=mock_llm):
            _run_testing(session_id)

    control = mock_db.ct_controls.find_one({"_id": control_id})
    tickmarks = [
        step["tickmark"]
        for sample in control["sample_results"]
        for step in sample["step_results"]
    ]
    assert "W" not in tickmarks
    assert tickmarks == ["PASS"]


def test_exception_refs_consistent_between_sample_results_and_exceptions(mock_db):
    from utils.control_assurance.pipeline.stage4_testing import _run_testing

    session_id, control_id = _insert_testing_control(mock_db, "sess-s4-xrefs", "ctrl-s4-xrefs")
    fake_testing_response = json.dumps(
        {
            "todi_results": {
                "design": {"conclusion": "Effective", "rationale": "Design is adequate."},
                "implementation": {"conclusion": "Effective", "rationale": "Implemented."},
            },
            "sample_results": [_sample_result(1, "X7")],
            "exceptions": [],
            "conclusions": {
                "d_and_i": "Effective",
                "oe": "Effective",
                "deficiencies_noted": True,
                "rationale": "Exception noted.",
                "issues_log_refs": [],
            },
            "testing_methods": {
                "inquiry": False,
                "observation": False,
                "inspection": True,
                "reperformance": False,
            },
        }
    )
    fake_issue_response = json.dumps(
        {
            "issues": [
                {
                    "title": "Generated exception",
                    "severity": "Medium",
                    "summary": "Generated exception.",
                    "detail": "Generated from tickmark.",
                    "root_cause": "Root cause pending auditor review.",
                    "recommendation": "Review generated exception.",
                    "issues_log_ref": "CTI-001",
                    "exception_refs": ["X7"],
                }
            ]
        }
    )
    mock_llm = MagicMock()
    mock_llm.invoke.side_effect = [
        MagicMock(content=fake_testing_response),
        MagicMock(content=fake_issue_response),
    ]

    with patch("utils.control_assurance.pipeline.stage4_testing._get_db", return_value=mock_db):
        with patch("utils.control_assurance.pipeline.stage4_testing.get_llm", return_value=mock_llm):
            _run_testing(session_id)

    control = mock_db.ct_controls.find_one({"_id": control_id})
    x_marks = {
        step["tickmark"]
        for sample in control["sample_results"]
        for step in sample["step_results"]
        if step["tickmark"].startswith("X")
    }
    exception_refs = {exception["ref"] for exception in control["exceptions"]}
    assert x_marks.issubset(exception_refs)


def test_oe_not_forced_at_exactly_20_percent(mock_db):
    from utils.control_assurance.pipeline.stage4_testing import _run_testing

    session_id, control_id = _insert_testing_control(mock_db, "sess-s4-boundary", "ctrl-s4-boundary")
    samples = [
        _sample_result(sample_num, f"X{sample_num}" if sample_num <= 4 else "PASS")
        for sample_num in range(1, 21)
    ]
    exceptions = [
        {
            "ref": f"X{sample_num}",
            "sample_num": sample_num,
            "description": f"Sample {sample_num} failed.",
            "root_cause": "Boundary test root cause.",
            "auditor_disposition": "Issue drafted",
            "issues_log_ref": None,
        }
        for sample_num in range(1, 5)
    ]
    fake_testing_response = json.dumps(
        {
            "todi_results": {
                "design": {"conclusion": "Effective", "rationale": "Design is adequate."},
                "implementation": {"conclusion": "Effective", "rationale": "Implemented."},
            },
            "sample_results": samples,
            "exceptions": exceptions,
            "conclusions": {
                "d_and_i": "Effective",
                "oe": "Effective",
                "deficiencies_noted": True,
                "rationale": "Exactly 20 percent exceptions.",
                "issues_log_refs": [],
            },
            "testing_methods": {
                "inquiry": False,
                "observation": False,
                "inspection": True,
                "reperformance": False,
            },
        }
    )
    fake_issue_response = json.dumps(
        {
            "issues": [
                {
                    "title": "Boundary issue",
                    "severity": "Medium",
                    "summary": "Four exceptions noted.",
                    "detail": "Exactly 20 percent exception rate.",
                    "root_cause": "Boundary test root cause.",
                    "recommendation": "Monitor remediation.",
                    "issues_log_ref": "CTI-001",
                    "exception_refs": ["X1", "X2", "X3", "X4"],
                }
            ]
        }
    )
    mock_llm = MagicMock()
    mock_llm.invoke.side_effect = [
        MagicMock(content=fake_testing_response),
        MagicMock(content=fake_issue_response),
    ]

    with patch("utils.control_assurance.pipeline.stage4_testing._get_db", return_value=mock_db):
        with patch("utils.control_assurance.pipeline.stage4_testing.get_llm", return_value=mock_llm):
            _run_testing(session_id)

    control = mock_db.ct_controls.find_one({"_id": control_id})
    assert control["conclusions"]["oe"] == "Effective"
    assert "exceeds the 20%" not in control["conclusions"]["rationale"]


def test_stage4_writes_checkpoint_before_each_control(mock_db):
    from utils.control_assurance.pipeline.stage4_testing import _run_testing

    session_id = "sess-s4-checkpoint"
    for index in range(3):
        _insert_testing_control(mock_db, session_id, f"ctrl-s4-checkpoint-{index}")
    fake_testing_response = json.dumps(
        {
            "todi_results": {
                "design": {"conclusion": "Effective", "rationale": "Design is adequate."},
                "implementation": {"conclusion": "Effective", "rationale": "Implemented."},
            },
            "sample_results": [_sample_result(1)],
            "exceptions": [],
            "conclusions": {
                "d_and_i": "Effective",
                "oe": "Effective",
                "deficiencies_noted": False,
                "rationale": "No exceptions.",
                "issues_log_refs": [],
            },
            "testing_methods": {
                "inquiry": False,
                "observation": False,
                "inspection": True,
                "reperformance": False,
            },
        }
    )
    mock_llm = MagicMock()
    mock_llm.invoke.side_effect = [MagicMock(content=fake_testing_response) for _ in range(3)]
    original_update_one = mock_db.ct_sessions.update_one
    checkpoint_writes = []

    def recording_update_one(*args, **kwargs):
        update = args[1] if len(args) > 1 else kwargs.get("update", {})
        set_payload = update.get("$set", {})
        if "stage_checkpoint" in set_payload:
            checkpoint_writes.append(set_payload["stage_checkpoint"])
        return original_update_one(*args, **kwargs)

    with patch("utils.control_assurance.pipeline.stage4_testing._get_db", return_value=mock_db):
        with patch("utils.control_assurance.pipeline.stage4_testing.get_llm", return_value=mock_llm):
            with patch.object(mock_db.ct_sessions, "update_one", side_effect=recording_update_one):
                _run_testing(session_id)

    per_control = [entry for entry in checkpoint_writes if entry.get("step", "").startswith("processing_")]
    assert len(per_control) == 3
    assert checkpoint_writes[-1]["step"] == "complete"


def test_run_testing_dispatches_stage5_and_records_task_id(mock_db):
    from utils.control_assurance.pipeline.stage4_testing import run_testing

    session_id = "sess-s4-dispatch"
    mock_db.ct_sessions.insert_one({"_id": session_id, "stage": "testing"})
    fake_task = MagicMock(id="stage5-task-001")

    with patch("utils.control_assurance.pipeline.stage4_testing._get_db", return_value=mock_db):
        with patch("utils.control_assurance.pipeline.stage4_testing._run_testing"):
            with patch(
                "utils.control_assurance.pipeline.stage5_workbook.generate_workbooks.apply_async",
                return_value=fake_task,
            ):
                result = run_testing(session_id)

    assert result["next_task"] == "stage5-task-001"
    assert mock_db.ct_sessions.find_one({"_id": session_id})["celery_task_id"] == "stage5-task-001"


def test_run_testing_failure_writes_error_checkpoint(mock_db):
    from utils.control_assurance.pipeline.stage4_testing import run_testing

    session_id = "sess-s4-failed"
    mock_db.ct_sessions.insert_one({"_id": session_id, "stage": "testing"})

    with patch("utils.control_assurance.pipeline.stage4_testing._get_db", return_value=mock_db):
        with patch("utils.control_assurance.pipeline.stage4_testing._run_testing", side_effect=RuntimeError("LLM down")):
            with pytest.raises(RuntimeError):
                run_testing(session_id)

    session = mock_db.ct_sessions.find_one({"_id": session_id})
    assert session["stage"] == "failed"
    assert session["stage_checkpoint"]["error"] == "LLM down"


def _completed_control(session_id: str = "sess-s5-001", control_id: str = "ctrl-s5-001") -> dict:
    return {
        "_id": control_id,
        "session_id": session_id,
        "control_id": "ITGC-001",
        "control_name": "Password Policy",
        "control_type": "Preventive",
        "domain": "Access Mgmt",
        "test_steps": [{"label": "A", "description": "Inspect policy", "evidence_required": "Screenshot"}],
        "evidence_files": [
            {
                "filename": "policy.txt",
                "file_type": "txt",
                "mapped_step_labels": ["A"],
                "identified_value": "Minimum length: 12",
                "gridfs_id": "evidence-gridfs-001",
            }
        ],
        "todi_results": {
            "design": {"conclusion": "Effective", "rationale": "Design aligns to policy."},
            "implementation": {"conclusion": "Effective", "rationale": "Policy is implemented."},
        },
        "sample_results": [
            _sample_result(1),
            _sample_result(2, "X1"),
            _sample_result(3),
        ],
        "exceptions": [
            {
                "ref": "X1",
                "sample_num": 2,
                "description": "Password minimum length was below baseline.",
                "root_cause": "Password baseline was not enforced.",
                "auditor_disposition": "Issue drafted",
                "issues_log_ref": "CTI-001",
            }
        ],
        "conclusions": {
            "d_and_i": "Effective",
            "oe": "Ineffective",
            "deficiencies_noted": True,
            "rationale": "One exception noted.",
            "issues_log_refs": ["CTI-001"],
            "testing_summary": "",
        },
        "testing_methods": {"inquiry": False, "observation": False, "inspection": True, "reperformance": False},
        "status": "complete",
        "workbook_output_id": None,
    }


def test_build_control_workbook_uses_testing_sheet_template_and_visible_cells():
    from utils.control_assurance.workbook_builder import build_control_workbook

    session = {
        "_id": "sess-s5-001",
        "title": "Q1 ITGC",
        "entity": "Tech",
        "testing_period": {"from": "2026-01-01", "to": "2026-03-31"},
        "framework": "SOX",
        "sign_off": {"preparer": {"name": "Jane Doe"}},
    }
    control = _completed_control()
    control["control_description"] = "Password parameters are configured and reviewed."
    control["walkthrough_performed"] = True
    control["sampling"] = {
        "population_description": "Active directory password policy population",
        "population_count": 24,
        "selected_size": 3,
        "selection_strategy": "random",
        "sample_period": "Q1 2026",
    }

    content = build_control_workbook(session, control, {"evidence-gridfs-001": b"Minimum length: 12"})
    workbook = openpyxl.load_workbook(filename=io.BytesIO(content), data_only=True)

    assert workbook.sheetnames == ["Test of controls", "Sample 1", "Sample 2", "Sample 3", "Auditor override"]
    sheet = workbook["Test of controls"]
    assert sheet["C4"].value == "ITGC-001 - Password Policy"
    assert sheet["C5"].value == "2026-01-01 to 2026-03-31"
    assert sheet["C6"].value == "Jane Doe"
    assert sheet["C7"].value == "Tech"
    assert sheet["C12"].value == "ITGC-001"
    assert sheet["C13"].value == "Password Policy"
    assert sheet["C14"].value == "Password parameters are configured and reviewed."
    assert sheet["C15"].value == "Preventive"
    assert sheet["M12"].value == "Y"
    assert sheet["M13"].value == "Y"
    assert sheet["M14"].value == "Effective"
    assert sheet["M15"].value == "Ineffective"
    assert sheet["C21"].value == "N"
    assert sheet["C22"].value == "N"
    assert sheet["C23"].value == "Y"
    assert sheet["C24"].value == "N"
    assert sheet["M21"].value == "Active directory password policy population"
    assert sheet["M22"].value == 24
    assert sheet["M23"].value == 3
    assert sheet["M24"].value == "random"
    assert "Inspect policy" in sheet["C33"].value
    assert sheet["B49"].value == "AD"
    assert sheet["C49"].value == "POL-002"
    assert sheet["D49"].value == "X1"
    assert "X1: Password minimum length was below baseline." in sheet["E49"].value
    assert sheet["C78"].value == "Yes"
    assert sheet["C79"].value == "Effective"
    assert sheet["C80"].value == "Ineffective"
    assert "One exception noted." in sheet["C81"].value


def test_build_control_workbook_summarises_binary_excel_evidence_without_crashing():
    from utils.control_assurance.workbook_builder import build_control_workbook

    session = {
        "_id": "sess-s5-binary",
        "title": "Q1 ITGC",
        "entity": "Tech",
        "testing_period": {"from": "2026-01-01", "to": "2026-03-31"},
        "framework": "SOX",
    }
    control = _completed_control("sess-s5-binary")
    control["evidence_files"][0]["filename"] = "workbook_ITGC_SAP_CM_01.xlsx"
    control["evidence_files"][0]["file_type"] = "excel"

    source = openpyxl.Workbook()
    source.active["A1"] = "Sample workbook evidence"
    evidence_bytes = io.BytesIO()
    source.save(evidence_bytes)

    content = build_control_workbook(
        session,
        control,
        {"evidence-gridfs-001": evidence_bytes.getvalue()},
    )
    workbook = openpyxl.load_workbook(filename=io.BytesIO(content), data_only=True)

    assert workbook.sheetnames == ["Test of controls", "Sample 1", "Sample 2", "Sample 3", "Auditor override"]
    assert "Inspect policy" in workbook["Test of controls"]["C33"].value


def test_build_control_workbook_expands_dynamic_ta_attributes_and_sampling_guidance():
    from utils.control_assurance.workbook_builder import build_control_workbook

    session = {
        "_id": "sess-s5-dynamic",
        "title": "Q1 ITGC",
        "entity": "Tech",
        "testing_period": {"from": "2026-01-01", "to": "2026-03-31"},
        "framework": "SOX",
    }
    control = _completed_control("sess-s5-dynamic")
    control["test_steps"] = [
        {
            "attribute_id": f"TA-{index:03d}",
            "label": f"TA-{index:03d}",
            "test_attribute": f"Attribute {index}",
            "description": f"Procedure {index}",
            "evidence_required": f"Evidence {index}",
        }
        for index in range(1, 12)
    ]
    control["sampling"] = {
        "population_description": "Ticket population",
        "population_count": 100,
        "adjusted_population_count": 87,
        "selected_size": 5,
        "selection_strategy": "random",
        "additional_context": "Exclude tickers beginning with 013, 014, or 105.",
    }
    control["sample_results"] = [
        {
            "sample_num": 1,
            "application": "SAP",
            "item_reference": "REQ-001",
            "step_results": [
                {"label": f"TA-{index:03d}", "tickmark": "PASS", "notes": ""}
                for index in range(1, 12)
            ],
        }
    ]

    content = build_control_workbook(session, control, {})
    workbook = openpyxl.load_workbook(filename=io.BytesIO(content), data_only=True)
    sheet = workbook["Test of controls"]

    ta_011_cells = [
        cell.coordinate
        for row in sheet.iter_rows()
        for cell in row
        if cell.value == "TA-011"
    ]
    assert ta_011_cells
    assert sheet["M25"].value == "Exclude tickers beginning with 013, 014, or 105."

    testing_header_row = next(
        row[0].row
        for row in sheet.iter_rows()
        if row[0].value == "Sample #" and row[3].value == "TA-001"
    )
    header_values = [sheet.cell(testing_header_row, col).value for col in range(4, 15)]
    assert header_values == [f"TA-{index:03d}" for index in range(1, 12)]
    assert sheet.cell(testing_header_row, 15).value == "Notes (Define any sample exceptions)"
    assert sheet.cell(testing_header_row + 1, 14).value == "√"


def test_build_control_workbook_adds_sample_evidence_tab_with_red_boxed_image():
    from PIL import Image

    from utils.control_assurance.workbook_builder import build_control_workbook

    session = {
        "_id": "sess-s5-image",
        "title": "Q1 ITGC",
        "entity": "Tech",
        "testing_period": {"from": "2026-01-01", "to": "2026-03-31"},
        "framework": "SOX",
    }
    control = _completed_control("sess-s5-image")
    control["sample_results"] = [_sample_result(1)]
    control["evidence_files"][0].update(
        {
            "filename": "policy_screenshot.png",
            "file_type": "image",
            "identified_value": "Minimum length: 12",
            "annotation_regions": [
                {
                    "step_label": "A",
                    "annotation_type": "bbox",
                    "bbox": {"x": 8, "y": 8, "w": 30, "h": 20},
                    "supporting_value": "Minimum length: 12",
                }
            ],
        }
    )

    source = Image.new("RGB", (80, 50), "white")
    image_bytes = io.BytesIO()
    source.save(image_bytes, format="PNG")

    content = build_control_workbook(
        session,
        control,
        {"evidence-gridfs-001": image_bytes.getvalue()},
    )
    workbook = openpyxl.load_workbook(filename=io.BytesIO(content), data_only=True)

    assert "Sample 1" in workbook.sheetnames
    sample_sheet = workbook["Sample 1"]
    assert sample_sheet["A1"].value == "Sample 1 Evidence"
    assert sample_sheet["B4"].value == "POL-001"
    assert sample_sheet["A8"].value == "policy_screenshot.png"
    assert sample_sheet["B9"].value == "Minimum length: 12"
    assert len(sample_sheet._images) == 1

    embedded = Image.open(io.BytesIO(sample_sheet._images[0]._data())).convert("RGB")
    assert embedded.getpixel((8, 8)) == (255, 0, 0)
    assert embedded.getpixel((38, 28)) == (255, 0, 0)


def test_workbook_contains_override_log_tab_when_overrides_exist():
    from utils.control_assurance.workbook_builder import build_control_workbook

    session = {
        "_id": "sess-s5-override",
        "title": "Q1 ITGC",
        "entity": "Tech",
        "testing_period": {"from": "2026-01-01", "to": "2026-03-31"},
        "framework": "SOX",
        "override_log": [
            {
                "control_id": "ctrl-s5-001",
                "scope": "evidence",
                "field": "ca_verification",
                "reason": "Reviewer confirmed source completeness.",
                "updated_at": "2026-05-13T00:00:00+00:00",
                "user": "Jane Doe",
            }
        ],
    }
    content = build_control_workbook(session, _completed_control("sess-s5-override"), {})
    workbook = openpyxl.load_workbook(filename=io.BytesIO(content), data_only=True)

    assert "Auditor override" in workbook.sheetnames
    assert workbook["Auditor override"]["D8"].value == 1
    assert workbook["Auditor override"]["E8"].value == "Jane Doe"
    assert workbook["Auditor override"]["F8"].value == "2026-05-13T00:00:00+00:00"
    assert workbook["Auditor override"]["G8"].value == (
        "ctrl-s5-001 | evidence | ca_verification | Reviewer confirmed source completeness."
    )


def test_generate_workbooks_creates_gridfs_ids_and_metadata(mock_db):
    from utils.control_assurance.pipeline.stage5_workbook import _generate_workbooks

    session_id = "sess-s5-001"
    control_id = "ctrl-s5-001"
    mock_db.ct_sessions.insert_one(
        {
            "_id": session_id,
            "stage": "workbook",
            "title": "Q1 ITGC",
            "entity": "Tech",
            "testing_period": {"from": "2026-01-01", "to": "2026-03-31"},
            "framework": "SOX",
        }
    )
    mock_db.ct_controls.insert_one(_completed_control(session_id, control_id))

    mock_llm = MagicMock()
    mock_llm.invoke.return_value = MagicMock(
        content=json.dumps(
            {
                "testing_summary": "Tested 3 samples and noted one exception.",
                "d_and_i_statement": "D&I was effective based on policy inspection.",
            }
        )
    )
    upload_calls = []

    def fake_upload(content: bytes, filename: str, metadata: dict) -> str:
        upload_calls.append((content, filename, metadata))
        return "gridfs-wb-001"

    with patch("utils.control_assurance.pipeline.stage5_workbook._get_db", return_value=mock_db):
        with patch("utils.control_assurance.pipeline.stage5_workbook.get_llm", return_value=mock_llm):
            with patch(
                "utils.control_assurance.pipeline.stage5_workbook.download_from_gridfs",
                return_value=b"Minimum length: 12",
            ):
                with patch(
                    "utils.control_assurance.pipeline.stage5_workbook.upload_to_gridfs",
                    side_effect=fake_upload,
                ):
                    _generate_workbooks(session_id)

    control = mock_db.ct_controls.find_one({"_id": control_id})
    assert control["workbook_output_id"] == "gridfs-wb-001"
    assert control["conclusions"]["testing_summary"] == "Tested 3 samples and noted one exception."
    assert control["conclusions"]["d_and_i_statement"] == "D&I was effective based on policy inspection."
    assert upload_calls[0][1] == "ITGC-001_Testing_Workpaper.xlsx"
    assert upload_calls[0][2] == {
        "type": "workbook_output",
        "session_id": session_id,
        "control_id": control_id,
        "filename": "ITGC-001_Testing_Workpaper.xlsx",
    }
    session = mock_db.ct_sessions.find_one({"_id": session_id})
    assert session["stage"] == "complete"
