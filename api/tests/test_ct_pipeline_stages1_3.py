from __future__ import annotations

import io
import json
from unittest.mock import MagicMock, patch

import mongomock
import openpyxl
import pytest
from fastapi.testclient import TestClient


CREATE_PAYLOAD = {
    "title": "Q1 2026 ITGC",
    "entity": "Technology",
    "testing_period": {"from": "2026-01-01", "to": "2026-03-31"},
    "framework": "SOX s.404",
    "preparer": "Jane Doe",
}

CONTROL_PAYLOAD = [
    {
        "control_id": "ITGC-001",
        "control_name": "Password Complexity Policy",
        "control_type": "Preventive",
        "domain": "Access Management",
        "test_steps": [
            {
                "label": "A",
                "description": "Inspect AD password policy",
                "evidence_required": "AD policy screenshot",
            }
        ],
    }
]


@pytest.fixture()
def mock_db():
    client = mongomock.MongoClient()
    return client["trace_db"]


@pytest.fixture()
def api_client(mock_db):
    with patch("utils.control_assurance.ct_db._get_db", return_value=mock_db):
        with patch("utils.control_assurance.ct_gridfs.get_ct_bucket"):
            from api.main import app

            yield TestClient(app)


def _make_template_bytes() -> bytes:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Control Data"
    worksheet.append(
        [
            "Control ID",
            "Control Name",
            "Control Type",
            "Domain",
            "Framework Reference",
            "Inherent Risk Rating",
            "Control Owner",
            "Frequency",
            "Prior Period Result",
            "Walkthrough Performed",
            "Sampling Mode",
            "Step A Description",
            "Step A Evidence Required",
            "Step B Description",
            "Step B Evidence Required",
        ]
    )
    worksheet.append(
        [
            "ITGC-001",
            "Password Policy",
            "Preventive",
            "Access Mgmt",
            "SOX s.404",
            "High",
            "John Smith",
            "Continuous",
            "Effective",
            "No",
            "Sample",
            "Inspect AD policy",
            "AD screenshot",
            "Check lockout",
            "Lockout config",
        ]
    )
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _make_new_control_input_bytes() -> bytes:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Control Data"
    worksheet.append(
        [
            "Control ID",
            "Risk statement",
            "Control Name/ Control Title",
            "Control description",
            "Control Type",
            "Domain / Category",
            "Control Owner",
            "Frequency",
            "Walkthrough Performed",
            "Sampling Mode",
            "Test objectives",
            "Test Steps",
            "Evidence requirements",
            "Additional Sampling guidance",
        ]
    )
    worksheet.append(
        [
            "PWD-001",
            "Weak passwords could allow unauthorized access to systems",
            "Password Complexity Policy",
            "Password complexity policy is configured and enforced on all systems",
            "Preventive",
            "Password settings",
            "IT Security Team",
            "Continuous",
            "No",
            "Walkthrough (1 sample)",
            "Verify that password complexity requirements are configured and enforced",
            (
                "1. Review the AD configuration screenshot. "
                "2. Compare minimum length against policy. "
                "3. Confirm password history is enabled. "
                "4. Test enforcement by attempting to set a weak password."
            ),
            "",
            "Exclude tickers beginning with 013, 014, or 105.",
        ]
    )
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_parse_template_creates_controls(mock_db):
    from utils.control_assurance.pipeline.stage1_parse import _parse_and_create_controls

    session_id = "sess-001"
    mock_db.ct_sessions.insert_one({"_id": session_id, "stage": "input", "stage_checkpoint": None})

    with patch("utils.control_assurance.pipeline.stage1_parse.ct_db._get_db", return_value=mock_db):
        with patch(
            "utils.control_assurance.pipeline.stage1_parse.download_from_gridfs",
            return_value=_make_template_bytes(),
        ):
            _parse_and_create_controls(session_id, "fake-gridfs-id")

    controls = list(mock_db.ct_controls.find({"session_id": session_id}))
    assert len(controls) == 1
    ctrl = controls[0]
    assert ctrl["control_id"] == "ITGC-001"
    assert ctrl["control_name"] == "Password Policy"
    assert ctrl["sampling"]["mode"] == "sample"
    assert len(ctrl["test_steps"]) == 2
    assert ctrl["test_steps"][0]["label"] == "A"
    assert ctrl["test_steps"][1]["label"] == "B"

    session = mock_db.ct_sessions.find_one({"_id": session_id})
    assert session["stage"] == "analysing"


def test_parse_new_control_input_template_creates_dynamic_ta_draft_controls(mock_db):
    from utils.control_assurance.pipeline.stage1_parse import _parse_and_create_controls

    session_id = "sess-new-template"
    mock_db.ct_sessions.insert_one({"_id": session_id, "stage": "input", "stage_checkpoint": None})

    with patch("utils.control_assurance.pipeline.stage1_parse.ct_db._get_db", return_value=mock_db):
        with patch(
            "utils.control_assurance.pipeline.stage1_parse.download_from_gridfs",
            return_value=_make_new_control_input_bytes(),
        ):
            _parse_and_create_controls(session_id, "new-template-gridfs-id")

    ctrl = mock_db.ct_controls.find_one({"session_id": session_id})
    assert ctrl["control_id"] == "PWD-001"
    assert ctrl["risk"] == "Weak passwords could allow unauthorized access to systems"
    assert ctrl["control_name"] == "Password Complexity Policy"
    assert ctrl["control_description"] == "Password complexity policy is configured and enforced on all systems"
    assert ctrl["test_objectives"] == "Verify that password complexity requirements are configured and enforced"
    assert ctrl["sampling"]["mode"] == "walkthrough"
    assert ctrl["sampling"]["additional_context"] == "Exclude tickers beginning with 013, 014, or 105."
    assert ctrl["controls_finalized"] is False
    assert len(ctrl["test_steps"]) == 4
    assert [step["attribute_id"] for step in ctrl["test_steps"]] == [
        "TA-001",
        "TA-002",
        "TA-003",
        "TA-004",
    ]
    assert ctrl["test_steps"][0]["label"] == "TA-001"
    assert ctrl["test_steps"][0]["test_attribute"]
    assert ctrl["test_steps"][0]["evidence_required"] == "AD configuration screenshot"
    assert ctrl["field_sources"]["risk"] == "user_provided"
    assert ctrl["field_sources"]["test_steps.0.evidence_required"] == "llm_derived"


def test_upload_template_dispatches_parse_task(api_client, mock_db):
    session = api_client.post("/ct/sessions", json=CREATE_PAYLOAD).json()

    with patch("api.routers.ct_v2.upload_to_gridfs", return_value="template-gridfs-id"):
        response = api_client.post(
            f"/ct/sessions/{session['id']}/upload-template",
            files=[
                (
                    "file",
                    (
                        "controls.xlsx",
                        _make_template_bytes(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    ),
                )
            ],
        )

    assert response.status_code == 202
    body = response.json()
    assert body["gridfs_id"] == "template-gridfs-id"
    assert body["status"] == "parsing"
    assert body["celery_task_id"]

    stored = mock_db.ct_sessions.find_one({"_id": session["id"]})
    assert stored["input_template_gridfs_id"] == "template-gridfs-id"
    assert stored["celery_task_id"] == body["celery_task_id"]


def test_upload_template_rejects_non_xlsx(api_client):
    session = api_client.post("/ct/sessions", json=CREATE_PAYLOAD).json()

    response = api_client.post(
        f"/ct/sessions/{session['id']}/upload-template",
        files=[("file", ("controls.csv", b"not,xlsx", "text/csv"))],
    )

    assert response.status_code == 422


def test_begin_analysis_queues_manual_controls(api_client, mock_db):
    session = api_client.post("/ct/sessions", json=CREATE_PAYLOAD).json()
    api_client.post(f"/ct/sessions/{session['id']}/controls", json=CONTROL_PAYLOAD)

    response = api_client.post(f"/ct/sessions/{session['id']}/begin-analysis")

    assert response.status_code == 202
    body = response.json()
    assert body["status"] == "analysis_queued"
    assert body["celery_task_id"]

    stored = mock_db.ct_sessions.find_one({"_id": session["id"]})
    assert stored["stage"] == "analysing"
    assert stored["celery_task_id"] == body["celery_task_id"]
    assert stored["stage_checkpoint"]["stage"] == "manual_input"
    assert stored["stage_checkpoint"]["step"] == "queued_llm_review"


def test_begin_analysis_rejects_with_no_controls(api_client):
    session = api_client.post("/ct/sessions", json=CREATE_PAYLOAD).json()

    response = api_client.post(f"/ct/sessions/{session['id']}/begin-analysis")

    assert response.status_code == 409
    assert response.json()["detail"] == "No controls available for analysis"


def _insert_review_session(mock_db, session_id: str = "sess-review"):
    mock_db.ct_sessions.insert_one(
        {
            "_id": session_id,
            "stage": "analysing",
            "title": "Q1 ITGC",
            "entity": "Tech",
            "testing_period": {"from": "2026-01-01", "to": "2026-03-31"},
            "framework": "SOX",
            "llm_suggestions": [],
            "llm_questions": [],
        }
    )
    mock_db.ct_controls.insert_one(
        {
            "_id": "ctrl-001",
            "session_id": session_id,
            "control_id": "ITGC-001",
            "control_name": "Password Policy",
            "control_type": "Preventive",
            "risk": "",
            "domain": "Access Mgmt",
            "frequency": "Continuous",
            "inherent_risk_rating": "High",
            "prior_period_result": "Effective",
            "walkthrough_performed": False,
            "sampling": {"mode": "sample"},
            "test_steps": [{"label": "A", "description": "Inspect policy"}],
        }
    )
    return session_id


def test_llm_review_writes_suggestions(mock_db):
    from utils.control_assurance.pipeline.stage2_review import _run_llm_review

    session_id = _insert_review_session(mock_db, "sess-002")
    fake_llm_response = """{
      "case_questions": [{"question_id": "q1", "question": "What is the entity scope?"}],
      "controls": [{"control_id": "ITGC-001", "suggestions": [{"suggestion_id": "s1", "text": "Consider adding step C"}], "questions": []}]
    }"""
    mock_llm = MagicMock()
    mock_llm.invoke.return_value = MagicMock(content=fake_llm_response)

    with patch("utils.control_assurance.pipeline.stage2_review._get_db", return_value=mock_db):
        with patch("utils.control_assurance.pipeline.stage2_review.get_llm", return_value=mock_llm):
            _run_llm_review(session_id)

    session = mock_db.ct_sessions.find_one({"_id": session_id})
    assert len(session["llm_questions"]) == 1
    assert session["llm_questions"][0]["question"] == "What is the entity scope?"
    assert len(session["llm_suggestions"]) == 1
    assert session["llm_suggestions"][0]["text"] == "Consider adding step C"


def test_llm_review_applies_derived_fields_and_moves_to_control_review(mock_db):
    from utils.control_assurance.pipeline.stage2_review import _run_llm_review

    session_id = _insert_review_session(mock_db, "sess-derived-review")
    fake_llm_response = """{
      "case_questions": [],
      "controls": [{
        "control_id": "ITGC-001",
        "suggestions": [],
        "questions": [],
        "derived_fields": {
          "risk": "Weak passwords could allow unauthorized access to systems.",
          "domain": "Access Management",
          "control_type": "Preventive",
          "test_steps": [{
            "attribute_id": "TA-001",
            "label": "TA-001",
            "test_attribute": "Password policy is configured according to baseline",
            "description": "Review the AD configuration screenshot.",
            "evidence_required": "AD configuration screenshot"
          }]
        }
      }]
    }"""
    mock_llm = MagicMock()
    mock_llm.invoke.return_value = MagicMock(content=fake_llm_response)

    with patch("utils.control_assurance.pipeline.stage2_review._get_db", return_value=mock_db):
        with patch("utils.control_assurance.pipeline.stage2_review.get_llm", return_value=mock_llm):
            _run_llm_review(session_id)

    control = mock_db.ct_controls.find_one({"session_id": session_id})
    assert control["risk"] == "Weak passwords could allow unauthorized access to systems."
    assert control["domain"] == "Access Management"
    assert control["control_type"] == "Preventive"
    assert control["test_steps"][0]["attribute_id"] == "TA-001"
    assert control["test_steps"][0]["test_attribute"] == "Password policy is configured according to baseline"
    assert control["test_steps"][0]["evidence_required"] == "AD configuration screenshot"
    assert control["field_sources"]["risk"] == "llm_derived"
    assert control["field_sources"]["test_steps.0.test_attribute"] == "llm_derived"

    session = mock_db.ct_sessions.find_one({"_id": session_id})
    assert session["stage"] == "control_review"
    assert session["controls_finalized"] is False
    assert session["stage_checkpoint"]["step"] == "awaiting_control_finalization"


def test_llm_review_retries_bad_json_and_records_parse_error(mock_db):
    from utils.control_assurance.pipeline.stage2_review import _run_llm_review

    session_id = _insert_review_session(mock_db, "sess-retry")
    mock_llm = MagicMock()
    mock_llm.invoke.side_effect = [
        MagicMock(content="not-json"),
        MagicMock(content='{"case_questions": [], "controls": []}'),
    ]

    with patch("utils.control_assurance.pipeline.stage2_review._get_db", return_value=mock_db):
        with patch("utils.control_assurance.pipeline.stage2_review.get_llm", return_value=mock_llm):
            _run_llm_review(session_id)

    session = mock_db.ct_sessions.find_one({"_id": session_id})
    assert len(session["parse_errors"]) == 1
    assert session["stage_checkpoint"]["step"] == "awaiting_control_finalization"
    assert mock_llm.invoke.call_count == 2


def test_get_suggestions_returns_review_state(api_client, mock_db):
    session_id = "sess-api-review"
    mock_db.ct_sessions.insert_one(
        {
            "_id": session_id,
            "llm_suggestions": [{"suggestion_id": "s1", "text": "Add step", "status": "pending"}],
            "llm_questions": [{"question_id": "q1", "question": "Scope?", "answered": False}],
        }
    )

    response = api_client.get(f"/ct/sessions/{session_id}/suggestions")

    assert response.status_code == 200
    body = response.json()
    assert body["suggestions"][0]["suggestion_id"] == "s1"
    assert body["questions"][0]["question_id"] == "q1"


def test_update_suggestion_status(api_client, mock_db):
    session_id = "sess-suggestion"
    mock_db.ct_sessions.insert_one(
        {
            "_id": session_id,
            "llm_suggestions": [{"suggestion_id": "s1", "text": "Add step", "status": "pending"}],
        }
    )

    response = api_client.patch(
        f"/ct/sessions/{session_id}/suggestions/s1",
        json={"status": "accepted"},
    )

    assert response.status_code == 200
    assert response.json()["status"] == "accepted"
    stored = mock_db.ct_sessions.find_one({"_id": session_id})
    assert stored["llm_suggestions"][0]["status"] == "accepted"


def test_answer_question_marks_answered(api_client, mock_db):
    session_id = "sess-question"
    mock_db.ct_sessions.insert_one(
        {
            "_id": session_id,
            "llm_questions": [{"question_id": "q1", "question": "Scope?", "answered": False}],
        }
    )

    response = api_client.post(
        f"/ct/sessions/{session_id}/questions/q1/answer",
        json={"answer": "Technology only"},
    )

    assert response.status_code == 200
    stored = mock_db.ct_sessions.find_one({"_id": session_id})
    assert stored["llm_questions"][0]["answer"] == "Technology only"
    assert stored["llm_questions"][0]["answered"] is True


def test_confirm_review_blocks_unanswered_questions(api_client, mock_db):
    session_id = "sess-blocked"
    mock_db.ct_sessions.insert_one(
        {
            "_id": session_id,
            "llm_questions": [{"question_id": "q1", "question": "Scope?", "answered": False}],
        }
    )

    response = api_client.post(f"/ct/sessions/{session_id}/confirm-review")

    assert response.status_code == 409


def test_confirm_review_queues_evidence_mapping(api_client, mock_db):
    session_id = "sess-confirm"
    mock_db.ct_sessions.insert_one(
        {
            "_id": session_id,
            "llm_questions": [{"question_id": "q1", "question": "Scope?", "answered": True}],
        }
    )

    response = api_client.post(f"/ct/sessions/{session_id}/confirm-review")

    assert response.status_code == 202
    body = response.json()
    assert body["status"] == "evidence_mapping_queued"
    assert body["celery_task_id"]
    stored = mock_db.ct_sessions.find_one({"_id": session_id})
    assert stored["stage"] == "population"
    assert stored["celery_task_id"] == body["celery_task_id"]


def _make_population_bytes() -> bytes:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Population"
    worksheet.append(["UserID", "Date", "Approver"])
    for idx in range(1, 6):
        worksheet.append([f"user-{idx}", f"2026-01-0{idx}", "manager"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _insert_stage3_control(mock_db, *, sampling_mode: str = "sample", evidence_passed=None):
    session_id = "sess-stage3"
    control_id = "ctrl-stage3"
    gridfs_id = "evidence-gridfs-id"
    mock_db.ct_sessions.insert_one(
        {
            "_id": session_id,
            "stage": "population",
            "testing_period": {"from": "2026-01-01", "to": "2026-03-31"},
            "override_log": [],
        }
    )
    ca_state = {
        "completeness_passed": evidence_passed,
        "accuracy_passed": evidence_passed,
        "issues": [],
        "overridden": False,
        "override_reason": None,
    }
    mock_db.ct_controls.insert_one(
        {
            "_id": control_id,
            "session_id": session_id,
            "control_id": "ITGC-001",
            "control_name": "Password Policy",
            "control_type": "Preventive",
            "domain": "Access Mgmt",
            "frequency": "Continuous",
            "inherent_risk_rating": "High",
            "prior_period_result": "Effective",
            "sampling": {
                "mode": sampling_mode,
                "population_file_id": "population-gridfs-id",
                "population_count": 0,
                "population_description": "All active users",
                "sample_period": "Jan-Mar 2026",
                "population_ca_verification": {
                    "completeness_passed": None,
                    "accuracy_passed": None,
                    "issues": [],
                    "overridden": False,
                    "override_reason": None,
                },
            },
            "test_steps": [
                {"label": "A", "description": "Inspect policy", "evidence_required": "Screenshot"}
            ],
            "evidence_files": [
                {
                    "gridfs_id": gridfs_id,
                    "filename": "policy.png",
                    "file_type": "image",
                    "mapped_step_labels": [],
                    "identified_value": "",
                    "annotation_regions": [],
                    "ca_verification": ca_state,
                }
            ],
        }
    )
    return session_id, control_id, gridfs_id


def test_evidence_ca_verification_stores_result(mock_db):
    from utils.control_assurance.pipeline.stage3_evidence import _verify_evidence_ca

    session_id, control_id, gridfs_id = _insert_stage3_control(mock_db)
    mock_llm = MagicMock()
    mock_llm.invoke.side_effect = [
        MagicMock(
            content=json.dumps(
                {
                    "completeness_passed": True,
                    "accuracy_passed": True,
                    "issues": [],
                    "identified_value": "Min length: 12",
                    "annotation_hint": "Top-right panel",
                }
            )
        ),
        MagicMock(
            content=json.dumps(
                {
                    "mapped_steps": ["A"],
                    "confidence": "high",
                    "per_step": [
                        {
                            "step_label": "A",
                            "supporting_value": "Min length: 12",
                            "annotation_type": "text_highlight",
                            "text_offset": {"start": 0, "end": 14},
                        }
                    ],
                }
            )
        ),
    ]

    with patch("utils.control_assurance.pipeline.stage3_evidence._get_db", return_value=mock_db):
        with patch("utils.control_assurance.pipeline.stage3_evidence.get_llm", return_value=mock_llm):
            with patch(
                "utils.control_assurance.pipeline.stage3_evidence.download_from_gridfs",
                return_value=b"fake",
            ):
                with patch(
                    "utils.control_assurance.pipeline.stage3_evidence.extract_file_content",
                    return_value=("text content", {}),
                ):
                    _verify_evidence_ca(session_id, control_id, gridfs_id)

    ctrl = mock_db.ct_controls.find_one({"_id": control_id})
    ev = ctrl["evidence_files"][0]
    assert ev["ca_verification"]["completeness_passed"] is True
    assert ev["ca_verification"]["accuracy_passed"] is True
    assert ev["identified_value"] == "Min length: 12"
    assert ev["mapped_step_labels"] == ["A"]
    assert ev["annotation_regions"][0]["step_label"] == "A"


def test_population_ca_and_sampling_store_conflict_fields(mock_db):
    from utils.control_assurance.pipeline.stage3_evidence import _process_population

    session_id, control_id, _ = _insert_stage3_control(mock_db)
    mock_llm = MagicMock()
    mock_llm.invoke.side_effect = [
        MagicMock(
            content=json.dumps(
                {
                    "completeness_passed": True,
                    "accuracy_passed": True,
                    "issues": [],
                    "summary": "Population looks complete.",
                }
            )
        ),
        MagicMock(
            content=json.dumps(
                {
                    "recommended_strategy": "random",
                    "recommended_size": 3,
                    "rationale": "High risk continuous control.",
                    "conflicts_with_user_input": False,
                    "conflict_reason": None,
                }
            )
        ),
    ]

    with patch("utils.control_assurance.pipeline.stage3_evidence._get_db", return_value=mock_db):
        with patch("utils.control_assurance.pipeline.stage3_evidence.get_llm", return_value=mock_llm):
            with patch(
                "utils.control_assurance.pipeline.stage3_evidence.download_from_gridfs",
                return_value=_make_population_bytes(),
            ):
                _process_population(session_id, control_id)

    ctrl = mock_db.ct_controls.find_one({"_id": control_id})
    sampling = ctrl["sampling"]
    assert sampling["population_count"] == 5
    assert sampling["population_ca_verification"]["completeness_passed"] is True
    assert sampling["llm_suggested_strategy"] == "random"
    assert sampling["llm_suggested_size"] == 3
    assert sampling["conflicts_with_user_input"] is False
    assert sampling["conflict_reason"] is None
    assert sampling["selected_size"] == 3
    assert len(sampling["selected_items"]) == 3


def test_population_ca_uses_stored_file_type_and_support_files(mock_db):
    from utils.control_assurance.pipeline.stage3_evidence import _process_population

    session_id, control_id, _ = _insert_stage3_control(mock_db)
    mock_db.ct_controls.update_one(
        {"_id": control_id},
        {
            "$set": {
                "sampling.population_file_id": "population-pdf-id",
                "sampling.population_filename": "suim_log.pdf",
                "sampling.population_file_type": "pdf",
                "sampling.population_support_files": [
                    {
                        "gridfs_id": "query-screenshot-id",
                        "filename": "suim_selection_screen.png",
                        "file_type": "image",
                        "support_type": "query_screenshot",
                        "comments": "Match unique USER_ID count to the screenshot.",
                        "reconciliation": {
                            "unique_key_columns": ["USER_ID"],
                            "expected_count": 42,
                        },
                    }
                ],
            }
        },
    )
    mock_llm = MagicMock()
    mock_llm.invoke.return_value = MagicMock(
        content=json.dumps(
            {
                "completeness_passed": False,
                "accuracy_passed": True,
                "issues": [
                    {
                        "check": "Completeness",
                        "finding": "Screenshot count must be reconciled to export.",
                        "severity": "medium",
                    }
                ],
                "summary": "Support screenshot reviewed.",
            }
        )
    )

    with patch("utils.control_assurance.pipeline.stage3_evidence._get_db", return_value=mock_db):
        with patch("utils.control_assurance.pipeline.stage3_evidence.get_llm", return_value=mock_llm):
            with patch(
                "utils.control_assurance.pipeline.stage3_evidence.download_from_gridfs",
                side_effect=[b"pdf-bytes", b"image-bytes"],
            ):
                with patch(
                    "utils.control_assurance.pipeline.stage3_evidence.extract_file_content",
                    side_effect=[
                        ("SUIM export contents", {"modDate": "2026-07-01"}),
                        ("Transaction SUIM Rows returned 42 Run timestamp 2026-07-01 09:15", {}),
                    ],
                ) as mock_extract:
                    _process_population(session_id, control_id)

    assert mock_extract.call_args_list[0].args == (b"pdf-bytes", "pdf", "suim_log.pdf")
    assert mock_extract.call_args_list[1].args == (
        b"image-bytes",
        "image",
        "suim_selection_screen.png",
    )
    prompt = mock_llm.invoke.call_args.args[0]
    assert "supporting_files" in prompt
    assert "Rows returned 42" in prompt
    assert "USER_ID" in prompt
    assert "expected_count" in prompt


def test_evidence_ca_includes_source_support_files_for_reconciliation(mock_db):
    from utils.control_assurance.pipeline.stage3_evidence import _verify_evidence_ca

    session_id, control_id, gridfs_id = _insert_stage3_control(mock_db)
    mock_db.ct_controls.update_one(
        {"_id": control_id, "evidence_files.gridfs_id": gridfs_id},
        {
            "$set": {
                "evidence_files.$.filename": "sql_export.xlsx",
                "evidence_files.$.file_type": "excel",
                "evidence_files.$.support_files": [
                    {
                        "gridfs_id": "sql-screenshot-id",
                        "filename": "query_execution.png",
                        "file_type": "image",
                        "support_type": "sql_query_screenshot",
                        "comments": "Compare unique DOCUMENT_ID count to rows returned.",
                        "reconciliation": {
                            "unique_key_columns": ["DOCUMENT_ID"],
                            "expected_count": 128,
                        },
                    }
                ],
            }
        },
    )
    mock_llm = MagicMock()
    mock_llm.invoke.return_value = MagicMock(
        content=json.dumps(
            {
                "completeness_passed": False,
                "accuracy_passed": True,
                "issues": [
                    {
                        "check": "Completeness",
                        "finding": "Export row count does not reconcile to SQL screenshot.",
                        "severity": "high",
                    }
                ],
                "identified_value": "SQL export with query support",
                "annotation_hint": "query screenshot",
            }
        )
    )

    with patch("utils.control_assurance.pipeline.stage3_evidence._get_db", return_value=mock_db):
        with patch("utils.control_assurance.pipeline.stage3_evidence.get_llm", return_value=mock_llm):
            with patch(
                "utils.control_assurance.pipeline.stage3_evidence.download_from_gridfs",
                side_effect=[b"excel-bytes", b"image-bytes"],
            ):
                with patch(
                    "utils.control_assurance.pipeline.stage3_evidence.extract_file_content",
                    side_effect=[
                        ("DOCUMENT_ID\tAMOUNT\nD1\t100\nD2\t200", {}),
                        (
                            "SELECT * FROM documents WHERE period='Q2'; Rows returned: 128; Executed at 2026-07-01 09:14",
                            {},
                        ),
                    ],
                ) as mock_extract:
                    _verify_evidence_ca(session_id, control_id, gridfs_id)

    assert mock_extract.call_args_list[0].args == (b"excel-bytes", "excel", "sql_export.xlsx")
    assert mock_extract.call_args_list[1].args == (b"image-bytes", "image", "query_execution.png")
    prompt = mock_llm.invoke.call_args.args[0]
    assert "supporting_files" in prompt
    assert "Rows returned: 128" in prompt
    assert "DOCUMENT_ID" in prompt
    assert "sql_query_screenshot" in prompt


def test_mapping_and_ca_override_endpoints_log_audit_trail(api_client, mock_db):
    session_id, control_id, gridfs_id = _insert_stage3_control(mock_db)

    mapping = api_client.patch(
        f"/ct/sessions/{session_id}/controls/{control_id}/mapping",
        json={"file_gridfs_id": gridfs_id, "mapped_step_labels": ["A"], "reason": "Manual mapping"},
    )
    ca = api_client.patch(
        f"/ct/sessions/{session_id}/controls/{control_id}/ca-override",
        json={"file_gridfs_id": gridfs_id, "reason": "Reviewer confirmed original source evidence."},
    )

    assert mapping.status_code == 200
    assert ca.status_code == 200
    ctrl = mock_db.ct_controls.find_one({"_id": control_id})
    assert ctrl["evidence_files"][0]["mapped_step_labels"] == ["A"]
    assert ctrl["evidence_files"][0]["ca_verification"]["overridden"] is True
    session = mock_db.ct_sessions.find_one({"_id": session_id})
    assert [entry["override_type"] for entry in session["override_log"]] == [
        "evidence_mapping",
        "ca_check",
    ]


def test_ca_override_rejects_reason_under_10_chars(api_client, mock_db):
    session_id, control_id, gridfs_id = _insert_stage3_control(mock_db)

    response = api_client.patch(
        f"/ct/sessions/{session_id}/controls/{control_id}/ca-override",
        json={"file_gridfs_id": gridfs_id, "reason": "short"},
    )

    assert response.status_code == 422


def test_sampling_override_writes_to_override_log(api_client, mock_db):
    session_id, control_id, _ = _insert_stage3_control(mock_db)
    mock_db.ct_controls.update_one(
        {"_id": control_id},
        {
            "$set": {
                "sampling.population_count": 10,
                "sampling.selection_strategy": "random",
                "sampling.selected_items": [1, 2, 3],
                "sampling.selected_size": 3,
            }
        },
    )

    response = api_client.patch(
        f"/ct/sessions/{session_id}/controls/{control_id}/sampling",
        json={
            "selection_strategy": "full",
            "reason": "Reviewer selected full population due to exception history.",
        },
    )

    assert response.status_code == 200
    session = mock_db.ct_sessions.find_one({"_id": session_id})
    assert session["override_log"][-1]["override_type"] == "sampling_methodology"
    assert session["override_log"][-1]["override_value"] == "full"


def test_sampling_update_stores_additional_context(api_client, mock_db):
    session_id, control_id, _ = _insert_stage3_control(mock_db)

    response = api_client.patch(
        f"/ct/sessions/{session_id}/controls/{control_id}/sampling",
        json={
            "selection_strategy": "random",
            "selected_size": 2,
            "additional_context": "Exclude tickers beginning with 013, 014, or 105.",
        },
    )

    assert response.status_code == 200
    control = mock_db.ct_controls.find_one({"_id": control_id})
    assert control["sampling"]["additional_context"] == "Exclude tickers beginning with 013, 014, or 105."
    assert control["sampling"]["selected_size"] == 2


def test_confirm_mapping_blocks_unresolved_ca(api_client, mock_db):
    session_id, _, _ = _insert_stage3_control(mock_db, sampling_mode="none", evidence_passed=None)

    response = api_client.post(f"/ct/sessions/{session_id}/confirm-mapping")

    assert response.status_code == 409
    assert response.json()["detail"]["blocked"] is True


def test_confirm_mapping_blocks_when_no_evidence_uploaded(api_client, mock_db):
    session_id, control_id, _ = _insert_stage3_control(mock_db, sampling_mode="none", evidence_passed=True)
    mock_db.ct_controls.update_one({"_id": control_id}, {"$set": {"evidence_files": []}})

    response = api_client.post(f"/ct/sessions/{session_id}/confirm-mapping")

    assert response.status_code == 409
    detail = response.json()["detail"]
    assert detail["blocked"] is True
    assert detail["unresolved"][0]["unresolved_files"][0]["filename"] == "No evidence files uploaded"


def test_confirm_mapping_blocks_when_population_ca_null(api_client, mock_db):
    session_id, control_id, _ = _insert_stage3_control(mock_db, sampling_mode="sample", evidence_passed=True)
    mock_db.ct_controls.update_one(
        {"_id": control_id},
        {"$set": {"evidence_files.0.ca_verification.completeness_passed": True}},
    )

    response = api_client.post(f"/ct/sessions/{session_id}/confirm-mapping")

    assert response.status_code == 409
    unresolved_files = response.json()["detail"]["unresolved"][0]["unresolved_files"]
    assert {"filename": "Population C&A not resolved"} in unresolved_files


def test_confirm_mapping_passes_when_population_overridden(api_client, mock_db):
    session_id, control_id, _ = _insert_stage3_control(mock_db, sampling_mode="sample", evidence_passed=True)
    mock_db.ct_controls.update_one(
        {"_id": control_id},
        {
            "$set": {
                "sampling.population_ca_verification.completeness_passed": False,
                "sampling.population_ca_verification.accuracy_passed": False,
                "sampling.population_ca_verification.overridden": True,
                "evidence_files.0.mapped_step_labels": ["A"],
            }
        },
    )

    response = api_client.post(f"/ct/sessions/{session_id}/confirm-mapping")

    assert response.status_code == 202


def test_confirm_mapping_409_body_shape(api_client, mock_db):
    session_id, _, _ = _insert_stage3_control(mock_db, sampling_mode="none", evidence_passed=None)

    response = api_client.post(f"/ct/sessions/{session_id}/confirm-mapping")

    detail = response.json()["detail"]
    assert response.status_code == 409
    assert set(detail) == {"blocked", "reason", "unresolved"}
    assert detail["blocked"] is True
    assert isinstance(detail["reason"], str)
    assert {"control_id", "control_name", "unresolved_files"}.issubset(detail["unresolved"][0])


def test_confirm_mapping_accepts_overridden_ca(api_client, mock_db):
    session_id, control_id, gridfs_id = _insert_stage3_control(
        mock_db, sampling_mode="none", evidence_passed=None
    )
    mock_db.ct_controls.update_one(
        {"_id": control_id, "evidence_files.gridfs_id": gridfs_id},
        {
            "$set": {
                "evidence_files.$.ca_verification.overridden": True,
                "evidence_files.$.mapped_step_labels": ["A"],
            }
        },
    )

    response = api_client.post(f"/ct/sessions/{session_id}/confirm-mapping")

    assert response.status_code == 202
    assert response.json()["status"] == "testing_queued"
    stored = mock_db.ct_sessions.find_one({"_id": session_id})
    assert stored["stage"] == "testing"


def test_stage3_writes_checkpoint_per_control(mock_db):
    from utils.control_assurance.pipeline.stage3_evidence import _run_evidence_mapping

    session_id = "sess-stage3-checkpoint"
    mock_db.ct_sessions.insert_one(
        {
            "_id": session_id,
            "stage": "population",
            "testing_period": {"from": "2026-01-01", "to": "2026-03-31"},
            "override_log": [],
        }
    )
    for index in range(2):
        mock_db.ct_controls.insert_one(
            {
                "_id": f"ctrl-stage3-checkpoint-{index}",
                "session_id": session_id,
                "control_id": f"ITGC-00{index}",
                "control_name": "Password Policy",
                "control_type": "Preventive",
                "domain": "Access Mgmt",
                "sampling": {"mode": "none"},
                "test_steps": [{"label": "A", "description": "Inspect policy"}],
                "evidence_files": [
                    {
                        "gridfs_id": f"evidence-gridfs-{index}",
                        "filename": "policy.png",
                        "file_type": "image",
                        "ca_verification": {
                            "completeness_passed": True,
                            "accuracy_passed": True,
                            "issues": [],
                        },
                    }
                ],
            }
        )
    original_update_one = mock_db.ct_sessions.update_one
    checkpoint_writes = []

    def recording_update_one(*args, **kwargs):
        update = args[1] if len(args) > 1 else kwargs.get("update", {})
        set_payload = update.get("$set", {})
        if "stage_checkpoint" in set_payload:
            checkpoint_writes.append(set_payload["stage_checkpoint"])
        return original_update_one(*args, **kwargs)

    with patch("utils.control_assurance.pipeline.stage3_evidence._get_db", return_value=mock_db):
        with patch("utils.control_assurance.pipeline.stage3_evidence._verify_evidence_ca"):
            with patch.object(mock_db.ct_sessions, "update_one", side_effect=recording_update_one):
                _run_evidence_mapping(session_id)

    per_control = [entry for entry in checkpoint_writes if entry.get("step", "").startswith("processing_")]
    assert len(per_control) == 2
    assert checkpoint_writes[-1]["step"] == "complete"
